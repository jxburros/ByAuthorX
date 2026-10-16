"""Review round trip through Word: send a review copy, merge the returned file as proposals and comments.

``create_review_package`` writes a DOCX review copy of the current revision with Track Changes switched
on (optionally restricted to tracked changes or comments through ``w:documentProtection``), pins that
revision as the merge base of the peer ``review:<id>`` and records the package in
``metadata.review_packages``. The copy carries the review identifier in its custom document properties,
a hidden ``_BX…`` bookmark and a ``w14:paraId`` on every paragraph, and the writer's open comments with
exact ranges and replies.

``plan_review_merge`` reads a returned copy and compares it three ways: the pinned snapshot is the base,
the returned file is theirs and the current document is ours. Each editor change becomes a pending
proposal with an exact range, kind and attribution (``w:author``/``w:date``); edits made with Track
Changes off are diffed against the snapshot and proposed as untracked. Comments and replies become
comment threads on the matching ranges. Nothing changes the text until the writer accepts a proposal,
and the writer's own later edits are never overwritten: a change whose range the writer has since
edited is proposed as a flagged conflict that replaces only that region if explicitly accepted.
Constructs the adapter cannot map (complex fields, SmartArt, charts, content controls, ink, equations,
text boxes, table structure and other revision kinds) are reported as located losses.

The returned file is untrusted data: it is read through the bounded, entity-safe archive reader and its
text never selects behaviour.
"""

from __future__ import annotations

import copy
import difflib
import hashlib
import io
import re
import tempfile
import zipfile
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

from .errors import ByaxError

W_NS = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
W = '{' + W_NS + '}'
W14 = '{http://schemas.microsoft.com/office/word/2010/wordml}'
W15_NS = 'http://schemas.microsoft.com/office/word/2012/wordml'
W15 = '{' + W15_NS + '}'
W16CID_NS = 'http://schemas.microsoft.com/office/word/2016/wordml/cid'
W16CID = '{' + W16CID_NS + '}'
W16CEX = '{http://schemas.microsoft.com/office/word/2018/wordml/cex}'
M = '{http://schemas.openxmlformats.org/officeDocument/2006/math}'
MC = '{http://schemas.openxmlformats.org/markup-compatibility/2006}'
A = '{http://schemas.openxmlformats.org/drawingml/2006/main}'
CUSTOM_NS = 'http://schemas.openxmlformats.org/officeDocument/2006/custom-properties'
VT_NS = 'http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes'
CP = '{http://schemas.openxmlformats.org/package/2006/metadata/core-properties}'
DC_TERMS = '{http://purl.org/dc/terms/}'

PEER_PREFIX = 'review:'
BOOKMARK_PREFIX = '_BX'
PROPERTY_NAMES = {
    'review_id': 'ByAuthorXReviewId',
    'document_id': 'ByAuthorXDocumentId',
    'revision': 'ByAuthorXRevision',
    'snapshot_hash': 'ByAuthorXSnapshotHash',
    'format': 'ByAuthorXReviewFormat',
}
REVIEW_FORMAT = '1'
RESTRICTIONS = {'tracked_changes': 'trackedChanges', 'comments': 'comments', 'none': None}
RSID = '00B1A7A5'
PARA = ' '  # an inserted or deleted paragraph mark inside a change
SUPPORTED_RUN_PROPS = ('bold', 'italic', 'underline', 'strikethrough', 'superscript', 'subscript')

_EXTENDED_PARTS = {
    'commentsExtended': (
        '/word/commentsExtended.xml',
        'application/vnd.openxmlformats-officedocument.wordprocessingml.commentsExtended+xml',
        'http://schemas.microsoft.com/office/2011/relationships/commentsExtended',
    ),
    'commentsIds': (
        '/word/commentsIds.xml',
        'application/vnd.openxmlformats-officedocument.wordprocessingml.commentsIds+xml',
        'http://schemas.microsoft.com/office/2016/09/relationships/commentsIds',
    ),
}
# ECMA-376 CT_Settings child order; Word refuses settings whose children are out of sequence.
_SETTINGS_ORDER = [
    'writeProtection',
    'view',
    'zoom',
    'removePersonalInformation',
    'removeDateAndTime',
    'doNotDisplayPageBoundaries',
    'displayBackgroundShape',
    'printPostScriptOverText',
    'printFractionalCharacterWidth',
    'printFormsData',
    'embedTrueTypeFonts',
    'embedSystemFonts',
    'saveSubsetFonts',
    'saveFormsData',
    'mirrorMargins',
    'alignBordersAndEdges',
    'bordersDoNotSurroundHeader',
    'bordersDoNotSurroundFooter',
    'gutterAtTop',
    'hideSpellingErrors',
    'hideGrammaticalErrors',
    'activeWritingStyle',
    'proofState',
    'formsDesign',
    'attachedTemplate',
    'linkStyles',
    'stylePaneFormatFilter',
    'stylePaneSortMethod',
    'documentType',
    'mailMerge',
    'revisionView',
    'trackRevisions',
    'doNotTrackMoves',
    'doNotTrackFormatting',
    'documentProtection',
    'autoFormatOverride',
    'styleLockTheme',
    'styleLockQFSet',
    'defaultTabStop',
]
# Revision kinds Word can write that this adapter does not map to proposals.
UNSUPPORTED_REVISIONS = {
    'cellIns': 'inserted table cell',
    'cellDel': 'deleted table cell',
    'cellMerge': 'merged table cells',
    'tblPrChange': 'table property change',
    'tblPrExChange': 'table property exception change',
    'trPrChange': 'table row property change',
    'tcPrChange': 'table cell property change',
    'tblGridChange': 'table grid change',
    'sectPrChange': 'section property change',
    'numberingChange': 'list numbering change',
    'customXmlInsRangeStart': 'custom XML insertion',
    'customXmlDelRangeStart': 'custom XML deletion',
    'customXmlMoveFromRangeStart': 'custom XML move',
    'customXmlMoveToRangeStart': 'custom XML move',
}


def _now() -> str:
    return datetime.now(UTC).replace(microsecond=0).isoformat().replace('+00:00', 'Z')


_DATE = re.compile(
    r'(19|20|21)\d\d-[01]\d-[0-3]\d(T[0-2]\d:[0-5]\d(:[0-5]\d(\.\d+)?)?(Z|[+-][0-2]\d:?[0-5]\d)?)?'
)


def _date(value) -> str:
    """An ISO 8601 date from the file, or empty (LibreOffice writes ``0-00-00T00:00:00Z`` when unknown)."""
    return value if isinstance(value, str) and _DATE.fullmatch(value) else ''


def _digest(*parts: str) -> str:
    return hashlib.sha256('\x1f'.join(parts).encode('utf-8')).hexdigest()


def unit_bookmark(block_id: str) -> str:
    """Hidden bookmark name (Word keeps names that start with ``_`` out of its bookmark list)."""
    return BOOKMARK_PREFIX + _digest('unit', block_id)[:20]


def _hex_id(*parts: str, used: set[str] | None = None) -> str:
    """An 8-digit hex identifier below 0x80000000, as Word requires for paraId and durableId."""
    salt = 0
    while True:
        value = '%08X' % (int(_digest(*parts, str(salt))[:8], 16) & 0x7FFFFFFE or 1)
        if used is None or value not in used:
            if used is not None:
                used.add(value)
            return value
        salt += 1


def _loss(code, message, target=None, **location):
    entry = {'code': code, 'severity': 'warning', 'message': message, 'target': target}
    if location:
        entry['location'] = location
    return entry


# -- Comments in DOCX (shared with the ordinary DOCX adapter) -----------------------------------


def comment_threads(root, extended=None, ids=None, extensible=None) -> dict[str, dict[str, Any]]:
    """Parse ``comments.xml`` (+ ``commentsExtended``/``commentsIds``/``commentsExtensible``) roots.

    Returns ``{w:id: {author, initials, date, text, para_id, parent, done, durable_id}}`` where
    ``parent`` is the ``w:id`` of the comment this one replies to (from ``w15:paraIdParent``).
    """
    comments, by_para = {}, {}
    for item in root.findall(W + 'comment'):
        paragraphs = item.findall('.//' + W + 'p')
        text = '\n'.join(_comment_text(p) for p in paragraphs).strip()
        para_id = paragraphs[-1].get(W14 + 'paraId') if paragraphs else None
        cid = item.get(W + 'id')
        comments[cid] = {
            'author': item.get(W + 'author', ''),
            'initials': item.get(W + 'initials', ''),
            'date': _date(item.get(W + 'date', '')),
            'text': text,
            'para_id': para_id,
            'parent': None,
            'done': False,
            'durable_id': None,
        }
        if para_id:
            by_para[para_id] = cid
    if extended is not None:
        for entry in extended.iter(W15 + 'commentEx'):
            cid = by_para.get(entry.get(W15 + 'paraId'))
            if cid is None:
                continue
            comments[cid]['done'] = entry.get(W15 + 'done') in ('1', 'true')
            parent = by_para.get(entry.get(W15 + 'paraIdParent'))
            if parent is not None and parent != cid:
                comments[cid]['parent'] = parent
    durable_dates = {}
    if extensible is not None:
        for entry in extensible.iter(W16CEX + 'commentExtensible'):
            durable_dates[entry.get(W16CEX + 'durableId')] = entry.get(W16CEX + 'dateUtc')
    if ids is not None:
        for entry in ids.iter(W16CID + 'commentId'):
            cid = by_para.get(entry.get(W16CID + 'paraId'))
            if cid is not None:
                durable = entry.get(W16CID + 'durableId')
                comments[cid]['durable_id'] = durable
                if _date(durable_dates.get(durable)) and not comments[cid]['date']:
                    comments[cid]['date'] = durable_dates[durable]
    return comments


def _comment_text(paragraph):
    parts = []
    for node in paragraph.iter():
        if node.tag == W + 't':
            parts.append(node.text or '')
        elif node.tag in (W + 'br', W + 'cr'):
            parts.append('\n')
        elif node.tag == W + 'tab':
            parts.append('\t')
    return ''.join(parts)


def comment_durable_id(comment_id: str) -> str:
    return _hex_id('durable', comment_id)


def export_comments(word, paragraphs, runs_by_block, comments, *, used_para_ids=None) -> list[dict]:
    """Write ``comments`` (with replies) onto exported runs with exact ranges and Word threading.

    Resolved threads are skipped. Replies are written as Word replies (``w15:paraIdParent`` in
    ``commentsExtended.xml``) on the same range. Each comment gets a ``w16cid:durableId`` derived from
    its ByAuthorX identity so a returned copy can be matched to the comment it came from. Returns the
    exported ``[{'id', 'docx_id', 'parent'}]``.
    """
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.text.run import Run

    used = set() if used_para_ids is None else used_para_ids
    written, extended, durable = [], [], []
    for comment in comments:
        paragraph = paragraphs.get(comment.get('target'))
        if paragraph is None or comment.get('status') == 'resolved':
            continue
        runs = _comment_runs(paragraph, runs_by_block.get(comment['target'], []), comment, OxmlElement, Run)
        thread = [(comment, None)]
        pending = [(reply, comment) for reply in comment.get('replies', [])]
        while pending:
            reply, parent = pending.pop(0)
            thread.append((reply, comment))
            pending.extend((child, reply) for child in reply.get('replies', []))
        root_para = None
        for item, parent in thread:
            made = word.add_comment(
                runs, text=str(item.get('text', '')), author=str(item.get('author', '') or 'ByAuthorX')
            )
            element = made._comment_elm
            # python-docx stamps the current time; keep output deterministic and the recorded date only.
            if item.get('date'):
                element.set(qn('w:date'), str(item['date']))
            elif qn('w:date') in element.attrib:
                del element.attrib[qn('w:date')]
            last = element.findall(qn('w:p'))[-1]
            para_id = _hex_id('comment', item['id'], used=used)
            last.set(W14 + 'paraId', para_id)
            last.set(W14 + 'textId', '77777777')
            entry = {'paraId': para_id, 'done': '0'}
            if parent is not None and root_para:
                entry['paraIdParent'] = root_para
            else:
                root_para = para_id
            extended.append(entry)
            durable.append((para_id, comment_durable_id(item['id'])))
            written.append(
                {'id': item['id'], 'docx_id': made.comment_id, 'parent': comment['id'] if parent else None}
            )
    if extended:
        _add_part(
            word,
            'commentsExtended',
            f'<w15:commentsEx xmlns:w15="{W15_NS}" xmlns:mc="{MC[1:-1]}" mc:Ignorable="w15">'
            + ''.join(
                '<w15:commentEx' + ''.join(f' w15:{key}="{value}"' for key, value in entry.items()) + '/>'
                for entry in extended
            )
            + '</w15:commentsEx>',
        )
        _add_part(
            word,
            'commentsIds',
            f'<w16cid:commentsIds xmlns:w16cid="{W16CID_NS}" xmlns:mc="{MC[1:-1]}" mc:Ignorable="w16cid">'
            + ''.join(
                f'<w16cid:commentId w16cid:paraId="{para}" w16cid:durableId="{value}"/>'
                for para, value in durable
            )
            + '</w16cid:commentsIds>',
        )
    return written


def _comment_runs(paragraph, runs, comment, OxmlElement, Run):
    """Runs covering the comment range; a point or empty range gets an empty anchor run."""
    start, end = comment.get('start'), comment.get('end')
    if not isinstance(start, int) or not isinstance(end, int):
        return paragraph.runs or [paragraph.add_run('')]
    covered = [run for a, b, run in runs if a >= start and b <= end and b > a]
    if covered and end > start:
        return covered
    following = next((run for a, b, run in runs if a >= start), None)
    if following is None:
        return [paragraph.add_run('')]
    empty = OxmlElement('w:r')
    following._r.addprevious(empty)
    return [Run(empty, paragraph)]


def _add_part(word, key, xml):
    from docx.opc.packuri import PackURI
    from docx.opc.part import Part

    partname, content_type, reltype = _EXTENDED_PARTS[key]
    blob = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n' + xml).encode('utf-8')
    part = Part(PackURI(partname), content_type, blob, word.part.package)
    word.part.relate_to(part, reltype)


# -- Review copy -------------------------------------------------------------------------------


def _clean(value):
    from .exchange import _XML_INVALID

    return _XML_INVALID.sub('�', value) if isinstance(value, str) else value


def _unit_kind(block):
    kind = block.get('type')
    if kind == 'table':
        return 'table'
    if kind == 'image':
        return 'image'
    if kind in ('page_break', 'section_break'):
        return 'break'
    return 'text'


def _review_candidate(snapshot):
    """The projection a review copy carries: ``(candidate, losses, units)``.

    ``units`` lists every block written to the copy in body order as ``{id, kind, text, cells}``; it is
    recomputed from the pinned snapshot when the copy comes back, so it must stay deterministic.
    """
    from .exchange import _exchange_losses, _xml_safe
    from .formats import FORMATS, output_document

    candidate, policy = output_document(snapshot, 'docx')
    losses = [
        _loss(
            'review_block_not_sent',
            'Excluded from the review copy by the authored output policy; it stays in your document.',
            entry.get('target'),
        )
        for entry in policy
    ]
    hidden = FORMATS['docx']['blocks']
    kept = []
    for block in candidate['blocks']:
        kind = block.get('type')
        if hidden.get(kind, 'full') == 'hidden':
            losses.append(
                _loss(
                    'review_block_not_sent',
                    f'{kind} is not sent in the Word review copy; it stays unchanged in your document.',
                    block['id'],
                )
            )
            continue
        rows = block.get('rows', [])
        if kind == 'table' and not (rows and max((len(r.get('cells', [])) for r in rows), default=0)):
            continue
        if kind == 'image' and block.get('attrs', {}).get('asset') not in candidate.get('assets', {}):
            losses.append(
                _loss(
                    'review_block_not_sent', 'Linked image without embedded bytes is not sent.', block['id']
                )
            )
            continue
        kept.append(block)
    candidate['blocks'] = kept
    texts = {b['id']: b.get('text', '') for b in kept if _unit_kind(b) == 'text'}
    review = snapshot.get('review', {})
    comments = []
    for comment in review.get('comments', []):
        if comment.get('status') == 'resolved':
            continue
        start, end = comment.get('start'), comment.get('end')
        text = texts.get(comment.get('target'))
        if (
            text is None
            or comment.get('status') == 'orphaned'
            or (isinstance(start, int) and not (isinstance(end, int) and 0 <= start <= end <= len(text)))
        ):
            losses.append(
                _loss(
                    'review_comment_not_sent',
                    'This comment is not anchored to text in the review copy and was not sent.',
                    comment.get('id'),
                )
            )
            continue
        comments.append(_clean_comment(comment))
    for change in review.get('tracked_changes', []):
        if change.get('status') == 'pending':
            losses.append(
                _loss(
                    'review_tracked_changes_flattened',
                    'A pending tracked change is sent as current text; accept or reject it in ByAuthorX.',
                    change.get('target'),
                )
            )
    if any(p.get('status') == 'pending' for p in review.get('proposals', [])):
        losses.append(
            _loss(
                'review_proposals_not_sent', 'Pending proposals stay in ByAuthorX and are not sent.', 'review'
            )
        )
    candidate['review'] = {'comments': comments, 'proposals': [], 'tracked_changes': [], 'findings': []}
    exchange_losses: list[dict] = []
    _exchange_losses(candidate, 'docx', exchange_losses)
    losses.extend(entry for entry in exchange_losses if entry['code'] != 'review_not_exported')
    candidate = _xml_safe(candidate, losses)
    units = []
    for block in candidate['blocks']:
        kind = _unit_kind(block)
        unit = {'id': block['id'], 'kind': kind, 'text': block.get('text', '') if kind == 'text' else ''}
        if kind == 'table':
            unit['cells'] = [
                [{'id': cell['id'], 'text': _clean(cell.get('text', ''))} for cell in row.get('cells', [])]
                for row in block.get('rows', [])
            ]
        units.append(unit)
    return candidate, losses, units


def _clean_comment(comment):
    cleaned = {k: _clean(v) for k, v in comment.items() if k != 'replies'}
    cleaned['replies'] = [_clean_comment(reply) for reply in comment.get('replies', [])]
    return cleaned


def snapshot_hash(units) -> str:
    from .model import canonical_json

    return hashlib.sha256(canonical_json(units).encode('utf-8')).hexdigest()


def review_identifier(doc, reviewer='') -> str:
    return 'rv_' + _digest('review', doc['id'], doc['revision'], reviewer or '')[:16]


def build_review_copy(
    snapshot, *, review_id, restrict='tracked_changes', reviewer='', created=None
) -> tuple[bytes, dict[str, Any]]:
    """DOCX bytes of the review copy and ``{losses, units, snapshot_hash, sha256}``. No I/O.

    The output depends only on the arguments, so the same snapshot and identifier give the same bytes.
    """
    from .exchange import _export_docx

    if restrict not in RESTRICTIONS:
        raise ByaxError(
            'invalid_input', 'restrict must be tracked_changes, comments or none', field='restrict'
        )
    candidate, losses, units = _review_candidate(snapshot)
    digest = snapshot_hash(units)
    properties = {
        PROPERTY_NAMES['review_id']: review_id,
        PROPERTY_NAMES['document_id']: snapshot['id'],
        PROPERTY_NAMES['revision']: snapshot['revision'],
        PROPERTY_NAMES['snapshot_hash']: digest,
        PROPERTY_NAMES['format']: REVIEW_FORMAT,
    }

    def customize(word, exported):
        if [block['id'] for block, _ in exported] != [unit['id'] for unit in units]:
            raise ByaxError('execution_failure', 'Review copy blocks do not match the review snapshot.')
        _mark_units(word, exported)
        _review_settings(word, restrict)
        _custom_properties(word, properties)
        core = word.core_properties
        core.last_modified_by = 'ByAuthorX'
        core.revision = 1
        core.comments = 'ByAuthorX review copy ' + review_id + (f' for {reviewer}' if reviewer else '')
        if created:
            stamp = datetime.fromisoformat(str(created))
            core.created = core.modified = stamp.astimezone(UTC).replace(tzinfo=None)

    buffer = io.BytesIO()
    _export_docx(candidate, buffer, customize)
    data = _normalized_zip(buffer.getvalue())
    return data, {
        'losses': losses,
        'units': units,
        'snapshot_hash': digest,
        'sha256': hashlib.sha256(data).hexdigest(),
        'comments_sent': _flatten_ids(candidate['review']['comments']),
    }


def _flatten_ids(comments):
    result, pending = [], list(comments)
    while pending:
        item = pending.pop(0)
        result.append(item['id'])
        pending.extend(item.get('replies', []))
    return result


def _mark_units(word, exported):
    """Identity bookmark, ``w14:paraId`` and rsid on every paragraph, as Word writes them."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    used: set[str] = set()
    for part in word.part.package.iter_parts():
        if str(part.partname) == '/word/comments.xml':
            used.update(p.get(W14 + 'paraId') for p in part.element.iter(qn('w:p')) if p.get(W14 + 'paraId'))
    body = word.element.body
    taken = [b.get(qn('w:id'), '') for b in body.iter(qn('w:bookmarkStart'))]
    next_id = max((int(v) for v in taken if v.isdigit()), default=-1) + 1
    owners = {}
    for block, element in exported:
        if element.tag != qn('w:p'):
            continue
        owners[element] = block['id']
        start, end = OxmlElement('w:bookmarkStart'), OxmlElement('w:bookmarkEnd')
        start.set(qn('w:id'), str(next_id))
        start.set(qn('w:name'), unit_bookmark(block['id']))
        end.set(qn('w:id'), str(next_id))
        next_id += 1
        position = 1 if len(element) and element[0].tag == qn('w:pPr') else 0
        element.insert(position, start)
        element.insert(position + 1, end)
    for index, paragraph in enumerate(body.iter(qn('w:p'))):
        owner = owners.get(paragraph)
        paragraph.set(W14 + 'paraId', _hex_id('para', owner or str(index), used=used))
        paragraph.set(W14 + 'textId', '77777777')
        paragraph.set(qn('w:rsidR'), RSID)
        paragraph.set(qn('w:rsidRDefault'), RSID)


def _settings_insert(settings, element):
    from docx.oxml.ns import qn

    local = element.tag.split('}')[-1]
    for existing in settings.findall(qn('w:' + local)):
        settings.remove(existing)
    later = set(_SETTINGS_ORDER[_SETTINGS_ORDER.index(local) + 1 :])
    known = set(_SETTINGS_ORDER)
    for child in settings:
        name = child.tag.split('}')[-1] if child.tag.startswith(W) else None
        if name is None or name in later or name not in known:
            child.addprevious(element)
            return
    settings.append(element)


def _review_settings(word, restrict):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    settings = word.settings.element
    _settings_insert(settings, OxmlElement('w:trackRevisions'))
    edit = RESTRICTIONS[restrict]
    if edit:
        protection = OxmlElement('w:documentProtection')
        protection.set(qn('w:edit'), edit)
        protection.set(qn('w:enforcement'), '1')
        _settings_insert(settings, protection)
    rsids = settings.find(qn('w:rsids'))
    if rsids is not None and not any(r.get(qn('w:val')) == RSID for r in rsids.findall(qn('w:rsid'))):
        entry = OxmlElement('w:rsid')
        entry.set(qn('w:val'), RSID)
        following = next((r for r in rsids.findall(qn('w:rsid')) if r.get(qn('w:val'), '') > RSID), None)
        if following is None:
            rsids.append(entry)
        else:
            following.addprevious(entry)


def _custom_properties(word, properties):
    import html

    from docx.opc.constants import CONTENT_TYPE, RELATIONSHIP_TYPE
    from docx.opc.packuri import PackURI
    from docx.opc.part import Part

    items = ''.join(
        f'<property fmtid="{{D5CDD505-2E9C-101B-9397-08002B2CF9AE}}" pid="{index + 2}" '
        f'name="{html.escape(name)}"><vt:lpwstr>{html.escape(str(value))}</vt:lpwstr></property>'
        for index, (name, value) in enumerate(properties.items())
    )
    blob = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        f'<Properties xmlns="{CUSTOM_NS}" xmlns:vt="{VT_NS}">{items}</Properties>'
    ).encode()
    part = Part(PackURI('/docProps/custom.xml'), CONTENT_TYPE.OFC_CUSTOM_PROPERTIES, blob, word.part.package)
    word.part.package.relate_to(part, RELATIONSHIP_TYPE.CUSTOM_PROPERTIES)


def _normalized_zip(data: bytes) -> bytes:
    """Rewrite with fixed member timestamps so identical inputs give identical bytes."""
    output = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(data)) as source,
        zipfile.ZipFile(output, 'w', zipfile.ZIP_DEFLATED) as target,
    ):
        for item in source.infolist():
            info = zipfile.ZipInfo(item.filename, date_time=(1980, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o644 << 16
            target.writestr(info, source.read(item))
    return output.getvalue()


def review_packages(doc) -> dict[str, Any]:
    """Review packages recorded in the document, with whether their snapshot pin is still held."""
    pins = doc.get('history', {}).get('peer_bases', {})
    packages = copy.deepcopy(doc.get('metadata', {}).get('review_packages', {}))
    for review_id, record in packages.items():
        record['pinned'] = pins.get(PEER_PREFIX + review_id) == record.get('revision')
    return {'review_packages': packages}


def prepare_review_package(
    doc, *, restrict='tracked_changes', reviewer='', file_name='', created=None
) -> tuple[bytes, list[dict[str, Any]], dict[str, Any]]:
    """Build a review copy of ``doc``'s current revision without writing anything.

    Returns ``(docx_bytes, operations, info)``. Applying ``operations`` through ``engine.apply`` pins
    the revision for the peer ``review:<id>`` and records the package in ``metadata.review_packages``.
    """
    from .model import snapshot

    if not isinstance(reviewer, str):
        raise ByaxError('invalid_input', 'reviewer must be text', field='reviewer')
    review_id = review_identifier(doc, reviewer)
    created = created or _now()
    data, info = build_review_copy(
        snapshot(doc), review_id=review_id, restrict=restrict, reviewer=reviewer, created=created
    )
    previous = doc.get('metadata', {}).get('review_packages', {}).get(review_id, {})
    record = {
        'id': review_id,
        'peer': PEER_PREFIX + review_id,
        'format': 'docx',
        'revision': doc['revision'],
        'snapshot_hash': info['snapshot_hash'],
        'restrict': restrict,
        'reviewer': reviewer,
        'file': file_name,
        'sha256': info['sha256'],
        'created': created,
        'status': previous.get('status', 'sent'),
        'comments_sent': len(info['comments_sent']),
        'returns': copy.deepcopy(previous.get('returns', [])),
    }
    packages = copy.deepcopy(doc.get('metadata', {}).get('review_packages', {}))
    packages[review_id] = record
    operations = [
        {'type': 'pin_merge_base', 'peer': record['peer'], 'revision': doc['revision']},
        {'type': 'set_metadata', 'values': {'review_packages': packages}},
    ]
    return (
        data,
        operations,
        {'review_id': review_id, 'record': record, 'losses': info['losses'], 'units': len(info['units'])},
    )


def create_review_package(
    doc,
    out,
    *,
    restrict='tracked_changes',
    reviewer='',
    owner=None,
    overwrite=False,
    dry_run=False,
    created=None,
) -> tuple[dict[str, Any], dict[str, Any], bytes | None]:
    """Plan the review copy for ``out``: returns ``(updated_doc, result, docx_bytes)``.

    Nothing is written here. The caller saves ``updated_doc`` (which holds the pin and the record) and
    then publishes the bytes, so a copy never exists without its recorded snapshot. On a dry run the
    original document and ``None`` come back with the planned losses.
    """
    from . import engine
    from .exchange import _report

    destination = Path(out)
    if destination.suffix.lower() != '.docx':
        raise ByaxError('invalid_input', 'A Word review copy must be a .docx file', field='out')
    if not overwrite and destination.exists() and not dry_run:
        raise ByaxError(
            'conflict', 'Output already exists; explicitly select overwrite', target=str(destination)
        )
    data, operations, info = prepare_review_package(
        doc, restrict=restrict, reviewer=reviewer, file_name=destination.name, created=created
    )
    report = _report('docx', info['losses'], [] if dry_run else [str(destination.resolve())])
    report['review'] = 'required' if info['losses'] else 'not_required'
    result = {'review_package': {**info['record'], 'units': info['units']}, 'export_report': report}
    if dry_run:
        return doc, {**result, 'dry_run': True}, None
    updated, applied = engine.apply(doc, {'expected_revision': doc['revision'], 'operations': operations})
    return updated, {**applied, **result, 'outputs': [str(destination.resolve())]}, data


def publish_review_copy(data: bytes, out, *, owner=None, overwrite=False) -> None:
    """Stage and atomically publish review copy bytes."""
    from .exchange import publish_file

    destination = Path(out)
    destination.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix='.byax-review-', dir=destination.parent) as temporary:
        stage = Path(temporary) / 'review.docx'
        stage.write_bytes(data)
        publish_file(stage, destination, overwrite=overwrite, owner=owner)


# -- Reading a returned copy -------------------------------------------------------------------

_REV_TAGS = {W + 'ins': 'ins', W + 'del': 'del', W + 'moveFrom': 'moveFrom', W + 'moveTo': 'moveTo'}
_INSERTED = ('ins', 'moveTo')
_DELETED = ('del', 'moveFrom')
_IGNORED_INLINE = {
    W + name
    for name in (
        'proofErr',
        'permStart',
        'permEnd',
        'bookmarkEnd',
        'lastRenderedPageBreak',
        'customXmlInsRangeEnd',
        'customXmlDelRangeEnd',
        'customXmlMoveFromRangeEnd',
        'customXmlMoveToRangeEnd',
    )
}
_IGNORED_RUN = {
    W + name
    for name in (
        'rPr',
        'annotationRef',
        'lastRenderedPageBreak',
        'separator',
        'continuationSeparator',
        'footnoteRef',
        'endnoteRef',
        'dayShort',
        'monthShort',
        'yearShort',
        'dayLong',
        'monthLong',
        'yearLong',
    )
}
_TOGGLES = {'b': 'bold', 'i': 'italic', 'strike': 'strikethrough'}


def _local(tag):
    return tag.rsplit('}', 1)[-1] if isinstance(tag, str) else ''


def _on(element):
    return element is not None and element.get(W + 'val', 'true') not in ('0', 'false', 'off', 'none')


def _run_props(rpr):
    props = {}
    if rpr is None:
        return props
    for tag, name in _TOGGLES.items():
        element = rpr.find(W + tag)
        if element is not None:
            props[name] = _on(element)
    underline = rpr.find(W + 'u')
    if underline is not None:
        props['underline'] = underline.get(W + 'val', 'single') != 'none'
    vertical = rpr.find(W + 'vertAlign')
    if vertical is not None:
        value = vertical.get(W + 'val')
        props['superscript'], props['subscript'] = value == 'superscript', value == 'subscript'
    other = sorted(
        _local(child.tag)
        for child in rpr
        if child.tag
        not in {W + t for t in (*_TOGGLES, 'u', 'vertAlign', 'rPrChange', 'rStyle', 'ins', 'del')}
    )
    if other:
        props['_other'] = other
    return props


def _paragraph_props(ppr):
    props = {'style': None, 'alignment': None, 'numbered': False}
    if ppr is None:
        return props
    style = ppr.find(W + 'pStyle')
    if style is not None:
        props['style'] = style.get(W + 'val')
    jc = ppr.find(W + 'jc')
    if jc is not None:
        props['alignment'] = {
            'left': 'left',
            'start': 'left',
            'center': 'center',
            'right': 'right',
            'end': 'right',
            'both': 'justify',
            'distribute': 'justify',
        }.get(jc.get(W + 'val'))
    props['numbered'] = ppr.find(W + 'numPr') is not None
    props['_other'] = sorted(
        _local(child.tag)
        for child in ppr
        if child.tag not in {W + t for t in ('pStyle', 'jc', 'numPr', 'rPr', 'pPrChange', 'sectPr')}
    )
    return props


class _BodyReader:
    """Flatten the returned body into a stream of text, paragraph marks and zero-width markers."""

    def __init__(self, rels, losses):
        self.rels, self.losses = rels, losses
        self.stream: list[dict[str, Any]] = []
        self.paragraph = -1
        self.move_names: dict[str, str | None] = {'moveFrom': None, 'moveTo': None}
        self.last_move: dict[str, str | None] = {'moveFrom': None, 'moveTo': None}
        self.move_ranges: dict[str | None, str] = {}
        self.fields: list[dict[str, Any]] = []
        self.table = -1
        self.in_table: int | None = None

    def loss(self, code, message, **location):
        where = {'table': self.in_table} if self.in_table is not None else {'paragraph': self.paragraph}
        self.losses.append(_loss(code, message, None, part='word/document.xml', **where, **location))

    def block(self, node):
        tag = node.tag
        if tag == W + 'p':
            self.read_paragraph(node)
        elif tag == W + 'tbl':
            self.table += 1
            table = self.read_table(node)
            self.stream.append({'t': 'table', 'table': table, 'el': self.paragraph, 'index': self.table})
        elif tag == W + 'sdt':
            self.loss(
                'docx_review_content_control', 'Content control: its text is compared, its binding is not.'
            )
            for child in node.findall(W + 'sdtContent'):
                for inner in child:
                    self.block(inner)
        elif tag == W + 'customXml':
            for inner in node:
                self.block(inner)
        elif tag in (W + 'sectPr', W + 'bookmarkEnd') or tag in _IGNORED_INLINE:
            return
        elif tag == W + 'altChunk':
            self.loss('docx_review_alt_chunk', 'Imported external content (altChunk) is not merged.')
        else:
            self.marker(node, None, None)

    def read_paragraph(self, node):
        self.paragraph += 1
        ppr = node.find(W + 'pPr')
        mark_element, change = None, None
        if ppr is not None:
            rpr = ppr.find(W + 'rPr')
            if rpr is not None:
                mark_element = next((child for child in rpr if child.tag in _REV_TAGS), None)
            found = ppr.find(W + 'pPrChange')
            if found is not None:
                change = {
                    'author': found.get(W + 'author', ''),
                    'date': _date(found.get(W + 'date', '')),
                    'id': found.get(W + 'id'),
                    'old': _paragraph_props(found.find(W + 'pPr')),
                }
            if ppr.find(W + 'sectPr') is not None:
                self.loss('docx_review_section_break', 'Section break and its page settings are not merged.')
        self.stream.append({'t': 'pstart', 'el': self.paragraph, 'para_id': node.get(W14 + 'paraId')})
        self.inline(node, None, None)
        # The mark's move name is known only after this paragraph's move range markers are read.
        mark_rev = self.revision(mark_element, mark=True) if mark_element is not None else None
        self.stream.append(
            {
                't': 'mark',
                'rev': mark_rev,
                'props': _paragraph_props(ppr),
                'pchange': change,
                'el': self.paragraph,
            }
        )

    def revision(self, element, mark=False):
        kind = _REV_TAGS[element.tag]
        move = None
        if kind in ('moveFrom', 'moveTo'):
            move = self.move_names[kind] or (self.last_move[kind] if mark else None)
        return {
            'kind': kind,
            'author': element.get(W + 'author', ''),
            'date': _date(element.get(W + 'date', '')),
            'id': element.get(W + 'id'),
            'move': move,
        }

    def inline(self, node, rev, link):
        for child in node:
            tag = child.tag
            if tag == W + 'pPr':
                continue
            if tag == W + 'r':
                self.run(child, rev, link)
            elif tag in _REV_TAGS:
                if tag in (W + 'moveFrom', W + 'moveTo'):
                    # LibreOffice writes the named range start inside the move wrapper.
                    start = child.find(W + tag[len(W) :] + 'RangeStart')
                    if start is not None:
                        self.marker(start, rev, link)
                inner = self.revision(child)
                if rev is not None and rev['kind'] in _INSERTED and inner['kind'] in _DELETED:
                    inner = {**inner, 'kind': 'gone'}
                self.inline(child, inner, link)
            elif tag == W + 'hyperlink':
                anchor = child.get(W + 'anchor')
                rid = child.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
                self.inline(child, rev, '#' + anchor if anchor else self.rels.get(rid, ''))
            elif tag == W + 'fldSimple':
                instruction = (child.get(W + 'instr') or '').strip()
                self.field_loss(instruction)
                self.inline(child, rev, link)
            elif tag in (W + 'smartTag', W + 'customXml', W + 'dir', W + 'bdo'):
                self.inline(child, rev, link)
            elif tag == W + 'sdt':
                self.loss(
                    'docx_review_content_control',
                    'Content control: its text is compared, its binding is not.',
                )
                for content in child.findall(W + 'sdtContent'):
                    self.inline(content, rev, link)
            elif tag in _IGNORED_INLINE:
                continue
            else:
                self.marker(child, rev, link)

    def marker(self, child, rev, link):
        tag = child.tag
        if tag == W + 'bookmarkStart':
            self.stream.append({'t': 'bookmark', 'name': child.get(W + 'name', '')})
        elif tag == W + 'commentRangeStart':
            self.stream.append({'t': 'cstart', 'id': child.get(W + 'id')})
        elif tag == W + 'commentRangeEnd':
            self.stream.append({'t': 'cend', 'id': child.get(W + 'id')})
        elif tag in (W + 'moveFromRangeStart', W + 'moveToRangeStart'):
            kind = 'moveFrom' if tag == W + 'moveFromRangeStart' else 'moveTo'
            self.move_ranges[child.get(W + 'id')] = kind
            self.move_names[kind] = self.last_move[kind] = child.get(W + 'name') or child.get(W + 'id')
        elif tag in (W + 'moveFromRangeEnd', W + 'moveToRangeEnd'):
            kind = self.move_ranges.pop(child.get(W + 'id'), None)
            if kind:
                self.move_names[kind] = None
        elif tag in (M + 'oMath', M + 'oMathPara'):
            self.loss(
                'docx_review_equation', 'Equation (Office Math) is not merged; it stays in the returned file.'
            )
        elif tag == MC + 'AlternateContent':
            self.alternate(child)
        elif _local(tag) in UNSUPPORTED_REVISIONS:
            self.loss(
                'docx_review_revision_unsupported',
                f'Tracked {UNSUPPORTED_REVISIONS[_local(tag)]} is not converted to a proposal.',
                element=_local(tag),
            )
        elif tag == W + 'subDoc':
            self.loss('docx_review_subdocument', 'Linked subdocument is not merged.')
        else:
            self.loss(
                'docx_review_unsupported_element',
                f'Unsupported Word element {_local(tag)} is not merged.',
                element=_local(tag),
            )

    def alternate(self, node):
        tags = {element.tag for element in node.iter()}
        if any(_local(tag) == 'contentPart' for tag in tags) or any(
            'ink' in str(tag).lower() for tag in tags
        ):
            self.loss('docx_review_ink', 'Ink annotation is not merged; review it in Word.')
        elif any(_local(tag) == 'txbxContent' for tag in tags):
            self.loss('docx_review_textbox', 'Text box content is not merged.')
        else:
            self.loss('docx_review_alternate_content', 'Alternate-content drawing or shape is not merged.')
        self.stream.append({'t': 'object', 'kind': 'alternate', 'rev': None, 'el': self.paragraph})

    def field_loss(self, instruction):
        name = instruction.split(' ')[0].upper() if instruction else '(empty)'
        self.loss(
            'docx_review_field',
            f'Field {name}: its displayed result is compared as text; the field code is not merged.',
            instruction=instruction[:200],
        )

    def text(self, value, rev, link, props, fmt):
        if not value or (self.fields and not self.fields[-1]['result']):
            return
        last = self.stream[-1] if self.stream else None
        if (
            last
            and last['t'] == 'text'
            and last['rev'] == rev
            and last['link'] == link
            and last['props'] == props
            and last['fmt'] == fmt
        ):
            last['text'] += value
            return
        self.stream.append(
            {
                't': 'text',
                'text': value,
                'rev': rev,
                'link': link,
                'props': props,
                'fmt': fmt,
                'el': self.paragraph,
            }
        )

    def run(self, node, rev, link):
        if rev is not None and rev['kind'] == 'gone':
            return
        rpr = node.find(W + 'rPr')
        props = _run_props(rpr)
        fmt = None
        change = rpr.find(W + 'rPrChange') if rpr is not None else None
        if change is not None:
            fmt = {
                'author': change.get(W + 'author', ''),
                'date': _date(change.get(W + 'date', '')),
                'id': change.get(W + 'id'),
                'old': _run_props(change.find(W + 'rPr')),
            }
        for child in node:
            tag = child.tag
            if tag in (W + 't', W + 'delText'):
                self.text(child.text or '', rev, link, props, fmt)
            elif tag in (W + 'tab', W + 'ptab'):
                self.text('\t', rev, link, props, fmt)
            elif tag == W + 'br':
                if child.get(W + 'type') in (None, 'textWrapping'):
                    self.text('\n', rev, link, props, fmt)
                else:
                    self.stream.append(
                        {'t': 'object', 'kind': 'page_break', 'rev': rev, 'el': self.paragraph}
                    )
            elif tag == W + 'cr':
                self.text('\n', rev, link, props, fmt)
            elif tag == W + 'noBreakHyphen':
                self.text('‑', rev, link, props, fmt)
            elif tag == W + 'softHyphen':
                self.text('­', rev, link, props, fmt)
            elif tag == W + 'fldChar':
                kind = child.get(W + 'fldCharType')
                if kind == 'begin':
                    self.fields.append({'instruction': '', 'result': False})
                elif kind == 'separate' and self.fields:
                    self.fields[-1]['result'] = True
                elif kind == 'end' and self.fields:
                    self.field_loss(self.fields.pop()['instruction'].strip())
            elif tag in (W + 'instrText', W + 'delInstrText'):
                if self.fields:
                    self.fields[-1]['instruction'] += child.text or ''
            elif tag == W + 'drawing':
                self.stream.append(
                    {'t': 'object', 'kind': _drawing_kind(child), 'rev': rev, 'el': self.paragraph}
                )
            elif tag in (W + 'pict', W + 'object'):
                self.loss('docx_review_embedded_object', 'Embedded VML or OLE object is not merged.')
                self.stream.append({'t': 'object', 'kind': 'object', 'rev': rev, 'el': self.paragraph})
            elif tag in (W + 'footnoteReference', W + 'endnoteReference'):
                self.loss('docx_review_note', 'Footnote or endnote added in Word is not merged.')
            elif tag == W + 'commentReference':
                self.stream.append({'t': 'cref', 'id': child.get(W + 'id')})
            elif tag == W + 'sym':
                self.loss('docx_review_symbol', 'Symbol-font character is not merged.')
            elif tag == MC + 'AlternateContent':
                self.alternate(child)
            elif tag in _IGNORED_RUN:
                continue
            else:
                self.loss(
                    'docx_review_unsupported_element',
                    f'Unsupported run content {_local(tag)} is not merged.',
                    element=_local(tag),
                )

    def read_table(self, node):
        self.in_table = self.table
        try:
            return self._read_table(node)
        finally:
            self.in_table = None

    def _read_table(self, node):
        for element in node.iter():
            local = _local(element.tag)
            if local in UNSUPPORTED_REVISIONS:
                self.loss(
                    'docx_review_revision_unsupported',
                    f'Tracked {UNSUPPORTED_REVISIONS[local]} is not converted to a proposal.',
                    element=local,
                )
        rows = []
        for tr in node.findall(W + 'tr'):
            trpr = tr.find(W + 'trPr')
            if trpr is not None and (trpr.find(W + 'ins') is not None or trpr.find(W + 'del') is not None):
                self.loss(
                    'docx_review_table_structure', 'Tracked table row insertion or deletion is not merged.'
                )
            cells = []
            for tc in tr.findall(W + 'tc'):
                reader = _BodyReader(self.rels, self.losses)
                reader.paragraph, reader.in_table = self.paragraph, self.in_table
                for child in tc:
                    if child.tag == W + 'tbl':
                        self.loss('docx_review_table_structure', 'Nested table is not merged.')
                    elif child.tag != W + 'tcPr':
                        reader.block(child)
                self.paragraph = reader.paragraph
                cells.append(_cell_summary(reader.stream))
            rows.append(cells)
        return {'rows': rows}


def _drawing_kind(drawing):
    for data in drawing.iter(A + 'graphicData'):
        uri = data.get('uri', '')
        if 'diagram' in uri:
            return 'smartart'
        if 'chart' in uri:
            return 'chart'
        if 'picture' in uri:
            return 'picture'
        if 'wordprocessingShape' in uri or 'wordprocessingGroup' in uri or 'wordprocessingCanvas' in uri:
            return 'shape'
        if 'ink' in uri.lower():
            return 'ink'
    return 'drawing'


def _cell_summary(stream):
    base, final, authors, dates, comments = [], [], set(), set(), []
    for item in stream:
        if item['t'] == 'text':
            rev = item['rev']
            if rev is None or rev['kind'] in _DELETED:
                base.append(item['text'])
            if rev is None or rev['kind'] in _INSERTED:
                final.append(item['text'])
            if rev is not None:
                authors.add(rev['author'])
                dates.add(rev['date'])
        elif item['t'] == 'mark':
            rev = item['rev']
            if rev is None or rev['kind'] in _DELETED:
                base.append('\n')
            if rev is None or rev['kind'] in _INSERTED:
                final.append('\n')
        elif item['t'] in ('cstart', 'cref'):
            comments.append(item['id'])
    return {
        'base': ''.join(base).removesuffix('\n'),
        'final': ''.join(final).removesuffix('\n'),
        'authors': sorted(a for a in authors if a),
        'dates': sorted(d for d in dates if d),
        'comments': comments,
    }


def read_returned(raw: bytes) -> dict[str, Any]:
    """Parse a returned review copy (bounded, entity-safe) into properties, a body stream and comments."""
    from .exchange import _parse_xml, _safe_zip

    archive = _safe_zip(raw)
    try:
        names = set(archive.namelist())

        def part(name):
            return _parse_xml(archive.read(name), name) if name in names else None

        if 'word/document.xml' not in names:
            raise ByaxError('invalid_input', 'DOCX package has no word/document.xml part')
        losses: list[dict[str, Any]] = []
        properties = {}
        custom = part('docProps/custom.xml')
        if custom is not None:
            for prop in custom.findall(f'{{{CUSTOM_NS}}}property'):
                value = next((child.text or '' for child in prop), '')
                properties[prop.get('name', '')] = value.strip()
        core = {}
        core_root = part('docProps/core.xml')
        if core_root is not None:
            for key, tag in (
                ('last_modified_by', CP + 'lastModifiedBy'),
                ('modified', DC_TERMS + 'modified'),
            ):
                element = core_root.find(tag)
                if element is not None and element.text:
                    core[key] = element.text.strip()
        rels = {}
        rel_root = part('word/_rels/document.xml.rels')
        if rel_root is not None:
            for rel in rel_root:
                if rel.get('Type', '').endswith('/hyperlink'):
                    rels[rel.get('Id')] = rel.get('Target', '')
        document = part('word/document.xml')
        body = document.find(W + 'body') if document is not None else None
        if body is None:
            raise ByaxError('invalid_input', 'DOCX document part has no body')
        reader = _BodyReader(rels, losses)
        for child in body:
            reader.block(child)
        comments = {}
        root = part('word/comments.xml')
        if root is not None:
            comments = comment_threads(
                root,
                part('word/commentsExtended.xml'),
                part('word/commentsIds.xml'),
                part('word/commentsExtensible.xml'),
            )
        for name in sorted(names):
            if re.fullmatch(r'word/(header|footer)\d*\.xml', name) or name in (
                'word/footnotes.xml',
                'word/endnotes.xml',
            ):
                root = part(name)
                if root is not None and any(e.tag in _REV_TAGS for e in root.iter()):
                    losses.append(
                        _loss(
                            'docx_review_region_changes',
                            'Tracked changes in headers, footers or notes are not merged.',
                            None,
                            part=name,
                        )
                    )
            elif name.startswith('word/vbaProject'):
                losses.append(
                    _loss(
                        'docx_review_macros', 'Macro project is ignored and never executed.', None, part=name
                    )
                )
        return {
            'properties': properties,
            'core': core,
            'stream': reader.stream,
            'comments': comments,
            'losses': losses,
        }
    finally:
        archive.close()


# -- Planning the merge ------------------------------------------------------------------------


class _Offsets:
    """Map offsets in ``before`` to ``after`` through word-level replacements."""

    def __init__(self, before, after):
        from .compare import word_changes

        self.splices = (
            []
            if before == after
            else [
                (c['before_start'], c['before_start'] + len(c['before']), c['after'])
                for c in word_changes(before, after)
            ]
        )

    def clean(self, start, end):
        from .text_merge import overlaps

        return not any(overlaps((start, end, ''), splice) for splice in self.splices)

    def start(self, position):
        """Position after any replacement ending at it (an insertion there comes first)."""
        return position + sum(
            len(r) - (b - a)
            for a, b, r in self.splices
            if b <= position and a < position or a == b == position
        )

    def end(self, position):
        """Position before an insertion exactly at it."""
        return position + sum(
            len(r) - (b - a) for a, b, r in self.splices if b < position or (b == position and a < b)
        )

    def region(self, start, end):
        """Smallest range containing ``[start, end)`` and every replacement touching it."""
        changed = True
        while changed:
            changed = False
            for a, b, _ in self.splices:
                touches = (
                    (max(a, start) < min(b, end))
                    or (a == b and start <= a <= end)
                    or (start == end and a <= start <= b)
                )
                if touches and (a < start or b > end):
                    start, end, changed = min(start, a), max(end, b), True
        outer_start = start + sum(
            len(r) - (b - a) for a, b, r in self.splices if b < start or (b == start and a < b)
        )
        outer_end = end + sum(len(r) - (b - a) for a, b, r in self.splices if b <= end)
        return start, end, outer_start, outer_end


def _short(text, limit=60):
    text = text.replace(PARA, ' ¶ ').replace('\n', ' ')
    return text if len(text) <= limit else text[: limit - 1] + '…'


def _describe(deleted, inserted):
    if deleted and inserted:
        return f'Replace “{_short(deleted)}” with “{_short(inserted)}”'
    if inserted:
        return f'Insert “{_short(inserted)}”'
    return f'Delete “{_short(deleted)}”'


def _snap(text, start, end):
    from .engine import grapheme_boundaries

    if start == end == 0 or not text:
        return start, end
    bounds = grapheme_boundaries(text)
    low = max((b for b in bounds if b <= start), default=0)
    high = min((b for b in bounds if b >= end), default=len(text))
    return low, high


def _text_ops(target, text, start, end, replacement):
    low, high = _snap(text, start, end)
    if (low, high) != (start, end):
        replacement = text[low:start] + replacement + text[end:high]
        start, end = low, high
    if start == end:
        return [{'type': 'insert_text', 'target': target, 'offset': start, 'text': replacement}]
    if not replacement:
        return [
            {
                'type': 'delete_text',
                'target': target,
                'start': start,
                'end': end,
                'expected_text': text[start:end],
            }
        ]
    match = text[start:end]
    positions = [m.start() for m in re.finditer(re.escape(match), text)]
    if start in positions:
        return [
            {
                'type': 'replace_text',
                'target': target,
                'match': match,
                'replacement': replacement,
                'occurrence': positions.index(start) + 1,
            }
        ]
    return [
        {'type': 'delete_text', 'target': target, 'start': start, 'end': end, 'expected_text': match},
        {'type': 'insert_text', 'target': target, 'offset': start, 'text': replacement},
    ]


def _whole(stream, indices, kind, name):
    """True when ``indices`` cover whole paragraphs that end with a mark of this move."""
    first, last = indices[0], indices[-1]
    if stream[last]['t'] != 'mark':
        return False
    position = first - 1
    while position >= 0 and stream[position]['t'] not in ('text', 'mark', 'table'):
        position -= 1
    if position >= 0:
        previous = stream[position]
        if previous['t'] == 'text' or (
            previous['t'] == 'mark' and previous['rev'] and previous['rev']['kind'] in _INSERTED
        ):
            return False
    allowed = {kind} | ({'ins', 'del'} if kind == 'moveTo' else set())
    for item in stream[first : last + 1]:
        if item['t'] in ('text', 'mark'):
            rev = item['rev']
            if rev is None and item is stream[last] and kind == 'moveFrom':
                continue  # LibreOffice leaves the moved paragraph's own mark unmarked
            if rev is None or rev['kind'] not in allowed or (rev['kind'] == kind and rev.get('move') != name):
                return False
            if item['t'] == 'mark' and rev['kind'] != kind:
                return False
        elif item['t'] == 'table':
            return False
    return True


def _relocate_moves(stream):
    """Turn whole-paragraph tracked moves into the moved paragraphs at their destination."""
    sides: dict[str, dict[str, list[int]]] = {}
    for index, item in enumerate(stream):
        rev = item.get('rev') if item['t'] in ('text', 'mark') else None
        if rev and rev['kind'] in ('moveFrom', 'moveTo') and rev.get('move'):
            sides.setdefault(rev['move'], {'moveFrom': [], 'moveTo': []})[rev['kind']].append(index)
    drop, relabel, before = set(), set(), {}
    for name, found in sides.items():
        source, target = found['moveFrom'], found['moveTo']
        if not source or not target:
            continue
        if stream[source[-1]]['t'] == 'text':
            following = next(
                (
                    i
                    for i in range(source[-1] + 1, len(stream))
                    if stream[i]['t'] in ('text', 'mark', 'table')
                ),
                None,
            )
            if (
                following is not None
                and stream[following]['t'] == 'mark'
                and stream[following]['rev'] is None
            ):
                source = [*source, following]
        if not (_whole(stream, source, 'moveFrom', name) and _whole(stream, target, 'moveTo', name)):
            continue
        moved_text = ''.join(stream[i]['text'] for i in source if stream[i]['t'] == 'text')
        landed = ''.join(
            stream[i]['text']
            for i in range(target[0], target[-1] + 1)
            if stream[i]['t'] == 'text' and stream[i]['rev']['kind'] in ('moveTo', 'del')
        )
        first = source[0]
        while first > 0 and stream[first - 1]['t'] not in ('text', 'mark', 'table'):
            first -= 1  # the paragraph's own bookmarks and paraId travel with it
        source_groups: list[list[dict]] = [[]]
        for index in range(first, source[-1] + 1):
            item = stream[index]
            if item['t'] in ('bookmark', 'pstart'):
                source_groups[-1].append(item)
                drop.add(index)
            elif item['t'] in ('text', 'mark'):
                drop.add(index)
                if item['t'] == 'mark':
                    source_groups.append([])
        target_starts = [target[0]] + [
            i + 1
            for i in range(target[0], target[-1])
            if stream[i]['t'] == 'mark' and stream[i]['rev']['kind'] == 'moveTo'
        ]
        if moved_text != landed or len(target_starts) != len(source_groups) - 1:
            drop -= set(range(first, source[-1] + 1))
            continue
        destination = stream[target[0]]['rev']
        for group, start in zip(source_groups, target_starts, strict=False):
            before.setdefault(start, []).extend(
                [{'t': 'moved', 'name': name, 'rev': destination}]
                + [{**item, 'moved_identity': True} for item in group]
            )
        for index in range(target[0], target[-1] + 1):
            item = stream[index]
            if item['t'] in ('text', 'mark') and item['rev'] and item['rev']['kind'] == 'moveTo':
                relabel.add(index)
    result = []
    for index, item in enumerate(stream):
        result.extend(before.get(index, []))
        if index in drop:
            continue
        result.append({**item, 'rev': None} if index in relabel else item)
    return result


def _split_units(stream):
    units: list[dict[str, Any]] = []
    current: list[dict[str, Any]] = []

    def flush():
        nonlocal current
        if any(item['t'] in ('text', 'mark') for item in current):
            units.append({'kind': 'text', 'items': current})
        current = []

    for item in stream:
        if item['t'] == 'table':
            flush()
            units.append({'kind': 'table', 'table': item['table'], 'items': [item], 'el': item['el']})
            continue
        current.append(item)
        if item['t'] == 'mark' and not (item['rev'] and item['rev']['kind'] in _INSERTED):
            flush()
    flush()
    for unit in units:
        if unit['kind'] != 'text':
            continue
        offset, base, final = 0, [], []
        for item in unit['items']:
            item['b'] = offset
            rev = item.get('rev') if item['t'] in ('text', 'mark', 'object') else None
            if item['t'] == 'text':
                if rev is None or rev['kind'] in _DELETED:
                    base.append(item['text'])
                    offset += len(item['text'])
                if rev is None or rev['kind'] in _INSERTED:
                    final.append(item['text'])
            elif item['t'] == 'mark' and (rev is None or rev['kind'] in _INSERTED):
                final.append(PARA)
        unit['base'] = ''.join(base)
        unit['final'] = ''.join(final).removesuffix(PARA)
        unit['bookmarks'] = [item['name'] for item in unit['items'] if item['t'] == 'bookmark']
        unit['moved'] = next((item for item in unit['items'] if item['t'] == 'moved'), None)
        unit['objects'] = [item for item in unit['items'] if item['t'] == 'object']
        marks = [item for item in unit['items'] if item['t'] == 'mark']
        unit['mark'] = (
            marks[-1] if marks and not (marks[-1]['rev'] and marks[-1]['rev']['kind'] in _INSERTED) else None
        )
        unit['final_marks'] = [
            item['props'] for item in marks if item['rev'] is None or item['rev']['kind'] in _INSERTED
        ]
        unit['el'] = next((item['el'] for item in unit['items'] if 'el' in item), None)
    return units


def _ratio(a, b):
    if not a and not b:
        return 1.0
    matcher = difflib.SequenceMatcher(a=a, b=b, autojunk=False)
    return matcher.ratio() if matcher.quick_ratio() >= 0.5 else 0.0


class _Planner:
    def __init__(self, doc, snapshot, units, parsed, record, source):
        self.doc, self.units, self.parsed, self.record, self.source = doc, units, parsed, record, source
        self.snapshot_blocks = {b['id']: b for b in snapshot['blocks']}
        self.snapshot = snapshot
        self.ours = {b['id']: b for b in doc['blocks']}
        self.order = [b['id'] for b in doc['blocks']]
        self.position = {identity: index for index, identity in enumerate(self.order)}
        self.losses = [dict(entry) for entry in parsed['losses']]
        self.changes: list[dict[str, Any]] = []
        self.conflicts: list[dict[str, Any]] = []
        self.untracked = parsed['core'].get('last_modified_by') or 'Reviewer'
        self.ids: set[str] = set()
        self._offsets: dict[Any, _Offsets] = {}

    # identities and coordinates

    def new_id(self, prefix, *parts):
        salt = 0
        while True:
            value = f'{prefix}_' + _digest(self.record['id'], self.source['sha256'], *parts, str(salt))[:32]
            if value not in self.ids:
                self.ids.add(value)
                return value
            salt += 1

    def snap_text(self, i):
        return self.snapshot_blocks[self.units[i]['id']].get('text', '')

    def to_ours(self, i):
        key = ('ours', i)
        if key not in self._offsets:
            ours = self.ours.get(self.units[i]['id'])
            self._offsets[key] = _Offsets(self.snap_text(i), ours.get('text', '') if ours else '')
        return self._offsets[key]

    def to_snapshot(self, r, i):
        key = ('snap', r)
        if key not in self._offsets:
            self._offsets[key] = _Offsets(self.returned[r]['base'], self.units[i]['text'])
        return self._offsets[key]

    def located(self, entry, r=None):
        """Attach the snapshot block the returned paragraph maps to, when it maps to one."""
        if entry.get('target') is None and r is not None and r in self.match:
            entry['target'] = self.units[self.match[r]]['id']
        return entry

    # main flow

    def run(self):
        stream = _relocate_moves(self.parsed['stream'])
        self.returned = _split_units(stream)
        self.align()
        paragraph_unit = {}
        for r, unit in enumerate(self.returned):
            for item in unit['items']:
                if 'el' in item:
                    paragraph_unit.setdefault(item['el'], r)
        tables = [r for r, unit in enumerate(self.returned) if unit['kind'] == 'table']
        for entry in self.losses:
            location = entry.get('location', {})
            if location.get('table') is not None and location['table'] < len(tables):
                self.located(entry, tables[location['table']])
            elif location.get('paragraph') is not None:
                self.located(entry, paragraph_unit.get(location['paragraph']))
        self.tracked_hunks()
        self.untracked_text()
        self.formatting()
        self.objects()
        self.tables()
        self.moves()
        self.structure()
        comment_ops, comment_report = self.comments()
        proposals = self.proposals()
        return proposals, comment_ops, comment_report

    def align(self):
        units, returned = self.units, self.returned
        keys = {unit_bookmark(u['id']): i for i, u in enumerate(units) if u['kind'] != 'table'}
        self.match: dict[int, int] = {}
        used: set[int] = set()
        for r, unit in enumerate(returned):
            if unit['kind'] == 'table':
                continue
            for name in unit['bookmarks']:
                i = keys.get(name)
                if i is not None and i not in used and self.compatible(unit, units[i]):
                    self.match[r] = i
                    used.add(i)
                    break
        snap_tables = [i for i, u in enumerate(units) if u['kind'] == 'table']
        ret_tables = [r for r, u in enumerate(returned) if u['kind'] == 'table']
        for r, i in zip(ret_tables, snap_tables, strict=False):
            self.match[r] = i
            used.add(i)
        if len(snap_tables) != len(ret_tables):
            self.losses.append(
                _loss(
                    'docx_review_table_structure',
                    'The returned file has a different number of tables; added or removed tables are not merged.',
                    None,
                )
            )
        free_snap = [i for i, u in enumerate(units) if i not in used and u['kind'] != 'table']
        free_ret = [r for r, u in enumerate(returned) if r not in self.match and u['kind'] != 'table']
        matcher = difflib.SequenceMatcher(
            a=[units[i]['text'] for i in free_snap], b=[returned[r]['base'] for r in free_ret], autojunk=False
        )
        for tag, a1, a2, b1, b2 in matcher.get_opcodes():
            pairs = list(zip(range(a1, a2), range(b1, b2), strict=False))
            for a, b in pairs:
                i, r = free_snap[a], free_ret[b]
                if not self.compatible(returned[r], units[i]):
                    continue
                if tag == 'equal' or _ratio(units[i]['text'], returned[r]['base']) >= 0.6:
                    self.match[r] = i

    @staticmethod
    def compatible(unit, snap):
        kinds = {item['kind'] for item in unit.get('objects', [])}
        if snap['kind'] == 'image':
            return 'picture' in kinds or 'drawing' in kinds
        if snap['kind'] == 'break':
            return 'page_break' in kinds
        return snap['kind'] == 'text'

    # change records

    def add(
        self,
        kind,
        ops,
        *,
        blocks,
        tracked,
        authors,
        dates=(),
        ids=(),
        summary,
        base=None,
        conflict=None,
        el=None,
        anchorless=False,
        pos=-1,
    ):
        authors = sorted({a for a in authors if a}) or ([self.untracked] if not tracked else ['Reviewer'])
        record = {
            'kind': kind,
            'ops': ops,
            'blocks': set(blocks),
            'tracked': tracked,
            'authors': authors,
            'dates': sorted({d for d in dates if d}),
            'revision_ids': sorted({str(x) for x in ids if x is not None}),
            'summary': summary,
            'base': base,
            'conflict': conflict,
            'el': el,
            'anchorless': anchorless,
            'pos': pos,
        }
        self.changes.append(record)
        return record

    def conflict(self, code, message, target, **details):
        entry = {'code': code, 'message': message, 'target': target, **details}
        self.conflicts.append(entry)
        return entry

    def text_change(self, i, start, end, inserted, meta):
        unit = self.units[i]
        block_id = unit['id']
        ours = self.ours.get(block_id)
        snapshot_text = self.snap_text(i)
        deleted = snapshot_text[start:end]
        base = {'target': block_id, 'start': start, 'end': end, 'text': deleted}
        if inserted == deleted:
            return
        if ours is None:
            self.conflict(
                'docx_review_conflict_block_removed',
                f'{_describe(deleted, inserted)}: the paragraph was removed after the review copy was sent.',
                block_id,
                base=base,
                inserted=inserted,
                authors=meta['authors'],
            )
            return
        if PARA in inserted:
            self.paragraph_change(i, start, end, inserted, meta, base)
            return
        offsets = self.to_ours(i)
        text = ours.get('text', '')
        conflict = None
        if offsets.clean(start, end):
            os_, oe = offsets.start(start), (offsets.end(end) if end > start else offsets.start(start))
            replacement = inserted
            if text[os_:oe] != deleted:
                conflict = {
                    'reason': 'The passage no longer matches the review snapshot.',
                    'base': deleted,
                    'ours': text[os_:oe],
                    'theirs': inserted,
                }
        else:
            rs, re_, os_, oe = offsets.region(start, end)
            replacement = snapshot_text[rs:start] + inserted + snapshot_text[end:re_]
            if text[os_:oe] == replacement:
                return  # the writer already made the same change
            conflict = {
                'reason': 'You changed this passage after sending the review copy; accepting replaces your newer text here.',
                'base': snapshot_text[rs:re_],
                'ours': text[os_:oe],
                'theirs': replacement,
            }
        self.add(
            meta.get(
                'kind', 'replacement' if deleted and inserted else 'insertion' if inserted else 'deletion'
            ),
            _text_ops(block_id, text, os_, oe, replacement),
            blocks=[block_id],
            summary=_describe(deleted, inserted),
            base=base,
            conflict=conflict,
            pos=os_,
            **{k: meta[k] for k in ('tracked', 'authors', 'dates', 'ids', 'el') if k in meta},
        )

    def new_block(self, text, props):
        style, level = 'body', None
        name = (props or {}).get('style') or ''
        heading = re.match(r'Heading\s*(\d)$', name, re.IGNORECASE)
        if heading:
            level = max(1, min(6, int(heading[1])))
        elif name in self.doc.get('styles', {}):
            style = name
        block = {
            'id': self.new_id('p', 'block', text, str(len(self.ids))),
            'type': 'heading' if level else 'paragraph',
            'text': text.replace(PARA, '\n'),
            'style': style,
            'marks': [],
            'attrs': {},
        }
        if level:
            block['attrs']['level'] = level
        if (props or {}).get('alignment'):
            block['attrs']['alignment'] = props['alignment']
        if (props or {}).get('numbered'):
            block['attrs']['list'] = 'ordered' if 'Number' in name else 'bullet'
            block['attrs']['level'] = block['attrs'].get('level', 1)
        return block

    def insert_chain(self, after, texts_props):
        ops, previous = [], after
        for text, props in texts_props:
            block = self.new_block(text, props)
            ops.append({'type': 'insert_block', 'block': block, 'after': previous})
            previous = block['id']
        return ops

    def ours_before(self, block_id):
        index = self.position.get(block_id)
        return self.order[index - 1] if index else None

    def paragraph_change(self, i, start, end, inserted, meta, base):
        block_id = self.units[i]['id']
        pieces = inserted.split(PARA)
        props = list(meta.get('marks', [])) + [meta.get('tail')]
        length = len(self.snap_text(i))
        count = len(pieces) - 1
        kind, split_at = 'paragraph_insertion', -1
        if start == end == length and pieces[0] == '':
            ops = self.insert_chain(block_id, [(pieces[j], props[j]) for j in range(1, count + 1)])
        elif start == end == 0 and pieces[-1] == '':
            ops = self.insert_chain(self.ours_before(block_id), [(pieces[j], props[j]) for j in range(count)])
        else:
            kind = 'split'
            offsets = self.to_ours(i)
            text = self.ours[block_id].get('text', '')
            if not offsets.clean(start, end):
                self.conflict(
                    'docx_review_conflict_split',
                    'A new paragraph break falls inside text you changed after sending the review copy.',
                    block_id,
                    base=base,
                    inserted=inserted,
                    authors=meta['authors'],
                )
                return
            os_, oe = offsets.start(start), (offsets.end(end) if end > start else offsets.start(start))
            ops = []
            if text[os_:oe] != pieces[0] + pieces[-1]:
                ops.extend(_text_ops(block_id, text, os_, oe, pieces[0] + pieces[-1]))
            split_at = os_ + len(pieces[0])
            ops.append(
                {
                    'type': 'split_block',
                    'target': block_id,
                    'offset': split_at,
                    'id': self.new_id('p', 'split', block_id),
                }
            )
            ops.extend(self.insert_chain(block_id, [(pieces[j], props[j]) for j in range(1, count)]))
        self.add(
            kind,
            ops,
            blocks=[block_id] if kind == 'split' else [],
            pos=split_at if kind == 'split' else -1,
            summary=(
                f'Start a new paragraph at “{_short(pieces[-1] + self.snap_text(i)[end:], 40)}”'
                + (
                    f' and insert {count - 1} paragraph(s): “{_short(PARA.join(pieces[1:-1]))}”'
                    if count > 1
                    else ''
                )
                if kind == 'split'
                else f'Insert {count} paragraph(s): “{_short(inserted.strip(PARA))}”'
            ),
            base=base,
            **{k: meta[k] for k in ('tracked', 'authors', 'dates', 'ids', 'el') if k in meta},
        )

    def span_change(self, r, start, r2, end, inserted, meta):
        """A change that crosses deleted paragraph marks, from returned unit ``r`` to ``r2``."""
        span = list(range(r, r2 + 1))
        mapped = [self.match.get(x) for x in span]
        last_is_end = r2 >= len(self.returned) or self.returned[r2]['kind'] != 'text'
        if last_is_end:
            mapped = mapped[:-1] + [None]
        if any(m is None for m in mapped[:-1]) or (not last_is_end and mapped[-1] is None):
            self.losses.append(
                _loss(
                    'docx_review_complex_change',
                    f'{_describe("", inserted) if inserted else "A deletion"} spans paragraphs that cannot be matched to your document; review it in Word.',
                    self.units[mapped[0]]['id'] if mapped[0] is not None else None,
                )
            )
            return
        ids = [self.units[m]['id'] if m is not None else None for m in mapped]
        first = ids[0]
        length = len(self.snap_text(mapped[0]))
        deleted_text = PARA.join(
            [self.snap_text(mapped[0])[start:]]
            + [self.snap_text(m) for m in mapped[1:-1]]
            + ([self.snap_text(mapped[-1])[:end]] if mapped[-1] is not None else [])
        )
        base = {'target': first, 'start': start, 'end': end, 'text': deleted_text, 'through': ids[-1]}
        whole = (start == 0 and (end == 0 or last_is_end)) or (
            start == length and (last_is_end or end == len(self.snap_text(mapped[-1])))
        )
        deleting = ids[:-1] if whole and start == 0 else ids[1:] if whole and not inserted else []
        # A paragraph the writer already removed needs no deletion; any other missing one blocks the change.
        missing = [x for x in ids if x is not None and x not in self.ours and x not in deleting]
        if missing:
            self.conflict(
                'docx_review_conflict_block_removed',
                'A change across paragraphs touches a paragraph you removed after sending the review copy.',
                missing[0],
                base=base,
                authors=meta['authors'],
            )
            return
        if whole and start == 0:
            removed = [x for x in ids[:-1] if x and x in self.ours]
            if not removed and not inserted:
                return
            ops = []
            if inserted:
                # Replacement paragraphs go where the removed ones were; text without a final mark
                # joins the start of the following paragraph.
                pieces = inserted.split(PARA)
                if last_is_end:
                    new, tail = [p for p in pieces if p], ''
                else:
                    new, tail = pieces[:-1], pieces[-1]
                if removed:
                    anchor = self.ours_before(removed[0])
                elif ids[-1]:
                    anchor = self.ours_before(ids[-1])
                else:
                    anchor = self.order[-1] if self.order else None
                ops.extend(self.insert_chain(anchor, [(p, None) for p in new]))
                if tail:
                    ops.append({'type': 'insert_text', 'target': ids[-1], 'offset': 0, 'text': tail})
            ops.extend({'type': 'delete_block', 'target': x} for x in removed)
            changed = [x for x in removed if self.ours[x] != self.snapshot_blocks.get(x)]
            self.add(
                'paragraph_deletion' if not inserted else 'paragraph_replacement',
                ops,
                blocks=removed + ([ids[-1]] if ids[-1] else []),
                summary=(f'Delete {len(removed)} paragraph(s): “{_short(deleted_text)}”'),
                base=base,
                conflict={'reason': 'You edited this paragraph after sending the review copy.'}
                if changed
                else None,
                anchorless=True,
                **{k: meta[k] for k in ('tracked', 'authors', 'dates', 'ids', 'el') if k in meta},
            )
            return
        if whole and start == length and not inserted:
            removed = [x for x in ids[1:] if x and x in self.ours]
            if not removed:
                return
            changed = [x for x in removed if self.ours[x] != self.snapshot_blocks.get(x)]
            self.add(
                'paragraph_deletion',
                [{'type': 'delete_block', 'target': x} for x in removed],
                blocks=removed,
                summary=f'Delete {len(removed)} paragraph(s): “{_short(deleted_text)}”',
                base=base,
                conflict={'reason': 'You edited this paragraph after sending the review copy.'}
                if changed
                else None,
                anchorless=True,
                **{k: meta[k] for k in ('tracked', 'authors', 'dates', 'ids', 'el') if k in meta},
            )
            return
        if PARA in inserted or last_is_end:
            self.losses.append(
                _loss(
                    'docx_review_complex_change',
                    'A change that joins paragraphs and adds new ones is not converted; review it in Word.',
                    first,
                )
            )
            return
        second = ids[-1]
        mids = [x for x in ids[1:-1] if x]
        between = (
            self.order[self.position[first] + 1 : self.position[second]]
            if self.position[second] > self.position[first]
            else None
        )
        if between is None or sorted(between) != sorted(mids):
            self.conflict(
                'docx_review_conflict_join',
                'Joining these paragraphs no longer fits your document; the paragraphs moved or new ones were added between them.',
                first,
                base=base,
                authors=meta['authors'],
            )
            return
        first_offsets, second_offsets = self.to_ours(mapped[0]), self.to_ours(mapped[-1])
        first_text, second_text = self.ours[first].get('text', ''), self.ours[second].get('text', '')
        if not (first_offsets.clean(start, length) and second_offsets.clean(0, end)):
            self.conflict(
                'docx_review_conflict_join',
                'Joining these paragraphs touches text you changed after sending the review copy.',
                first,
                base=base,
                authors=meta['authors'],
            )
            return
        ops = []
        head_end = second_offsets.end(end) if end else 0
        if head_end:
            ops.extend(_text_ops(second, second_text, 0, head_end, ''))
        tail_start = first_offsets.start(start)
        if tail_start != len(first_text) or inserted:
            ops.extend(_text_ops(first, first_text, tail_start, len(first_text), inserted))
        ops.extend({'type': 'delete_block', 'target': x} for x in mids)
        ops.append({'type': 'join_blocks', 'target': first, 'next': second})
        self.add(
            'join',
            ops,
            blocks=[first, second, *mids],
            summary=f'Join paragraphs: {_describe(deleted_text, inserted)}',
            base=base,
            anchorless=True,
            **{k: meta[k] for k in ('tracked', 'authors', 'dates', 'ids', 'el') if k in meta},
        )

    # tracked changes

    def tracked_hunks(self):
        hunks: list[dict[str, Any]] = []
        current: dict[str, Any] | None = None

        def close():
            nonlocal current
            if current is not None:
                hunks.append(current)
            current = None

        for r, unit in enumerate(self.returned):
            if unit['kind'] != 'text':
                close()
                continue
            for index, item in enumerate(unit['items']):
                kind = item['t']
                if kind not in ('text', 'mark'):
                    continue
                rev = item['rev']
                if rev is None:
                    close()
                    continue
                family = 'move' if rev['kind'] in ('moveFrom', 'moveTo') and rev.get('move') else 'edit'
                key = (family, rev['author'])
                if current is None or current['key'] != key:
                    close()
                    current = {
                        'key': key,
                        'start': (r, item['b']),
                        'end': (r, item['b']),
                        'inserted': '',
                        'revs': [],
                        'marks': [],
                        'tail': None,
                        'el': item.get('el'),
                    }
                current['revs'].append(rev)
                current['last'] = (r, index)
                if kind == 'text':
                    if rev['kind'] in _DELETED:
                        current['end'] = (r, item['b'] + len(item['text']))
                    else:
                        current['inserted'] += item['text']
                elif rev['kind'] in _INSERTED:
                    current['inserted'] += PARA
                    current['marks'].append(item['props'])
                else:
                    current['end'] = (r + 1, 0)
        close()
        for hunk in hunks:
            # The paragraph mark that ends the last inserted piece gives that paragraph's properties.
            r_last, index = hunk['last']
            following = (
                item
                for unit in self.returned[r_last:]
                if unit['kind'] == 'text'
                for item in (unit['items'][index + 1 :] if unit is self.returned[r_last] else unit['items'])
            )
            hunk['tail'] = next((item['props'] for item in following if item['t'] == 'mark'), None)
            (r, start), (r2, end) = hunk['start'], hunk['end']
            revs = hunk['revs']
            kinds = {rev['kind'] for rev in revs}
            meta = {
                'tracked': True,
                'authors': [rev['author'] for rev in revs],
                'dates': [rev['date'] for rev in revs],
                'ids': [rev['id'] for rev in revs],
                'marks': hunk['marks'],
                'tail': hunk['tail'],
                'el': hunk['el'],
            }
            if kinds & {'moveFrom', 'moveTo'}:
                meta['kind'] = 'move_from' if kinds <= {'moveFrom'} else 'move_to'
            if r not in self.match:
                continue  # an untracked new paragraph carries its own final text
            i = self.match[r]
            to_snapshot = self.to_snapshot(r, i)
            snap_start = to_snapshot.start(start)
            if r2 == r:
                snap_end = to_snapshot.end(end) if end > start else snap_start
                self.text_change(i, snap_start, snap_end, hunk['inserted'], meta)
                continue
            if r2 < len(self.returned) and r2 in self.match and self.returned[r2]['kind'] == 'text':
                end = self.to_snapshot(r2, self.match[r2]).end(end) if end else 0
            self.span_change(r, snap_start, r2, end, hunk['inserted'], meta)

    def untracked_text(self):
        from .compare import word_changes

        for r, i in sorted(self.match.items()):
            unit = self.returned[r]
            if unit['kind'] != 'text' or self.units[i]['kind'] != 'text':
                continue
            snapshot_text = self.units[i]['text']
            if unit['base'] == snapshot_text:
                continue
            for change in word_changes(snapshot_text, unit['base']):
                start = change['before_start']
                self.text_change(
                    i,
                    start,
                    start + len(change['before']),
                    change['after'],
                    {'tracked': False, 'authors': [self.untracked], 'el': unit['el']},
                )

    def formatting(self):
        for r, i in sorted(self.match.items()):
            unit = self.returned[r]
            if unit['kind'] != 'text':
                continue
            groups: list[dict[str, Any]] = []
            for item in unit['items']:
                if item['t'] != 'text' or item['rev'] is not None or not item.get('fmt'):
                    continue
                fmt = item['fmt']
                delta = {
                    key: item['props'].get(key, False)
                    for key in SUPPORTED_RUN_PROPS
                    if item['props'].get(key, False) != fmt['old'].get(key, False)
                }
                other = set(item['props'].get('_other', [])) ^ set(fmt['old'].get('_other', []))
                if other:
                    self.losses.append(
                        self.located(
                            _loss(
                                'docx_review_formatting_unsupported',
                                'Tracked formatting change ('
                                + ', '.join(sorted(other))
                                + ') is not converted.',
                                None,
                                paragraph=item.get('el'),
                            ),
                            r,
                        )
                    )
                if not delta:
                    continue
                last = groups[-1] if groups else None
                if (
                    last
                    and last['delta'] == delta
                    and last['author'] == fmt['author']
                    and last['end'] == item['b']
                ):
                    last['end'] = item['b'] + len(item['text'])
                    last['ids'].append(fmt['id'])
                    continue
                groups.append(
                    {
                        'delta': delta,
                        'author': fmt['author'],
                        'date': fmt['date'],
                        'ids': [fmt['id']],
                        'start': item['b'],
                        'end': item['b'] + len(item['text']),
                        'el': item.get('el'),
                    }
                )
            if not groups:
                continue
            to_snapshot = self.to_snapshot(r, i)
            block_id = self.units[i]['id']
            ours = self.ours.get(block_id)
            for group in groups:
                start, end = to_snapshot.start(group['start']), to_snapshot.end(group['end'])
                label = ', '.join(f'{"" if v else "remove "}{k}' for k, v in sorted(group['delta'].items()))
                text = self.snap_text(i)[start:end]
                if ours is None:
                    self.conflict(
                        'docx_review_conflict_block_removed',
                        f'Formatting ({label}) on a removed paragraph.',
                        block_id,
                    )
                    continue
                offsets = self.to_ours(i)
                if not offsets.clean(start, end):
                    self.conflict(
                        'docx_review_conflict_formatting',
                        f'Formatting ({label}) applies to text you changed after sending the review copy.',
                        block_id,
                        base={'target': block_id, 'start': start, 'end': end, 'text': text},
                    )
                    continue
                os_, oe = offsets.start(start), offsets.end(end)
                os_, oe = _snap(ours.get('text', ''), os_, oe)
                self.add(
                    'formatting',
                    [
                        {
                            'type': 'format_text',
                            'target': block_id,
                            'start': os_,
                            'end': oe,
                            'marks': dict(group['delta']),
                        }
                    ],
                    blocks=[block_id],
                    tracked=True,
                    authors=[group['author']],
                    dates=[group['date']],
                    ids=group['ids'],
                    summary=f'Format “{_short(text)}”: {label}',
                    base={'target': block_id, 'start': start, 'end': end, 'text': text},
                    el=group['el'],
                )

    def objects(self):
        for r, unit in enumerate(self.returned):
            if unit['kind'] != 'text':
                continue
            i = self.match.get(r)
            snap_kind = self.units[i]['kind'] if i is not None else None
            for item in unit['objects']:
                rev = item['rev']
                expected = (snap_kind == 'image' and item['kind'] in ('picture', 'drawing')) or (
                    snap_kind == 'break' and item['kind'] == 'page_break'
                )
                if expected and rev is not None and rev['kind'] in _DELETED and i is not None:
                    block_id = self.units[i]['id']
                    if block_id in self.ours:
                        self.add(
                            'object_deletion',
                            [{'type': 'delete_block', 'target': block_id}],
                            blocks=[block_id],
                            tracked=True,
                            authors=[rev['author']],
                            dates=[rev['date']],
                            ids=[rev['id']],
                            summary='Delete the ' + ('image' if snap_kind == 'image' else 'page break'),
                            anchorless=True,
                            el=item['el'],
                        )
                    continue
                if expected or item['kind'] == 'alternate' or (rev is not None and rev['kind'] in _DELETED):
                    continue
                names = {
                    'smartart': (
                        'docx_review_smartart',
                        'SmartArt diagram is not merged; it stays in the returned file.',
                    ),
                    'chart': ('docx_review_chart', 'Chart is not merged; it stays in the returned file.'),
                    'ink': ('docx_review_ink', 'Ink annotation is not merged; review it in Word.'),
                    'shape': ('docx_review_shape', 'Shape or text box is not merged.'),
                    'page_break': (
                        'docx_review_page_break',
                        'A page or column break added in Word is not merged.',
                    ),
                    'object': ('docx_review_embedded_object', 'Embedded object is not merged.'),
                }
                code, message = names.get(
                    item['kind'],
                    ('docx_review_image', 'An image added in Word is not merged; insert it in ByAuthorX.'),
                )
                self.losses.append(
                    self.located(
                        _loss(code, message, None, part='word/document.xml', paragraph=item['el']), r
                    )
                )

    def tables(self):
        for r, unit in enumerate(self.returned):
            if unit['kind'] != 'table' or r not in self.match:
                continue
            snap = self.units[self.match[r]]
            rows = unit['table']['rows']
            if [len(row) for row in rows] != [len(row) for row in snap['cells']]:
                self.losses.append(
                    _loss(
                        'docx_review_table_structure',
                        'Rows or columns were added or removed in Word; table structure changes are not merged.',
                        snap['id'],
                    )
                )
                continue
            for row_index, (returned_row, snap_row) in enumerate(zip(rows, snap['cells'], strict=False)):
                for column, (cell, snap_cell) in enumerate(zip(returned_row, snap_row, strict=False)):
                    if cell['final'] == snap_cell['text'] and cell['base'] == snap_cell['text']:
                        continue
                    if cell['final'] == snap_cell['text']:
                        continue
                    current = self._cell_text(snap['id'], snap_cell['id'])
                    raw = self._snapshot_cell_text(snap['id'], snap_cell['id'])
                    if current is None:
                        self.conflict(
                            'docx_review_conflict_block_removed',
                            'Edited table cell no longer exists.',
                            snap_cell['id'],
                        )
                        continue
                    conflict = None
                    if current != raw:
                        if current == cell['final']:
                            continue
                        conflict = {
                            'reason': 'You changed this cell after sending the review copy; accepting replaces it.',
                            'base': raw,
                            'ours': current,
                            'theirs': cell['final'],
                        }
                    self.add(
                        'table_cell',
                        [
                            {
                                'type': 'update_cell',
                                'target': snap_cell['id'],
                                'text': cell['final'],
                                'expected_text': current,
                            }
                        ],
                        blocks=[snap['id']],
                        tracked=bool(cell['authors']),
                        authors=cell['authors'] or [self.untracked],
                        dates=cell['dates'],
                        summary=f'Table row {row_index + 1}, column {column + 1}: {_describe(raw, cell["final"])}',
                        base={'target': snap_cell['id'], 'text': raw},
                        conflict=conflict,
                        el=unit['el'],
                    )

    def _cell_text(self, table_id, cell_id):
        table = self.ours.get(table_id)
        for row in (table or {}).get('rows', []):
            for cell in row.get('cells', []):
                if cell['id'] == cell_id:
                    return cell.get('text', '')
        return None

    def _snapshot_cell_text(self, table_id, cell_id):
        for row in self.snapshot_blocks.get(table_id, {}).get('rows', []):
            for cell in row.get('cells', []):
                if cell['id'] == cell_id:
                    return cell.get('text', '')
        return ''

    # structure

    def inserted_at(self, r, position):
        """True when returned unit ``r`` gained new paragraphs at its start or end."""
        unit = self.returned[r]
        if unit['kind'] != 'text':
            return False
        edge = 0 if position == 'start' else len(unit['base'])
        return any(
            item['t'] == 'mark' and item['rev'] and item['rev']['kind'] in _INSERTED and item['b'] == edge
            for item in unit['items']
        )

    def destination(self, r, skip):
        """Where returned unit ``r`` sits in ours: next to a matched neighbour that gained no new
        paragraphs on that side, so pending insertions keep their place whichever is accepted first."""
        previous = r - 1
        following = r + 1
        if (
            previous in self.match
            and following in self.match
            and self.inserted_at(previous, 'end')
            and not self.inserted_at(following, 'start')
        ):
            identity = self.units[self.match[following]]['id']
            if identity in self.ours and identity not in skip:
                return {'before': identity}
        for previous in range(r - 1, -1, -1):
            i = self.match.get(previous)
            if i is not None and self.units[i]['id'] in self.ours and self.units[i]['id'] not in skip:
                return {'after': self.units[i]['id']}
        for following in range(r + 1, len(self.returned)):
            i = self.match.get(following)
            if i is not None and self.units[i]['id'] in self.ours and self.units[i]['id'] not in skip:
                return {'before': self.units[i]['id']}
        return None

    def moves(self):
        from .compare import _moved

        tracked = {r for r, unit in enumerate(self.returned) if unit.get('moved') and r in self.match}
        order = [self.units[i]['id'] for r, i in sorted(self.match.items()) if r not in tracked]
        snap_order = [u['id'] for u in self.units if u['id'] in set(order)]
        untracked = set(_moved(snap_order, order))
        for r, i in sorted(self.match.items()):
            block_id = self.units[i]['id']
            if r not in tracked and block_id not in untracked:
                continue
            if block_id not in self.ours:
                self.conflict(
                    'docx_review_conflict_block_removed',
                    'A moved paragraph was removed after the review copy was sent.',
                    block_id,
                )
                continue
            where = self.destination(r, {block_id})
            if where is None:
                continue
            index = self.position[block_id]
            if 'after' in where and index > 0 and self.order[index - 1] == where['after']:
                continue
            if 'before' in where and index + 1 < len(self.order) and self.order[index + 1] == where['before']:
                continue
            moved = self.returned[r].get('moved')
            rev = moved['rev'] if moved else None
            self.add(
                'move',
                [{'type': 'move_block', 'target': block_id, **where}],
                blocks=[block_id],
                tracked=rev is not None,
                authors=[rev['author']] if rev else [self.untracked],
                dates=[rev['date']] if rev else [],
                ids=[rev['id']] if rev else [],
                summary=f'Move paragraph “{_short(self.snap_text(i))}”'
                + (
                    f' after “{_short(self.ours[where["after"]].get("text", ""), 30)}”'
                    if 'after' in where
                    else f' before “{_short(self.ours[where["before"]].get("text", ""), 30)}”'
                ),
                base={'target': block_id, 'move': moved['name'] if moved else None},
                anchorless=True,
                el=self.returned[r]['el'],
            )

    def structure(self):
        matched = set(self.match.values())
        for i, unit in enumerate(self.units):
            if i in matched or unit['kind'] == 'table':
                continue
            block_id = unit['id']
            if block_id not in self.ours:
                continue
            changed = self.ours[block_id] != self.snapshot_blocks.get(block_id)
            self.add(
                'paragraph_deletion',
                [{'type': 'delete_block', 'target': block_id}],
                blocks=[block_id],
                tracked=False,
                authors=[self.untracked],
                summary=f'Delete paragraph “{_short(unit["text"])}” (removed in Word without Track Changes)',
                base={'target': block_id, 'start': 0, 'end': len(unit['text']), 'text': unit['text']},
                conflict={'reason': 'You edited this paragraph after sending the review copy.'}
                if changed
                else None,
                anchorless=True,
            )
        r = 0
        while r < len(self.returned):
            unit = self.returned[r]
            if r in self.match or unit['kind'] != 'text':
                r += 1
                continue
            group = []
            while r < len(self.returned) and r not in self.match and self.returned[r]['kind'] == 'text':
                group.append(self.returned[r])
                r += 1
            paragraphs = []
            for item in group:
                pieces = item['final'].split(PARA)
                paragraphs.extend(zip(pieces, item['final_marks'] + [None] * len(pieces), strict=False))
            where = self.destination(r - len(group), set())
            if where is None:
                anchor = self.order[-1] if self.order else None
            elif 'after' in where:
                anchor = where['after']
            else:
                anchor = self.ours_before(where['before'])
            text = PARA.join(p for p, _ in paragraphs)
            if not text.replace(PARA, '').strip():
                continue  # empty paragraphs (or ones holding only a reported object) add nothing
            self.add(
                'paragraph_insertion',
                self.insert_chain(anchor, paragraphs),
                blocks=[],
                tracked=False,
                authors=[self.untracked],
                summary=f'Insert {len(paragraphs)} paragraph(s) added in Word without Track Changes: “{_short(text)}”',
                el=group[0]['el'],
            )
        for r, i in sorted(self.match.items()):
            unit = self.returned[r]
            mark = unit.get('mark') if unit['kind'] == 'text' else None
            if not mark or not mark.get('pchange'):
                continue
            block_id = self.units[i]['id']
            if block_id not in self.ours:
                continue
            old, new, change = mark['pchange']['old'], mark['props'], mark['pchange']
            values = {}
            old_heading = re.match(r'Heading\s*(\d)$', old.get('style') or '', re.IGNORECASE)
            new_heading = re.match(r'Heading\s*(\d)$', new.get('style') or '', re.IGNORECASE)
            if (old_heading and old_heading[1]) != (new_heading and new_heading[1]):
                values['level'] = int(new_heading[1]) if new_heading else None
            if old.get('alignment') != new.get('alignment'):
                values['alignment'] = new.get('alignment') or 'left'
            ignored = set(new.get('_other', [])) ^ set(old.get('_other', []))
            if (
                (old.get('style') != new.get('style') and 'level' not in values)
                or ignored
                or old.get('numbered') != new.get('numbered')
            ):
                self.losses.append(
                    _loss(
                        'docx_review_paragraph_format',
                        'Tracked paragraph formatting (style, spacing, indentation or numbering) is not converted.',
                        block_id,
                        paragraph=mark['el'],
                    )
                )
            if values:
                self.add(
                    'paragraph_format',
                    [{'type': 'set_attributes', 'target': block_id, 'values': values}],
                    blocks=[block_id],
                    tracked=True,
                    authors=[change['author']],
                    dates=[change['date']],
                    ids=[change['id']],
                    summary='Paragraph format: ' + ', '.join(f'{k} {v}' for k, v in values.items()),
                    base={'target': block_id},
                    anchorless=True,
                    el=mark['el'],
                )

    # comments

    def comment_position(self, cid, positions):
        found = positions.get(cid, {})
        start = found.get('start') or found.get('ref')
        end = found.get('end') or found.get('ref') or start
        return start, end

    def comments(self):
        threads = self.parsed['comments']
        report = {'added': 0, 'replies': 0, 'resolved': 0, 'existing': 0, 'marked_done': []}
        if not threads:
            return [], report
        positions: dict[str, dict[str, Any]] = {}
        cells: dict[str, str] = {}
        for r, unit in enumerate(self.returned):
            if unit['kind'] == 'table':
                if r in self.match:
                    snap = self.units[self.match[r]]
                    for row, snap_row in zip(unit['table']['rows'], snap['cells'], strict=False):
                        for cell, snap_cell in zip(row, snap_row, strict=False):
                            for cid in cell['comments']:
                                cells.setdefault(cid, snap_cell['id'])
                continue
            for item in unit['items']:
                if item['t'] in ('cstart', 'cend', 'cref'):
                    key = {'cstart': 'start', 'cend': 'end', 'cref': 'ref'}[item['t']]
                    positions.setdefault(item['id'], {}).setdefault(key, (r, item['b']))
        sent = {}
        for comment in self.snapshot.get('review', {}).get('comments', []):
            pending = [comment]
            while pending:
                item = pending.pop(0)
                sent[comment_durable_id(item['id'])] = item
                sent[(item.get('author') or 'ByAuthorX', (item.get('text') or '').strip())] = item
                pending.extend(item.get('replies', []))

        def root_of(cid):
            seen = set()
            while threads[cid]['parent'] in threads and cid not in seen:
                seen.add(cid)
                cid = threads[cid]['parent']
            return cid

        def existing(info):
            return sent.get(info.get('durable_id')) or sent.get((info['author'], info['text']))

        ops: list[dict[str, Any]] = []
        targets: dict[str, str] = {}
        for cid, info in threads.items():
            if root_of(cid) != cid:
                continue
            match = existing(info)
            if match is not None:
                targets[cid] = match['id']
                report['existing'] += 1
                if info['done'] and match.get('status') != 'resolved':
                    report['marked_done'].append(match['id'])
                continue
            anchor = self.comment_anchor(cid, positions, cells)
            if anchor is None:
                continue
            identity = self.new_id('comment', 'comment', info.get('durable_id') or cid)
            op = {
                'type': 'add_comment',
                'id': identity,
                'target': anchor['target'],
                'text': info['text'],
                'author': info['author'] or 'Reviewer',
                'source': 'docx_review',
            }
            if 'start' in anchor:
                op.update(start=anchor['start'], end=anchor['end'])
            if info['date']:
                op['date'] = info['date']
            ops.append(op)
            targets[cid] = identity
            report['added'] += 1
            if info['done']:
                ops.append({'type': 'resolve_comment', 'target': identity})
                report['resolved'] += 1
        for cid, info in threads.items():
            root = root_of(cid)
            if root == cid or root not in targets:
                continue
            if existing(info) is not None:
                continue
            op = {
                'type': 'reply_comment',
                'id': self.new_id('comment', 'reply', info.get('durable_id') or cid),
                'target': targets[root],
                'text': info['text'],
                'author': info['author'] or 'Reviewer',
            }
            if info['date']:
                op['date'] = info['date']
            ops.append(op)
            report['replies'] += 1
        return ops, report

    def comment_anchor(self, cid, positions, cells):
        if cid in cells:
            return {'target': cells[cid]}
        start, end = self.comment_position(cid, positions)
        if start is None:
            self.losses.append(
                _loss(
                    'docx_review_comment_unanchored',
                    'Comment without a location in the text; attached to the first paragraph.',
                    None,
                    comment=cid,
                )
            )
            first = next((u['id'] for u in self.units if u['kind'] == 'text' and u['id'] in self.ours), None)
            return {'target': first, 'start': 0, 'end': 0} if first else None
        r, offset = start
        r2, offset2 = end
        if r != r2:
            offset2 = len(self.returned[r]['base'])
            self.losses.append(
                self.located(
                    _loss(
                        'docx_review_comment_range',
                        'Comment range spans paragraphs; anchored to the end of its first paragraph.',
                        None,
                        comment=cid,
                    ),
                    r,
                )
            )
        if r not in self.match or self.units[self.match[r]]['kind'] != 'text':
            for previous in range(r - 1, -1, -1):
                i = self.match.get(previous)
                if i is not None and self.units[i]['kind'] == 'text' and self.units[i]['id'] in self.ours:
                    self.losses.append(
                        _loss(
                            'docx_review_comment_on_new_text',
                            'Comment on a paragraph added in Word; anchored to the end of the preceding paragraph.',
                            self.units[i]['id'],
                            comment=cid,
                        )
                    )
                    length = len(self.ours[self.units[i]['id']].get('text', ''))
                    return {'target': self.units[i]['id'], 'start': length, 'end': length}
            return None
        i = self.match[r]
        block_id = self.units[i]['id']
        if block_id not in self.ours:
            self.losses.append(
                _loss(
                    'docx_review_comment_unanchored',
                    'Comment on a paragraph you removed; it was not added.',
                    block_id,
                    comment=cid,
                )
            )
            return None
        to_snapshot = self.to_snapshot(r, i)
        s, e = to_snapshot.start(offset), to_snapshot.end(max(offset, offset2))
        e = max(s, e)
        offsets = self.to_ours(i)
        if not offsets.clean(s, e):
            self.losses.append(
                _loss(
                    'docx_review_comment_remapped',
                    'Commented text changed since the review copy was sent; check the comment range.',
                    block_id,
                    comment=cid,
                )
            )
        text = self.ours[block_id].get('text', '')
        os_, oe = offsets.start(s), max(offsets.start(s), offsets.end(e))
        os_, oe = _snap(text, min(os_, len(text)), min(oe, len(text)))
        return {'target': block_id, 'start': os_, 'end': oe}

    # proposals

    def proposals(self):
        """One proposal per change; changes on a paragraph that is also moved, joined, deleted or
        reformatted share that change's proposal, because accepting a separate text edit first would
        leave the structural edit stale."""
        parent = list(range(len(self.changes)))

        def find(index):
            while parent[index] != index:
                parent[index] = parent[parent[index]]
                index = parent[index]
            return index

        by_block: dict[str, list[int]] = {}
        for index, change in enumerate(self.changes):
            for block in change['blocks']:
                by_block.setdefault(block, []).append(index)
        for indices in by_block.values():
            structural = [i for i in indices if self.changes[i]['anchorless']]
            for index in indices if structural else []:
                parent[find(index)] = find(structural[0])
        groups: dict[int, list[dict[str, Any]]] = {}
        for index, change in enumerate(self.changes):
            groups.setdefault(find(index), []).append(change)
        return [self.proposal(group) for group in groups.values()]

    def proposal(self, group):
        if len(group) == 1:
            operations = list(group[0]['ops'])
        else:
            # Within one proposal edits apply in order without remapping: text edits run from the end of
            # each paragraph backwards, then structural edits, so every recorded offset stays exact.
            textual = [c for c in group if not c['anchorless']]
            structural = [c for c in group if c['anchorless']]
            textual.sort(key=lambda c: -c['pos'])
            operations = [op for c in textual for op in c['ops']] + [
                op for c in structural for op in c['ops']
            ]
        authors = sorted({a for c in group for a in c['authors']})
        dates = sorted({d for c in group for d in c['dates']})
        tracked = all(c['tracked'] for c in group)
        lead = group[0]
        summary = (
            lead['summary']
            if len(group) == 1
            else f'{len(group)} related changes: ' + '; '.join(c['summary'] for c in group)
        )
        attribution = ', '.join(authors) + (f', {dates[0][:10]}' if dates else '')
        explanation = f'{summary} — {attribution} ({"tracked change" if tracked else "edit made without Track Changes"})'
        conflicts = [c['conflict'] for c in group if c['conflict']]
        if conflicts:
            explanation = 'Conflict: ' + explanation + '. ' + conflicts[0]['reason']
        values: dict[str, Any] = {
            'kind': 'docx_review',
            'operations': operations,
            'explanation': explanation,
            'author': authors[0] if authors else self.untracked,
            'review_source': {
                'review_id': self.record['id'],
                'format': 'docx',
                'file': self.source['file'],
                'sha256': self.source['sha256'],
                'changes': [
                    {
                        'kind': c['kind'],
                        'tracked': c['tracked'],
                        'authors': c['authors'],
                        'dates': c['dates'],
                        'revision_ids': c['revision_ids'],
                        'summary': c['summary'],
                        **({'base': c['base']} if c['base'] else {}),
                        **({'paragraph': c['el']} if c['el'] is not None else {}),
                    }
                    for c in group
                ],
            },
        }
        if conflicts:
            values['conflict'] = conflicts[0]
        return {'type': 'propose', 'values': values}


def _read_reviewed(path):
    from .exchange import MAX_INPUT

    source = Path(path)
    if source.suffix.lower() != '.docx':
        raise ByaxError('invalid_input', 'The reviewed copy must be a .docx file', field='reviewed_path')
    if not source.is_file():
        raise ByaxError('not_found', 'Reviewed copy does not exist', path=str(source))
    if source.stat().st_size > MAX_INPUT:
        raise ByaxError('resource_limit', 'Reviewed copy exceeds 32 MiB limit', path=str(source))
    return source.read_bytes()


def plan_review_merge(doc, reviewed_path, *, review_id=None) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    """Canonical operations that merge a returned review copy into ``doc``, plus the merge report.

    The operations register one pending ``docx_review`` proposal per editor change (related changes on a
    paragraph that is also moved, joined or reformatted share one proposal), add the editor's comments
    and replies, and record the return in the review package. Applying them changes no text.
    """
    from . import engine

    raw = _read_reviewed(reviewed_path)
    sha = hashlib.sha256(raw).hexdigest()
    parsed = read_returned(raw)
    properties = parsed['properties']
    found = properties.get(PROPERTY_NAMES['review_id'])
    packages = doc.get('metadata', {}).get('review_packages', {})
    if review_id and found and review_id != found:
        raise ByaxError(
            'conflict',
            'The returned file belongs to a different review package.',
            target=found,
            expected=review_id,
        )
    chosen = review_id or found
    if not chosen:
        raise ByaxError(
            'missing_target',
            'The returned file has no ByAuthorX review identifier; pass the review package ID explicitly.',
            candidates=sorted(packages),
        )
    if properties.get(PROPERTY_NAMES['document_id']) not in (None, doc['id']):
        raise ByaxError(
            'conflict', 'The returned file is a review copy of a different document.', target=chosen
        )
    record = packages.get(chosen)
    if record is None:
        raise ByaxError(
            'missing_target',
            'This document has no record of that review package.',
            target=chosen,
            candidates=sorted(packages),
        )
    if doc['history'].get('peer_bases', {}).get(record['peer']) != record['revision']:
        raise ByaxError(
            'missing_target',
            'The review snapshot is no longer pinned in this document, so a safe three-way merge is not possible.',
            target=chosen,
        )
    source = {'file': Path(reviewed_path).name, 'sha256': sha}
    info = {
        'review_id': chosen,
        'matched_by': 'custom_property' if found else 'argument',
        'snapshot_revision': record['revision'],
        **source,
        'already_merged': False,
    }
    if any(entry.get('sha256') == sha for entry in record.get('returns', [])):
        info['already_merged'] = True
        return [], {
            'review_merge': info,
            'proposals': 0,
            'conflicts': [],
            'losses': [],
            'comments': {},
            'summary': ['This returned copy was already merged.'],
        }
    snapshot = engine.revision_document(doc, record['revision'])
    _, _, units = _review_candidate(snapshot)
    if snapshot_hash(units) != record['snapshot_hash']:
        raise ByaxError(
            'conflict', 'The pinned review snapshot does not reproduce the copy that was sent.', target=chosen
        )
    if properties.get(PROPERTY_NAMES['snapshot_hash']) not in (None, record['snapshot_hash']):
        raise ByaxError('conflict', 'The returned file was made from a different snapshot.', target=chosen)
    planner = _Planner(doc, snapshot, units, parsed, record, source)
    proposals, comment_ops, comment_report = planner.run()
    record = copy.deepcopy(record)
    record['status'] = 'merged'
    record['returns'] = [
        *record.get('returns', []),
        {
            **source,
            'merged_at': _now(),
            'base_revision': doc['revision'],
            'proposals': len(proposals),
            'comments': comment_report['added'],
            'replies': comment_report['replies'],
        },
    ]
    packages = {**copy.deepcopy(packages), chosen: record}
    operations = proposals + comment_ops + [{'type': 'set_metadata', 'values': {'review_packages': packages}}]
    summary = [
        f'{len(proposals)} proposal(s) from {info["file"]}',
        f'{comment_report["added"]} new comment(s), {comment_report["replies"]} repl(ies)',
    ]
    if planner.conflicts:
        summary.append(
            f'{len(planner.conflicts)} change(s) could not be placed because you changed or removed that text'
        )
    if planner.losses:
        summary.append(f'{len(planner.losses)} located loss(es): constructs that were not merged')
    report = {
        'review_merge': info,
        'proposals': len(proposals),
        'proposal_summaries': [op['values']['explanation'] for op in proposals],
        'conflicts': planner.conflicts,
        'losses': planner.losses,
        'comments': comment_report,
        'summary': summary,
    }
    return operations, report


def merge_reviewed_copy(
    doc, reviewed_path, *, review_id=None, dry_run=False
) -> tuple[dict[str, Any], dict[str, Any]]:
    """Plan and apply the merge; returns ``(document, result)``. Proposals stay pending for review."""
    from . import engine

    operations, report = plan_review_merge(doc, reviewed_path, review_id=review_id)
    if not operations:
        return doc, {**report, 'changed_ids': []}
    updated, result = engine.apply(
        doc, {'expected_revision': doc['revision'], 'operations': operations}, dry_run=dry_run
    )
    proposals = [
        p['id']
        for p in updated['review']['proposals']
        if p.get('kind') == 'docx_review'
        and p.get('review_source', {}).get('sha256') == report['review_merge']['sha256']
    ]
    report['proposal_ids'] = proposals
    return (doc if dry_run else updated), {**result, **report}


# -- Commands (CLI, MCP and REST through cli.execute) ------------------------------------------


def _check_revision(doc, expected):
    if expected and doc['revision'] != expected:
        raise ByaxError(
            'stale_revision', 'Document revision changed', expected=expected, actual=doc['revision']
        )


def review_package_command(a: dict[str, Any], owner: str | None) -> dict[str, Any]:
    """``review-package``: create a Word review copy (default) or list the recorded review packages."""
    from .cli import envelope
    from .storage import DocumentLock, load_document, save_document

    action = a.get('action', 'create')
    if action == 'list':
        doc = load_document(a['path'])
        return envelope(review_packages(doc), doc)
    if action != 'create':
        raise ByaxError('invalid_input', f'Unknown review-package action: {action}', field='action')
    if not a.get('out'):
        raise ByaxError('invalid_input', 'review-package requires --out', field='out')
    lock = DocumentLock(a['path'])
    with lock:
        doc = load_document(a['path'])
        _check_revision(doc, a.get('expect_revision'))
        overwrite = a.get('overwrite', False)
        updated, result, data = create_review_package(
            doc,
            a['out'],
            restrict=a.get('restrict', 'tracked_changes'),
            reviewer=a.get('reviewer', ''),
            overwrite=overwrite,
            dry_run=a.get('dry_run', False),
        )
        if data is None:
            return envelope(result, doc)
        # The snapshot is recorded before the copy exists, so a published copy can always be merged.
        save_document(updated, a['path'], owner=lock.owner)
        publish_review_copy(data, a['out'], owner=owner, overwrite=overwrite)
        return envelope(result, updated)


def merge_review_command(a: dict[str, Any], owner: str | None) -> dict[str, Any]:
    """``merge-review``: merge a returned Word review copy as pending proposals and comments."""
    from .cli import _mutate

    def merge(doc):
        _check_revision(doc, a.get('expect_revision'))
        updated, result = merge_reviewed_copy(
            doc, a['reviewed_path'], review_id=a.get('review_id'), dry_run=a.get('dry_run', False)
        )
        return updated, result, updated is not doc

    return _mutate(a['path'], merge)
