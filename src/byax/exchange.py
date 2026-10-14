"""Bounded, offline import and staged export adapters.

Markdown dialect: CommonMark plus markdown-it's table extension. Raw HTML is
retained in provenance but is never executed. Original imports are embedded as
opaque assets, and conversion cannot overwrite the imported source path.
"""

from __future__ import annotations

import base64
import copy
import hashlib
import html
import io
import itertools
import json
import os
import re
import subprocess
import tempfile
import uuid
import zipfile
import zlib
from html.parser import HTMLParser
from pathlib import Path, PurePosixPath

from .errors import ByaxError
from .model import link_external_image, new_document

MAX_INPUT = 32 * 1024 * 1024
MAX_EXPANDED = 128 * 1024 * 1024
MAX_REPEATED_SPACES = 1_000_000
MAX_TABLE_CELLS = 100_000
MAX_DEPTH = 256


def _id(prefix='block'):
    return f'{prefix}_{uuid.uuid4().hex}'


def _block(text='', kind='paragraph', **attrs):
    return {'id': _id(), 'type': kind, 'text': text, 'style': 'body', 'marks': [], 'attrs': attrs}


def _loss(losses, code, message, target=None):
    losses.append({'code': code, 'severity': 'warning', 'message': message, 'target': target})


def _asset(doc, data, mime, name):
    key = hashlib.sha256(data).hexdigest()
    doc.setdefault('assets', {})[key] = {
        'data': base64.b64encode(data).decode('ascii'),
        'mime': mime,
        'name': name,
    }
    return key


def _fingerprint(block):
    return hashlib.sha256(json.dumps(block, sort_keys=True, ensure_ascii=False).encode()).hexdigest()


def _list_kind(attrs):
    value = attrs.get('list')
    return value.get('kind', 'bullet') if isinstance(value, dict) else value


def _list_level(attrs):
    value = attrs.get('list')
    return int(value.get('level', 1)) if isinstance(value, dict) else int(attrs.get('level', 1))


def _safe_url(value):
    return value.startswith(('http://', 'https://', 'mailto:', '#')) or ':' not in value


def _decode(raw):
    if raw.startswith((b'\xff\xfe', b'\xfe\xff')):
        return raw.decode('utf-16'), 'utf-16'
    try:
        return raw.decode('utf-8-sig'), 'utf-8-sig' if raw.startswith(b'\xef\xbb\xbf') else 'utf-8'
    except UnicodeDecodeError:
        return raw.decode('cp1252', errors='replace'), 'cp1252'


def _raised_by_package_reader(exc):
    """True when the innermost frame is python-docx or zipfile rather than ByAuthorX code."""
    import traceback

    import docx

    frames = traceback.extract_tb(exc.__traceback__)
    origin = os.path.normcase(frames[-1].filename) if frames else ''
    zip_source = os.path.normcase(zipfile.__file__)
    if zip_source.endswith('__init__.py'):
        zip_source = os.path.dirname(zip_source) + os.sep
    return origin.startswith((os.path.normcase(os.path.dirname(docx.__file__)) + os.sep, zip_source))


def _safe_zip(raw):
    try:
        archive = zipfile.ZipFile(io.BytesIO(raw))
        entries = archive.infolist()
        if len(entries) > 10000 or sum(item.file_size for item in entries) > MAX_EXPANDED:
            raise ByaxError('resource_limit', 'Archive expanded size or member count exceeds import limit')
        names = set()
        for item in entries:
            path = PurePosixPath(item.filename)
            if path.is_absolute() or '..' in path.parts or '\\' in item.filename or item.filename in names:
                raise ByaxError(
                    'unsafe_archive', 'Archive contains unsafe or duplicate member', member=item.filename
                )
            names.add(item.filename)
            if item.flag_bits & 1:
                raise ByaxError('encrypted_input', 'Encrypted archive is not supported')
            if item.file_size > MAX_INPUT or (
                item.compress_size and item.file_size / item.compress_size > 500
            ):
                raise ByaxError(
                    'resource_limit', 'Archive member exceeds size or expansion limit', member=item.filename
                )
            if item.filename.endswith(('.xml', '.rels')):
                from defusedxml.ElementTree import fromstring

                try:
                    fromstring(archive.read(item))
                except Exception as exc:
                    raise ByaxError(
                        'unsafe_xml', 'Invalid XML or prohibited XML entity declaration', member=item.filename
                    ) from exc
        corrupt = archive.testzip()
        if corrupt:
            raise ByaxError('invalid_input', 'Corrupt ZIP member', member=corrupt)
        return archive
    except (zipfile.BadZipFile, NotImplementedError, EOFError, zlib.error) as exc:
        raise ByaxError('invalid_input', 'Invalid ZIP-based document') from exc


def _report(fmt, losses, outputs=None, direction='export'):
    return {
        'schema_version': 1,
        'execution': 'completed',
        'validation': 'not_run',
        'format': fmt,
        'direction': direction,
        'outputs': outputs or [],
        'losses': losses,
        'findings': losses,
        'coverage': {
            'content': 'supported text retained; see located losses',
            'structure': 'editable adapter subset',
            'visual': 'not verified against source application',
            'review': 'native only; exchange losses reported',
        },
        'fidelity': 'partial' if losses else 'supported_subset',
    }


def import_document(path):
    path = Path(path)
    fmt = path.suffix.lower().lstrip('.')
    if fmt == 'byax':
        from .storage import load_document

        return load_document(path), _report(fmt, [], direction='import')
    if not path.is_file():
        raise ByaxError('not_found', 'Import file does not exist', path=str(path))
    if path.stat().st_size > MAX_INPUT:
        raise ByaxError('resource_limit', 'Import file exceeds 32 MiB limit', path=str(path))
    raw, losses = path.read_bytes(), []
    doc = new_document(path.stem)
    doc['blocks'] = []
    source = None
    if fmt in ('txt', 'md', 'markdown', 'html', 'htm', 'rtf'):
        source, encoding = _decode(raw)
        if fmt == 'txt':
            doc['blocks'] = [_block(line) for line in source.splitlines()]
        elif fmt in ('md', 'markdown'):
            _import_markdown(doc, source, losses)
        elif fmt in ('html', 'htm'):
            parser = _HTMLImport(doc, losses)
            parser.feed(source)
            parser.close()
            parser.flush()
        else:
            _import_rtf(doc, source, losses)

    elif fmt == 'docx':
        _safe_zip(raw).close()
        try:
            _import_docx(doc, raw, losses)
        except (KeyError, AttributeError) as exc:
            # python-docx reports missing or mistyped package parts this way; keep our own bugs visible.
            if not _raised_by_package_reader(exc):
                raise
            raise ByaxError(
                'invalid_input', f'DOCX package is missing or has an invalid part: {exc}'
            ) from exc
    elif fmt == 'odt':
        _import_odt(doc, raw, losses)
    elif fmt == 'doc':
        doc, adapter_report = _import_doc(path)
        losses.extend(adapter_report['losses'])
    else:
        raise ByaxError('unsupported_format', f'Unsupported import format: {fmt or "(none)"}')
    for block in doc['blocks']:
        if block.get('type') == 'table':
            rows = block.get('rows', [])
            width = max((len(row.get('cells', [])) for row in rows), default=0)
            if not rows or not width:
                block['type'], block['text'] = 'paragraph', ''
                block.pop('rows', None)
                _loss(losses, 'empty_table', 'Empty table normalized to an empty paragraph.', block['id'])
                continue
            for row in rows:
                if len(row['cells']) < width:
                    row['cells'].extend(
                        {'id': _id('cell'), 'text': '', 'attrs': {}} for _ in range(width - len(row['cells']))
                    )
                    _loss(
                        losses,
                        'ragged_table',
                        'Short table row padded with empty cells to preserve an editable rectangular grid.',
                        row['id'],
                    )
    if not doc['blocks']:
        doc['blocks'] = [_block()]
    original_key = _asset(doc, raw, 'application/octet-stream', path.name)
    provenance = {
        'filename': path.name,
        'original_path': str(path.resolve()),
        'format': fmt,
        'original_asset': original_key,
        'sha256': hashlib.sha256(raw).hexdigest(),
        'losses': losses,
    }
    if source is not None:
        provenance.update(
            {
                'source': source,
                'encoding': encoding,
                'line_ending': '\r\n' if '\r\n' in source else '\n',
                'fingerprints': {b['id']: _fingerprint(b) for b in doc['blocks']},
                'block_order': [b['id'] for b in doc['blocks']],
            }
        )
    doc['metadata']['import'] = provenance
    _snap_marks(doc)
    from .model import validate_document

    validate_document(doc)
    return doc, _report(fmt, losses, direction='import')


def _snap_marks(doc):
    """Source runs may split grapheme clusters; widen marks to whole clusters."""
    import regex

    for block in doc['blocks']:
        for item in [block, *(cell for row in block.get('rows', []) for cell in row.get('cells', []))]:
            marks, text = item.get('marks'), item.get('text', '')
            if not marks:
                continue
            boundaries = sorted({0, len(text)} | {m.end() for m in regex.finditer(r'\X', text)})
            boundary_set = set(boundaries)
            for mark in marks:
                if mark.get('start') not in boundary_set:
                    mark['start'] = max(x for x in boundaries if x < mark['start'])
                if mark.get('end') not in boundary_set:
                    mark['end'] = min(x for x in boundaries if x > mark['end'])


def _inline(tokens):
    text, marks, opened, images = '', [], [], []
    for token in tokens or []:
        kind = token.type
        if kind in ('text', 'code_inline', 'softbreak', 'hardbreak', 'html_inline'):
            value = '\n' if kind in ('softbreak', 'hardbreak') else token.content
            start = len(text)
            text += value
            if kind == 'code_inline':
                marks.append({'start': start, 'end': len(text), 'type': 'code', 'value': True})
        elif kind == 'image':
            images.append({'src': token.attrGet('src'), 'alt': token.content})
            text += token.content
        elif kind.endswith('_open'):
            name = {
                'strong_open': 'bold',
                'em_open': 'italic',
                's_open': 'strikethrough',
                'link_open': 'link',
            }.get(kind)
            if name:
                opened.append((kind, name, len(text), token.attrGet('href') if name == 'link' else True))
        elif kind.endswith('_close') and opened:
            opening = kind.replace('_close', '_open')
            for index in range(len(opened) - 1, -1, -1):
                if opened[index][0] == opening:
                    _, name, start, value = opened.pop(index)
                    if name != 'link' or _safe_url(value or ''):
                        marks.append({'start': start, 'end': len(text), 'type': name, 'value': value})
                    break
    return text, marks, images


def _import_markdown(doc, source, losses):
    from markdown_it import MarkdownIt

    parser = MarkdownIt('commonmark', {'html': False}).enable('table')
    tokens = parser.parse(source)
    lists, quote_depth, i = [], 0, 0
    while i < len(tokens):
        token = tokens[i]
        kind = token.type
        if kind in ('bullet_list_open', 'ordered_list_open'):
            lists.append('ordered' if kind == 'ordered_list_open' else 'bullet')
        elif kind in ('bullet_list_close', 'ordered_list_close'):
            lists.pop()
        elif kind == 'blockquote_open':
            quote_depth += 1
        elif kind == 'blockquote_close':
            quote_depth -= 1
        elif kind in ('heading_open', 'paragraph_open'):
            inline = tokens[i + 1]
            text, marks, images = _inline(inline.children)
            block = _block(text, 'heading' if kind == 'heading_open' else 'paragraph')
            block['marks'] = marks
            if kind == 'heading_open':
                block['attrs']['level'] = int(token.tag[1:])
            if lists:
                block['attrs'].update({'list': lists[-1], 'level': len(lists)})
            if quote_depth:
                block['style'] = 'quotation'
                block['attrs']['quote_depth'] = quote_depth
            block['attrs']['source_lines'] = token.map
            doc['blocks'].append(block)
            for image in images:
                _loss(
                    losses,
                    'external_image',
                    'Markdown image is not fetched; source URL retained in provenance.',
                    block['id'],
                )
                block['attrs'].setdefault('external_images', []).append(image)
            i += 1
        elif kind in ('fence', 'code_block'):
            block = _block(token.content.rstrip('\n'), code_language=token.info, source_lines=token.map)
            block['style'] = 'code'
            block['marks'] = [{'start': 0, 'end': len(block['text']), 'type': 'code', 'value': True}]
            doc['blocks'].append(block)
        elif kind == 'hr':
            block = _block('―', thematic_break=True, source_lines=token.map)
            doc['blocks'].append(block)
        elif kind == 'table_open':
            block = _block(kind='table', header_rows=1, source_lines=token.map)
            rows, cells, cell = [], None, None
            i += 1
            while i < len(tokens) and tokens[i].type != 'table_close':
                current = tokens[i]
                if current.type == 'tr_open':
                    cells = []
                elif current.type in ('th_open', 'td_open'):
                    cell = {'id': _id('cell'), 'text': '', 'marks': [], 'attrs': {}}
                elif current.type == 'inline' and cell is not None:
                    cell['text'], cell['marks'], _ = _inline(current.children)
                elif current.type in ('th_close', 'td_close'):
                    cells.append(cell)
                    cell = None
                elif current.type == 'tr_close':
                    rows.append({'id': _id('row'), 'cells': cells})
                i += 1
            block['rows'] = rows
            doc['blocks'].append(block)
        i += 1


MAX_SPAN = 1000


def _heading_level(level, losses, target):
    """The model supports heading levels 1-6; deeper outline levels are reported."""
    if 1 <= level <= 6:
        return level
    _loss(losses, 'heading_level', f'Heading level {level} is outside 1-6 and was clamped.', target)
    return min(6, max(1, level))


def _span(value):
    """Row/column spans are small positive ASCII integers ("²".isdigit() is true)."""
    value = str(value).strip()
    if not value.isascii() or not value.isdecimal():
        return 1
    return max(1, min(int(value), MAX_SPAN))


class _HTMLImport(HTMLParser):
    def __init__(self, doc, losses):
        super().__init__(convert_charrefs=True)
        self.doc, self.losses = doc, losses
        self.current = None
        self.stack, self.lists, self.tables = [], [], []
        self.skip = 0

    def flush(self):
        if self.current is not None:
            if self.current['text'].strip() or self.current['type'] == 'heading':
                self.doc['blocks'].append(self.current)
            self.current = None
        self.stack = []

    def _end_cell(self):
        current = self.tables[-1]
        if current['cell'] is not None:
            current['cells'].append(current['cell'])
            current['cell'] = None

    def _end_row(self):
        # HTML lets </td>, </th> and </tr> be omitted; close them implicitly.
        self._end_cell()
        current = self.tables[-1]
        if current['row_open']:
            current['rows'].append({'id': _id('row'), 'cells': current['cells']})
            current['cells'], current['row_open'] = [], False

    def handle_starttag(self, tag, attrs):
        # Bare attributes such as <a href> have no value.
        attrs = {key: value or '' for key, value in attrs}
        if tag == 'title':
            # Document titles are metadata, never body text; </head> is optional.
            self.skip += 1
            return
        if tag in ('script', 'style', 'iframe', 'object', 'embed', 'svg', 'math'):
            if tag != 'embed':
                self.skip += 1
            _loss(
                self.losses,
                'html_inactive_content',
                f'<{tag}> was not executed or rendered; original HTML remains attached.',
                f'line:{self.getpos()[0]}',
            )
            return
        if self.skip:
            return
        if tag in ('p', 'div', 'h1', 'h2', 'h3', 'h4', 'h5', 'h6', 'blockquote', 'pre', 'li'):
            self.flush()
            self.current = _block(kind='heading' if re.fullmatch(r'h[1-6]', tag) else 'paragraph')
            if tag.startswith('h') and len(tag) == 2:
                self.current['attrs']['level'] = int(tag[1])
            if tag == 'blockquote':
                self.current['style'] = 'quotation'
            if tag == 'pre':
                self.current['style'] = 'code'
            if tag == 'li':
                self.current['attrs'].update(
                    {'list': self.lists[-1] if self.lists else 'bullet', 'level': len(self.lists)}
                )
        elif tag in ('ul', 'ol'):
            self.lists.append('ordered' if tag == 'ol' else 'bullet')
        elif tag == 'br':
            self.handle_data('\n')
        elif tag == 'table':
            self.flush()
            self.tables.append(
                {
                    'block': _block(kind='table', header_rows=0),
                    'rows': [],
                    'cells': [],
                    'cell': None,
                    'row_open': False,
                }
            )
        elif tag == 'tr' and self.tables:
            self._end_row()
            self.tables[-1]['row_open'] = True
        elif tag in ('th', 'td') and self.tables:
            self._end_cell()
            self.tables[-1]['row_open'] = True
            self.tables[-1]['cell'] = {
                'id': _id('cell'),
                'text': '',
                'attrs': {key: _span(attrs[key]) for key in ('rowspan', 'colspan') if attrs.get(key)},
            }
            if tag == 'th':
                self.tables[-1]['block']['attrs']['header_rows'] = 1
        elif tag == 'img':
            self.flush()
            block = _block(kind='image', alt=attrs.get('alt', ''), source=attrs.get('src', ''))
            self.doc['blocks'].append(block)
            data_uri = re.fullmatch(
                r'data:(image/(?:png|jpeg|gif|webp));base64,([A-Za-z0-9+/=\s]+)', attrs.get('src', '')
            )
            if data_uri and len(data_uri[2]) <= MAX_INPUT * 2:
                try:
                    raw = base64.b64decode(data_uri[2], validate=False)
                    block['attrs']['asset'] = _asset(self.doc, raw, data_uri[1], 'embedded-image')
                except ValueError:
                    _loss(self.losses, 'invalid_image', 'Malformed embedded image data.', block['id'])
            else:
                if attrs.get('src'):
                    link_external_image(self.doc, block, attrs['src'])
                _loss(
                    self.losses,
                    'external_image',
                    'HTML images are not fetched automatically; explicit embedding is required.',
                    block['id'],
                )
        else:
            mark = {
                'b': 'bold',
                'strong': 'bold',
                'i': 'italic',
                'em': 'italic',
                'u': 'underline',
                's': 'strikethrough',
                'del': 'strikethrough',
                'code': 'code',
                'sup': 'superscript',
                'sub': 'subscript',
                'a': 'link',
            }.get(tag)
            if mark:
                if self.current is None:
                    self.current = _block()
                self.stack.append(
                    (tag, mark, len(self.current['text']), attrs.get('href', '') if mark == 'link' else True)
                )
        if attrs.get('style') or attrs.get('class'):
            _loss(
                self.losses,
                'html_css',
                'CSS styling is not imported by the semantic HTML adapter.',
                f'line:{self.getpos()[0]}',
            )

    def handle_endtag(self, tag):
        if tag in ('script', 'style', 'iframe', 'object', 'svg', 'math', 'title') and self.skip:
            self.skip -= 1
            return
        if self.skip:
            return
        if tag in ('td', 'th') and self.tables:
            self._end_cell()
        elif tag == 'tr' and self.tables:
            self._end_row()
        elif tag == 'table' and self.tables:
            self._end_row()
            current = self.tables.pop()
            current['block']['rows'] = current['rows']
            self.doc['blocks'].append(current['block'])
        elif tag in ('ul', 'ol') and self.lists:
            self.lists.pop()
        elif tag in ('p', 'div', 'h1', 'h2', 'h3', 'h4', 'h5', 'h6', 'blockquote', 'pre', 'li'):
            self.flush()
        elif self.current:
            for index in range(len(self.stack) - 1, -1, -1):
                item_tag, mark, start, value = self.stack[index]
                if item_tag == tag:
                    self.stack.pop(index)
                    if mark != 'link' or _safe_url(value):
                        self.current['marks'].append(
                            {'start': start, 'end': len(self.current['text']), 'type': mark, 'value': value}
                        )
                    break

    def handle_data(self, data):
        if self.skip:
            return
        if self.tables and self.tables[-1]['cell'] is not None:
            self.tables[-1]['cell']['text'] += data
        else:
            if self.current is None:
                if not data.strip():
                    return
                self.current = _block()
            self.current['text'] += data


def _import_docx(doc, raw, losses):
    from docx import Document
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    word = Document(io.BytesIO(raw))
    for style in word.styles:
        if str(style.type) not in ('PARAGRAPH (1)', 'CHARACTER (2)'):
            continue
        values = {'family': 'text' if str(style.type) == 'CHARACTER (2)' else 'paragraph'}
        if style.base_style is not None:
            values['based_on'] = style.base_style.style_id
        for key, value in (
            ('bold', style.font.bold),
            ('italic', style.font.italic),
            ('underline', style.font.underline),
            ('strikethrough', style.font.strike),
            ('font_family', style.font.name),
            ('font_size', style.font.size.pt if style.font.size else None),
        ):
            if value is not None:
                values[key] = value
        if style.font.color.rgb:
            values['color'] = '#' + str(style.font.color.rgb)
        if str(style.type) == 'PARAGRAPH (1)':
            pf = style.paragraph_format
            for key, prop in (
                ('space_before', 'space_before'),
                ('space_after', 'space_after'),
                ('indent', 'left_indent'),
                ('first_line_indent', 'first_line_indent'),
            ):
                value = getattr(pf, prop)
                if value is not None:
                    values[key] = value.pt
        doc['styles'][style.style_id] = values
    if word.sections:
        section = word.sections[0]
        doc['layout'].update(
            {
                'width': section.page_width.pt,
                'height': section.page_height.pt,
                'margins': {
                    key: getattr(section, f'{key}_margin').pt for key in ('top', 'right', 'bottom', 'left')
                },
            }
        )
        doc['layout']['header'] = '\n'.join(p.text for p in section.header.paragraphs)
        doc['layout']['footer'] = '\n'.join(p.text for p in section.footer.paragraphs)
        if len(word.sections) > 1:
            _loss(
                losses,
                'docx_sections',
                'Only first section page geometry and repeated regions are imported.',
                'sections',
            )
    doc['metadata'].update(
        {
            'title': word.core_properties.title or doc['metadata']['title'],
            'author': word.core_properties.author or '',
        }
    )
    for child in word.element.body:
        if child.tag == qn('w:p'):
            paragraph = Paragraph(child, word)
            name = paragraph.style.name if paragraph.style else 'Normal'
            match = re.match(r'Heading (\d)', name)
            block = _block(paragraph.text, 'heading' if match else 'paragraph')
            if paragraph.style:
                block['style'] = paragraph.style.style_id
            if match:
                block['attrs']['level'] = _heading_level(int(match[1]), losses, block['id'])
            alignment = {0: 'left', 1: 'center', 2: 'right', 3: 'justify'}.get(paragraph.alignment)
            if alignment:
                block['attrs']['alignment'] = alignment
            for prop, target in (
                ('space_before', 'space_before'),
                ('space_after', 'space_after'),
                ('left_indent', 'indent'),
                ('first_line_indent', 'first_line_indent'),
            ):
                value = getattr(paragraph.paragraph_format, prop)
                if value is not None:
                    block['attrs'][target] = value.pt
            offset = 0
            content_runs = []
            if hasattr(paragraph, 'iter_inner_content'):
                for item in paragraph.iter_inner_content():
                    if hasattr(item, 'runs'):
                        content_runs.extend((run, item.url) for run in item.runs)
                    else:
                        content_runs.append((item, None))
            else:
                content_runs = [(run, None) for run in paragraph.runs]
            for run, hyperlink in content_runs:
                if hyperlink and _safe_url(hyperlink):
                    block['marks'].append(
                        {'start': offset, 'end': offset + len(run.text), 'type': 'link', 'value': hyperlink}
                    )
                for key, value in (
                    ('bold', run.bold),
                    ('italic', run.italic),
                    ('underline', run.underline),
                    ('strikethrough', run.font.strike),
                    ('superscript', run.font.superscript),
                    ('subscript', run.font.subscript),
                    ('font_family', run.font.name),
                    ('font_size', run.font.size.pt if run.font.size else None),
                    ('color', '#' + str(run.font.color.rgb) if run.font.color.rgb else None),
                ):
                    if value is not None:
                        block['marks'].append(
                            {'start': offset, 'end': offset + len(run.text), 'type': key, 'value': value}
                        )
                offset += len(run.text)
            if child.xpath('.//w:numPr'):
                block['attrs']['list'] = 'bullet'
                _loss(
                    losses,
                    'docx_list_numbering',
                    'List membership retained; original numbering definitions/restarts are not resolved.',
                    block['id'],
                )
            doc['blocks'].append(block)
            for drawing in child.xpath('.//a:blip'):
                relationship = drawing.get(qn('r:embed'))
                if relationship in word.part.related_parts:
                    part = word.part.related_parts[relationship]
                    key = _asset(doc, part.blob, part.content_type, str(part.partname).split('/')[-1])
                    image = _block(kind='image', asset=key, alt='')
                    doc['blocks'].append(image)
                    _loss(
                        losses,
                        'docx_image_anchor',
                        'Image bytes imported; original inline/anchor position, wrapping and alternative text require review.',
                        image['id'],
                    )
            for unsupported in (
                'ins',
                'del',
                'commentRangeStart',
                'footnoteReference',
                'fldChar',
                'bookmarkStart',
                'pict',
                'oMath',
            ):
                if child.xpath(f'.//w:{unsupported}' if unsupported != 'oMath' else './/m:oMath'):
                    _loss(
                        losses,
                        'docx_' + unsupported,
                        f'Word {unsupported} structure is retained only in original attachment; displayed paragraph text may omit this content.',
                        block['id'],
                    )
        elif child.tag == qn('w:tbl'):
            table = Table(child, word)
            block = _block(
                kind='table',
                header_rows=1 if table.rows and table.rows[0]._tr.xpath('./w:trPr/w:tblHeader') else 0,
            )
            block['rows'] = [
                {
                    'id': _id('row'),
                    'cells': [{'id': _id('cell'), 'text': cell.text, 'attrs': {}} for cell in row.cells],
                }
                for row in table.rows
            ]
            doc['blocks'].append(block)
            _loss(
                losses,
                'docx_table_format',
                'Table text and grid retained; cell marks, merged spans, widths and border details require review.',
                block['id'],
            )
        elif child.tag != qn('w:sectPr'):
            _loss(
                losses,
                'docx_unknown',
                f'Unsupported Word body element {child.tag}; original remains attached.',
                'body',
            )
    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        for name in archive.namelist():
            if name.startswith(('word/comments', 'word/footnotes', 'word/endnotes', 'word/vbaProject')):
                _loss(
                    losses,
                    'docx_attachment',
                    f'{name} is preserved in original attachment, not imported as editable content.',
                    name,
                )


def _import_odt(doc, raw, losses):
    from defusedxml.ElementTree import fromstring

    ns = {
        'text': 'urn:oasis:names:tc:opendocument:xmlns:text:1.0',
        'table': 'urn:oasis:names:tc:opendocument:xmlns:table:1.0',
        'office': 'urn:oasis:names:tc:opendocument:xmlns:office:1.0',
        'draw': 'urn:oasis:names:tc:opendocument:xmlns:drawing:1.0',
        'xlink': 'http://www.w3.org/1999/xlink',
        'style': 'urn:oasis:names:tc:opendocument:xmlns:style:1.0',
        'fo': 'urn:oasis:names:tc:opendocument:xmlns:xsl-fo-compatible:1.0',
        'svg': 'urn:oasis:names:tc:opendocument:xmlns:svg-compatible:1.0',
    }

    def attr(element, family, key, default=None):
        return element.get('{' + ns[family] + '}' + key, default)

    def integer(element, family, key, default):
        value = attr(element, family, key, default)
        if not re.fullmatch(r'-?[0-9]{1,9}', str(value).strip()):
            raise ByaxError('invalid_input', f'ODT {key} must be an integer', field=key)
        return int(value)

    def points(value):
        match = re.fullmatch(r'(-?[0-9]+(?:\.[0-9]+)?|-?\.[0-9]+)(pt|in|cm|mm|px)?', str(value))
        if not match:
            return None
        return (
            float(match[1])
            * {None: 1, 'pt': 1, 'in': 72, 'cm': 72 / 2.54, 'mm': 72 / 25.4, 'px': 0.75}[match[2]]
        )

    with _safe_zip(raw) as archive:
        if 'META-INF/manifest.xml' in archive.namelist() and b'encryption-data' in archive.read(
            'META-INF/manifest.xml'
        ):
            raise ByaxError('encrypted_input', 'Encrypted ODT input is not supported')
        if 'content.xml' not in archive.namelist():
            raise ByaxError('invalid_input', 'ODT lacks content.xml')
        roots = [fromstring(archive.read('content.xml'))]
        if 'styles.xml' in archive.namelist():
            roots.append(fromstring(archive.read('styles.xml')))
        fonts, list_styles = {}, {}
        for root in roots:
            for listing in root.findall('.//text:list-style', ns):
                list_styles[attr(listing, 'style', 'name')] = (
                    'ordered' if listing.find('text:list-level-style-number', ns) is not None else 'bullet'
                )
        for root in roots:
            for face in root.findall('.//style:font-face', ns):
                fonts[attr(face, 'style', 'name')] = attr(face, 'svg', 'font-family', '').strip("'")
            for style in root.findall('.//style:style', ns):
                name = attr(style, 'style', 'name')
                values = {'family': attr(style, 'style', 'family', 'paragraph')}
                parent = attr(style, 'style', 'parent-style-name')
                if parent:
                    values['based_on'] = parent
                properties = style.find('style:text-properties', ns)
                if properties is not None:
                    font = attr(properties, 'fo', 'font-family') or fonts.get(
                        attr(properties, 'style', 'font-name')
                    )
                    if font:
                        values['font_family'] = font.strip("'")
                    size = points(attr(properties, 'fo', 'font-size', ''))
                    if size:
                        values['font_size'] = size
                    for key, prop, setting in (
                        ('bold', 'font-weight', 'bold'),
                        ('italic', 'font-style', 'italic'),
                    ):
                        value = attr(properties, 'fo', prop)
                        if value:
                            values[key] = value == setting
                    for key in ('color', 'background-color'):
                        value = attr(properties, 'fo', key)
                        if value and value != 'transparent':
                            values['highlight' if key == 'background-color' else key] = value
                    if attr(properties, 'style', 'text-underline-style', 'none') != 'none':
                        values['underline'] = True
                    if attr(properties, 'style', 'text-line-through-style', 'none') != 'none':
                        values['strikethrough'] = True
                paragraph = style.find('style:paragraph-properties', ns)
                if paragraph is not None:
                    alignment = attr(paragraph, 'fo', 'text-align')
                    if alignment:
                        values['alignment'] = {'start': 'left', 'end': 'right'}.get(alignment, alignment)
                    for key, source in (
                        ('space_before', 'margin-top'),
                        ('space_after', 'margin-bottom'),
                        ('indent', 'margin-left'),
                        ('first_line_indent', 'text-indent'),
                    ):
                        value = points(attr(paragraph, 'fo', source, ''))
                        if value is not None:
                            values[key] = value
                    line_height = attr(paragraph, 'fo', 'line-height', '')
                    if re.fullmatch(r'[0-9]+(?:\.[0-9]+)?%', line_height):
                        values['line_spacing'] = float(line_height[:-1]) / 100
                if name:
                    doc['styles'][name] = values
        for root in roots:
            page = root.find('.//style:page-layout-properties', ns)
            if page is not None:
                for target, source in (('width', 'page-width'), ('height', 'page-height')):
                    value = points(attr(page, 'fo', source, ''))
                    if value:
                        doc['layout'][target] = value
                for side in ('top', 'right', 'bottom', 'left'):
                    value = points(attr(page, 'fo', 'margin-' + side, ''))
                    if value is not None:
                        doc['layout']['margins'][side] = value
            for region in ('header', 'footer'):
                element = root.find('.//style:' + region, ns)
                if element is not None:
                    doc['layout'][region] = '\n'.join(
                        ''.join(p.itertext()) for p in element.findall('text:p', ns)
                    )

        def style_values(name, seen=None):
            seen = set() if seen is None else seen
            if not name or name in seen:
                return {}
            seen.add(name)
            own = doc['styles'].get(name, {})
            return {**style_values(own.get('based_on'), seen), **own}

        inline_keys = {
            'bold',
            'italic',
            'underline',
            'strikethrough',
            'font_family',
            'font_size',
            'color',
            'highlight',
            'superscript',
            'subscript',
        }
        spaces = [0]

        def rich_text(element, block):
            def walk(node, depth=0):
                if depth > MAX_DEPTH:
                    raise ByaxError('resource_limit', 'ODT inline nesting is too deep')
                start = len(block['text'])
                if node.text:
                    block['text'] += node.text
                for child in node:
                    name = child.tag.rsplit('}', 1)[-1]
                    if name == 's':
                        count = integer(child, 'text', 'c', '1')
                        spaces[0] += max(count, 0)
                        if count < 0 or count > 100000 or spaces[0] > MAX_REPEATED_SPACES:
                            raise ByaxError('resource_limit', 'ODT repeated-space limit exceeded')
                        block['text'] += ' ' * count
                    elif name == 'tab':
                        block['text'] += '\t'
                    elif name == 'line-break':
                        block['text'] += '\n'
                    elif name == 'note':
                        body = child.find('text:note-body', ns)
                        note = _block(
                            '\n'.join(''.join(p.itertext()) for p in body.findall('text:p', ns))
                            if body is not None
                            else '',
                            'footnote',
                            reference=block['id'],
                        )
                        pending_notes.append(note)
                    elif name == 'frame':
                        picture = child.find('draw:image', ns)
                        if picture is not None:
                            source = attr(picture, 'xlink', 'href', '')
                            image = _block(
                                kind='image',
                                alt='',
                                width=points(attr(child, 'svg', 'width', '4in')),
                                height=points(attr(child, 'svg', 'height', '3in')),
                            )
                            if source in archive.namelist():
                                ext = Path(source).suffix.lower().lstrip('.')
                                mime = 'image/' + {'jpg': 'jpeg', 'svg': 'svg+xml'}.get(ext, ext)
                                image['attrs']['asset'] = _asset(
                                    doc, archive.read(source), mime, Path(source).name
                                )
                            else:
                                image['attrs']['source'] = source
                                link_external_image(doc, image, source)
                                _loss(
                                    losses,
                                    'odt_external_image',
                                    'External image was not fetched.',
                                    image['id'],
                                )
                            pending_images.append(image)
                    else:
                        child_start = len(block['text'])
                        walk(child, depth + 1)
                        if name == 'a':
                            href = attr(child, 'xlink', 'href', '')
                            if _safe_url(href):
                                block['marks'].append(
                                    {
                                        'start': child_start,
                                        'end': len(block['text']),
                                        'type': 'link',
                                        'value': href,
                                    }
                                )
                    if child.tail:
                        block['text'] += child.tail
                for key, value in style_values(attr(node, 'text', 'style-name')).items():
                    if key in inline_keys:
                        block['marks'].append(
                            {'start': start, 'end': len(block['text']), 'type': key, 'value': value}
                        )

            walk(element)

        pending_notes, pending_images = [], []

        def paragraph(element, list_level=0, list_kind='bullet'):
            name = element.tag.rsplit('}', 1)[-1]
            block = _block(kind='heading' if name == 'h' else 'paragraph')
            if name == 'h':
                block['attrs']['level'] = _heading_level(
                    integer(element, 'text', 'outline-level', '1'), losses, block['id']
                )
            style = attr(element, 'text', 'style-name')
            if style in doc['styles']:
                block['style'] = style
            if list_level:
                block['attrs'].update(list=list_kind, level=list_level)
            rich_text(element, block)
            return block

        def visit(element, level=0, list_kind='bullet', depth=0):
            if depth > MAX_DEPTH:
                raise ByaxError('resource_limit', 'ODT structure nesting is too deep')
            name = element.tag.rsplit('}', 1)[-1]
            if name in ('p', 'h'):
                doc['blocks'].append(paragraph(element, level, list_kind))
                doc['blocks'].extend(pending_images)
                doc['blocks'].extend(pending_notes)
                pending_images.clear()
                pending_notes.clear()
            elif name in ('list', 'list-item', 'section'):
                if name == 'list':
                    list_kind = list_styles.get(attr(element, 'text', 'style-name'), list_kind)
                for child in element:
                    visit(child, level + (1 if name == 'list' else 0), list_kind, depth + 1)
                if name == 'section':
                    _loss(
                        losses,
                        'odt_section_geometry',
                        'Section content retained; section-specific geometry requires review.',
                        'body/section',
                    )
            elif name == 'table':
                block = _block(kind='table', header_rows=0)
                block['rows'] = []
                row_elements = list(element.findall('table:table-row', ns))
                header = element.find('table:table-header-rows', ns)
                if header is not None:
                    headers = list(header.findall('table:table-row', ns))
                    row_elements = headers + row_elements
                    block['attrs']['header_rows'] = len(headers)
                total = 0
                for row in row_elements:
                    cells = []
                    row_repeat = integer(row, 'table', 'number-rows-repeated', '1')
                    if not 1 <= row_repeat <= 1000:
                        raise ByaxError('resource_limit', 'ODT repeated-row limit exceeded')
                    for cell in row:
                        count = integer(cell, 'table', 'number-columns-repeated', '1')
                        if not 1 <= count <= 1000:
                            raise ByaxError('resource_limit', 'ODT repeated-column limit exceeded')
                        # Check before expanding so repeats cannot allocate first.
                        if total + (len(cells) + count) * row_repeat > MAX_TABLE_CELLS:
                            raise ByaxError('resource_limit', 'ODT table exceeds 100,000 cells')
                        paragraphs = [paragraph(p) for p in cell.findall('text:p', ns)]
                        value = {'id': _id('cell'), 'text': '', 'marks': [], 'attrs': {}}
                        for p in paragraphs:
                            offset = len(value['text'])
                            value['text'] += ('\n' if offset else '') + p['text']
                            shift = offset + (1 if offset else 0)
                            value['marks'].extend(
                                {**m, 'start': m['start'] + shift, 'end': m['end'] + shift}
                                for m in p['marks']
                            )
                        for key, xml in (
                            ('rowspan', 'number-rows-spanned'),
                            ('colspan', 'number-columns-spanned'),
                        ):
                            if attr(cell, 'table', xml):
                                value['attrs'][key] = _span(attr(cell, 'table', xml))
                        for _ in range(count):
                            cells.append({**copy.deepcopy(value), 'id': _id('cell')})
                    for _ in range(row_repeat):
                        block['rows'].append(
                            {
                                'id': _id('row'),
                                'cells': [{**copy.deepcopy(c), 'id': _id('cell')} for c in cells],
                            }
                        )
                    total += len(cells) * row_repeat
                doc['blocks'].append(block)
            elif name in ('sequence-decls', 'user-field-decls', 'variable-decls'):
                _loss(
                    losses,
                    'odt_fields',
                    'ODT field definitions retained only in original attachment.',
                    'body/' + name,
                )
            else:
                text = ''.join(element.itertext()).strip()
                if text:
                    doc['blocks'].append(_block(text))
                _loss(
                    losses,
                    'odt_structure',
                    f'Unsupported ODT {name} flattened to text and preserved in original.',
                    'body/' + name,
                )

        body = roots[0].find('.//office:text', ns)
        if body is None:
            raise ByaxError('invalid_input', 'ODT lacks text body')
        for element in body:
            visit(element)


def _import_doc(path):
    configured = os.environ.get('BYAX_LIBREOFFICE')
    if not configured or not Path(configured).is_file():
        raise ByaxError(
            'adapter_unavailable',
            'DOC import requires the absolute path to a local LibreOffice executable in BYAX_LIBREOFFICE',
        )
    with tempfile.TemporaryDirectory(prefix='byax-doc-') as temporary:
        root = Path(temporary)
        profile = root / 'profile'
        (profile / 'user').mkdir(parents=True)
        (profile / 'user' / 'registrymodifications.xcu').write_text(
            '<?xml version="1.0" encoding="UTF-8"?><oor:items xmlns:oor="http://openoffice.org/2001/registry"><item oor:path="/org.openoffice.Office.Common/Security/Scripting"><prop oor:name="MacroSecurityLevel" oor:op="fuse"><value>3</value></prop></item></oor:items>',
            encoding='utf-8',
        )
        try:
            result = subprocess.run(
                [
                    str(Path(configured).resolve()),
                    '-env:UserInstallation=' + profile.as_uri(),
                    '--headless',
                    '--nologo',
                    '--nodefault',
                    '--norestore',
                    '--convert-to',
                    'docx',
                    '--outdir',
                    str(root),
                    str(path.resolve()),
                ],
                capture_output=True,
                timeout=90,
                check=False,
                shell=False,
            )
        except subprocess.TimeoutExpired as exc:
            raise ByaxError('adapter_timeout', 'Local DOC conversion exceeded 90 seconds') from exc
        output = root / (path.stem + '.docx')
        if result.returncode or not output.is_file():
            raise ByaxError('adapter_failed', 'Local DOC conversion failed', returncode=result.returncode)
        doc, report = import_document(output)
        # The outer import attaches the original .doc; drop the intermediate .docx copy.
        doc['assets'].pop(doc['metadata']['import']['original_asset'], None)
        _loss(
            report['losses'],
            'doc_adapter',
            'DOC converted by explicitly configured local LibreOffice; cross-application fidelity is not verified.',
            'source',
        )
        return doc, report


def _plain(doc):
    parts = []
    for b in doc['blocks']:
        if b.get('type') == 'table':
            parts.extend('\t'.join(c.get('text', '') for c in r.get('cells', [])) for r in b.get('rows', []))
        elif b.get('type') == 'image':
            parts.append('[Image: ' + b.get('attrs', {}).get('alt', '') + ']')
        else:
            parts.append(b.get('text', ''))
    return '\n'.join(parts)


def _html_inline(block):
    text = block.get('text', '')
    marks = block.get('marks', [])
    segments = sorted(
        {0, len(text), *[max(0, min(len(text), int(m.get(k, 0)))) for m in marks for k in ('start', 'end')]}
    )
    result = []
    for start, end in itertools.pairwise(segments):
        value = html.escape(text[start:end]).replace('\n', '<br>')
        for mark in marks:
            if mark.get('start', 0) <= start and mark.get('end', 0) >= end:
                kind, setting = mark.get('type'), mark.get('value', True)
                tags = {
                    'bold': 'strong',
                    'italic': 'em',
                    'underline': 'u',
                    'strikethrough': 's',
                    'code': 'code',
                    'superscript': 'sup',
                    'subscript': 'sub',
                }
                if kind in tags and setting:
                    value = f'<{tags[kind]}>{value}</{tags[kind]}>'
                elif kind == 'link' and _safe_url(str(setting)):
                    value = f'<a href="{html.escape(str(setting), quote=True)}">{value}</a>'
        result.append(value)
    return ''.join(result)


def _html(doc):
    result = [
        '<!doctype html><html lang="en"><head><meta charset="utf-8"><title>'
        + html.escape(doc.get('metadata', {}).get('title', ''))
        + '</title></head><body>'
    ]
    for block in doc['blocks']:
        kind, attrs = block.get('type'), block.get('attrs', {})
        if kind == 'table':
            result.append('<table>')
            for index, row in enumerate(block.get('rows', [])):
                tag = 'th' if index < attrs.get('header_rows', 0) else 'td'
                result.append(
                    '<tr>'
                    + ''.join(f'<{tag}>' + _html_inline(cell) + f'</{tag}>' for cell in row.get('cells', []))
                    + '</tr>'
                )
            result.append('</table>')
        elif kind == 'image':
            asset = doc.get('assets', {}).get(attrs.get('asset'), {})
            mime = asset.get('mime', '')
            if mime in ('image/png', 'image/jpeg', 'image/gif', 'image/webp'):
                result.append(
                    f'<img src="data:{mime};base64,{asset.get("data", "")}" alt="{html.escape(attrs.get("alt", ""), quote=True)}">'
                )
            else:
                result.append('<p>' + html.escape(attrs.get('alt', '[unavailable image]')) + '</p>')
        else:
            tag = (
                'h' + str(max(1, min(6, attrs.get('level', 1))))
                if kind == 'heading'
                else 'pre'
                if block.get('style') == 'code'
                else 'blockquote'
                if block.get('style') == 'quotation'
                else 'p'
            )
            content = _html_inline(block)
            if attrs.get('list'):
                lt = 'ol' if _list_kind(attrs) in ('ordered', 'numbered') else 'ul'
                result.append(f'<{lt}><li>{content}</li></{lt}>')
            else:
                result.append(f'<{tag}>{content}</{tag}>')
    return '\n'.join(result) + '</body></html>'


def _md_inline(block):
    text = block.get('text', '')

    def escape(value):
        return re.sub(r'([\\`*_{}\[\]<>])', r'\\\1', value)

    marks = block.get('marks', [])
    if not marks:
        return escape(text)
    # Segmenting marks keeps content intact even with overlapping rich ranges.
    boundaries = sorted(
        {0, len(text), *[max(0, min(len(text), int(m.get(k, 0)))) for m in marks for k in ('start', 'end')]}
    )
    output = []
    for start, end in itertools.pairwise(boundaries):
        value = escape(text[start:end])
        for mark in marks:
            if mark.get('start', 0) <= start and mark.get('end', 0) >= end:
                kind = mark.get('type')
                if kind in ('bold', 'italic', 'code') and mark.get('value', True):
                    delimiter = {'bold': '**', 'italic': '*', 'code': '`'}[kind]
                    value = delimiter + value + delimiter
                elif kind == 'link' and _safe_url(str(mark.get('value', ''))):
                    value = '[' + value + '](' + str(mark['value']).replace(')', '%29') + ')'
        output.append(value)
    return ''.join(output)


def _md_block(block):
    attrs, kind = block.get('attrs', {}), block.get('type')
    if kind == 'table':
        rows = [
            '| '
            + ' | '.join(c.get('text', '').replace('|', '\\|').replace('\n', ' ') for c in r.get('cells', []))
            + ' |'
            for r in block.get('rows', [])
        ]
        if rows:
            columns = len(block['rows'][0]['cells'])
            rows.insert(1, '| ' + ' | '.join(['---'] * columns) + ' |')
        return '\n'.join(rows)
    if kind == 'image':
        return '![' + attrs.get('alt', '').replace(']', '\\]') + '](' + attrs.get('source', '') + ')'
    if attrs.get('thematic_break'):
        return '---'
    if block.get('style') == 'code':
        fence = '`' * max(3, max((len(v) for v in re.findall(r'`+', block.get('text', ''))), default=0) + 1)
        return fence + attrs.get('code_language', '') + '\n' + block.get('text', '') + '\n' + fence
    text = _md_inline(block)
    if kind == 'heading':
        text = '#' * max(1, min(6, attrs.get('level', 1))) + ' ' + text
    if block.get('style') == 'quotation':
        text = '\n'.join('> ' + line for line in text.split('\n'))
    if attrs.get('list'):
        text = (
            '  ' * (_list_level(attrs) - 1)
            + ('1. ' if _list_kind(attrs) in ('ordered', 'numbered') else '- ')
            + text
        )
    return text


def _markdown(doc, losses):
    provenance = doc.get('metadata', {}).get('import', {})
    blocks = doc['blocks']
    if provenance.get('format') in ('md', 'markdown') and [b['id'] for b in blocks] == provenance.get(
        'block_order'
    ):
        source = provenance['source']
        fingerprints = provenance.get('fingerprints', {})
        changed = [b for b in blocks if _fingerprint(b) != fingerprints.get(b['id'])]
        if not changed:
            return source
        edits = []
        source_lines = source.splitlines(keepends=True)
        for block in changed:
            span = block.get('attrs', {}).get('source_lines')
            if not span or block.get('attrs', {}).get('quote_depth') or block.get('attrs', {}).get('list'):
                break
            ending = provenance.get('line_ending', '\n')
            replacement = _md_block(block).replace('\n', ending)
            if span[1] <= len(source_lines) and source_lines[span[1] - 1].endswith(('\n', '\r')):
                replacement += ending
            edits.append((span[0], span[1], replacement))
        else:
            for start, end, replacement in sorted(edits, reverse=True):
                source_lines[start:end] = [replacement]
            return ''.join(source_lines)
    _loss(
        losses,
        'markdown_normalized',
        'Markdown regenerated using CommonMark plus tables; source whitespace/syntax may change.',
        'document',
    )
    return '\n\n'.join(_md_block(b) for b in blocks) + '\n'


_XML_INVALID = re.compile('[\x00-\x08\x0b\x0c\x0e-\x1f\ufffe\uffff\ud800-\udfff]')


def _xml_safe(doc, losses):
    """Replace characters XML cannot carry one-for-one so mark offsets stay valid."""
    found = []

    def clean(value):
        if isinstance(value, str):
            if _XML_INVALID.search(value):
                found.append(True)
                return _XML_INVALID.sub('\ufffd', value)
            return value
        if isinstance(value, dict):
            return {key: clean(item) for key, item in value.items()}
        if isinstance(value, list):
            return [clean(item) for item in value]
        return value

    cleaned = {
        **doc,
        **{key: clean(doc[key]) for key in ('blocks', 'metadata', 'layout', 'styles') if key in doc},
    }
    if found:
        _loss(
            losses,
            'xml_control_characters',
            'Control characters that DOCX cannot store were replaced with U+FFFD.',
            'document',
        )
    return cleaned


def _export_docx(doc, path):
    from docx import Document
    from docx.shared import Pt, RGBColor

    word = Document()
    from docx.enum.style import WD_STYLE_TYPE

    named = {}
    for name, values in doc.get('styles', {}).items():
        actual = 'Normal' if name == 'body' else name
        style = (
            word.styles[actual]
            if actual in word.styles
            else word.styles.add_style(
                actual, WD_STYLE_TYPE.CHARACTER if values.get('family') == 'text' else WD_STYLE_TYPE.PARAGRAPH
            )
        )
        named[name] = style
        for key in ('bold', 'italic', 'underline'):
            if key in values:
                setattr(style.font, key, bool(values[key]))
        if values.get('font_family'):
            style.font.name = values['font_family']
        if values.get('font_size'):
            style.font.size = Pt(values['font_size'])
        if values.get('color') and re.fullmatch(r'#[0-9a-fA-F]{6}', values['color']):
            style.font.color.rgb = RGBColor.from_string(values['color'][1:])
        if style.type == WD_STYLE_TYPE.PARAGRAPH:
            for key, target in (
                ('space_before', 'space_before'),
                ('space_after', 'space_after'),
                ('indent', 'left_indent'),
                ('first_line_indent', 'first_line_indent'),
            ):
                if key in values:
                    setattr(style.paragraph_format, target, Pt(values[key]))
    for name, values in doc.get('styles', {}).items():
        if values.get('based_on') in named:
            named[name].base_style = named[values['based_on']]
    section = word.sections[0]
    layout = doc.get('layout', {})
    section.page_width, section.page_height = Pt(layout.get('width', 612)), Pt(layout.get('height', 792))
    for key, value in layout.get('margins', {}).items():
        if key in ('top', 'right', 'bottom', 'left'):
            setattr(section, key + '_margin', Pt(value))
    for region in ('header', 'footer'):
        value = layout.get(region, '')
        getattr(section, region).paragraphs[0].text = (
            value if isinstance(value, str) else value.get('text', '')
        )
    word.core_properties.title = doc.get('metadata', {}).get('title', '')
    word.core_properties.author = doc.get('metadata', {}).get('author', '')
    for block in doc['blocks']:
        kind, attrs = block.get('type'), block.get('attrs', {})
        if kind == 'table':
            rows = block.get('rows', [])
            columns = max((len(row['cells']) for row in rows), default=0)
            if not rows or not columns:
                continue
            table = word.add_table(rows=len(rows), cols=columns)
            table.style = 'Table Grid'
            for r, row in enumerate(rows):
                for c, cell in enumerate(row['cells']):
                    table.cell(r, c).text = cell.get('text', '')
            if attrs.get('header_rows'):
                from docx.oxml import OxmlElement

                for row in table.rows[: int(attrs['header_rows'])]:
                    row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
            continue
        if kind == 'image':
            asset = doc.get('assets', {}).get(attrs.get('asset'))
            if asset:
                word.add_picture(
                    io.BytesIO(base64.b64decode(asset['data'])), width=Pt(attrs.get('width', 288))
                )
            continue
        if kind in ('page_break', 'section_break'):
            word.add_page_break()
            continue
        paragraph = (
            word.add_heading(level=max(1, min(6, attrs.get('level', 1))))
            if kind == 'heading'
            else word.add_paragraph()
        )
        if block.get('style') in named and (kind != 'heading' or block['style'] != 'body'):
            paragraph.style = named[block['style']]
        if attrs.get('list'):
            paragraph.style = 'List Number' if _list_kind(attrs) in ('ordered', 'numbered') else 'List Bullet'
        if attrs.get('alignment'):
            from docx.enum.text import WD_ALIGN_PARAGRAPH

            paragraph.alignment = {
                'left': WD_ALIGN_PARAGRAPH.LEFT,
                'right': WD_ALIGN_PARAGRAPH.RIGHT,
                'center': WD_ALIGN_PARAGRAPH.CENTER,
                'justify': WD_ALIGN_PARAGRAPH.JUSTIFY,
            }.get(attrs['alignment'])
        for name, prop in (
            ('space_before', 'space_before'),
            ('space_after', 'space_after'),
            ('indent', 'left_indent'),
            ('first_line_indent', 'first_line_indent'),
        ):
            if name in attrs:
                setattr(paragraph.paragraph_format, prop, Pt(attrs[name]))
        text, marks = block.get('text', ''), block.get('marks', [])
        boundaries = sorted(
            {
                0,
                len(text),
                *[max(0, min(len(text), int(m.get(k, 0)))) for m in marks for k in ('start', 'end')],
            }
        )
        for start, end in itertools.pairwise(boundaries):
            run = paragraph.add_run(text[start:end])
            for mark in marks:
                if mark.get('start', 0) <= start and mark.get('end', 0) >= end:
                    name, value = mark.get('type'), mark.get('value', True)
                    if name in ('bold', 'italic', 'underline'):
                        setattr(run, name, bool(value))
                    elif name in ('strikethrough', 'superscript', 'subscript'):
                        setattr(run.font, 'strike' if name == 'strikethrough' else name, bool(value))
                    elif name in ('font_family', 'font'):
                        run.font.name = str(value)
                    elif name in ('font_size', 'size'):
                        run.font.size = Pt(value)
                    elif name == 'color' and re.fullmatch(r'#[0-9a-fA-F]{6}', str(value)):
                        run.font.color.rgb = RGBColor.from_string(value[1:])
            link = next(
                (
                    m.get('value')
                    for m in marks
                    if m.get('type') == 'link' and m.get('start', 0) <= start and m.get('end', 0) >= end
                ),
                None,
            )
            if link and _safe_url(str(link)):
                from docx.opc.constants import RELATIONSHIP_TYPE
                from docx.oxml import OxmlElement
                from docx.oxml.ns import qn

                anchor = OxmlElement('w:hyperlink')
                if str(link).startswith('#'):
                    anchor.set(qn('w:anchor'), str(link)[1:])
                else:
                    anchor.set(
                        qn('r:id'),
                        paragraph.part.relate_to(str(link), RELATIONSHIP_TYPE.HYPERLINK, is_external=True),
                    )
                anchor.append(run._r)
                paragraph._p.append(anchor)
    word.save(path)


def _export_odt(doc, path):
    from odf.dc import Title
    from odf.draw import Frame, Image
    from odf.meta import InitialCreator
    from odf.opendocument import OpenDocumentText
    from odf.style import (
        Footer,
        Header,
        MasterPage,
        PageLayout,
        PageLayoutProperties,
        ParagraphProperties,
        Style,
        TextProperties,
    )
    from odf.table import CoveredTableCell, Table, TableCell, TableHeaderRows, TableRow
    from odf.text import (
        A,
        H,
        List,
        ListItem,
        ListLevelStyleBullet,
        ListLevelStyleNumber,
        ListStyle,
        Note,
        NoteBody,
        NoteCitation,
        P,
        Span,
    )

    output = OpenDocumentText()
    output.meta.addElement(Title(text=doc.get('metadata', {}).get('title', '')))
    output.meta.addElement(InitialCreator(text=doc.get('metadata', {}).get('author', '')))
    style_map, mark_styles = {}, {}

    def text_properties(values):
        properties = {}
        for name, odf_name in (
            ('font_family', 'fontfamily'),
            ('color', 'color'),
            ('highlight', 'backgroundcolor'),
        ):
            if name in values:
                properties[odf_name] = values[name]
        if values.get('font_size'):
            properties['fontsize'] = f'{values["font_size"]}pt'
        if 'bold' in values:
            properties['fontweight'] = 'bold' if values['bold'] else 'normal'
        if 'italic' in values:
            properties['fontstyle'] = 'italic' if values['italic'] else 'normal'
        if 'underline' in values:
            properties['textunderlinestyle'] = 'solid' if values['underline'] else 'none'
        if 'strikethrough' in values:
            properties['textlinethroughstyle'] = 'solid' if values['strikethrough'] else 'none'
        if values.get('superscript'):
            properties['textposition'] = 'super 58%'
        if values.get('subscript'):
            properties['textposition'] = 'sub 58%'
        return properties

    for name, values in doc.get('styles', {}).items():
        if not isinstance(values, dict):
            continue
        style = Style(name=name, family='text' if values.get('family') == 'text' else 'paragraph')
        properties = text_properties(values)
        if properties:
            style.addElement(TextProperties(**properties))
        paragraph = {}
        if values.get('alignment'):
            paragraph['textalign'] = values['alignment']
        for key, target in (
            ('space_before', 'margintop'),
            ('space_after', 'marginbottom'),
            ('indent', 'marginleft'),
            ('first_line_indent', 'textindent'),
        ):
            if key in values:
                paragraph[target] = f'{values[key]}pt'
        if paragraph and values.get('family') != 'text':
            style.addElement(ParagraphProperties(**paragraph))
        output.styles.addElement(style)
        style_map[name] = style
    for name, values in doc.get('styles', {}).items():
        if values.get('based_on') in style_map and name in style_map:
            style_map[name].setAttribute('parentstylename', style_map[values['based_on']])
    layout = doc.get('layout', {})
    page = PageLayout(name='ByaxPage')
    props = {'pagewidth': f'{layout.get("width", 612)}pt', 'pageheight': f'{layout.get("height", 792)}pt'}
    props.update(
        {
            'margin' + side: f'{layout.get("margins", {}).get(side, 72)}pt'
            for side in ('top', 'right', 'bottom', 'left')
        }
    )
    page.addElement(PageLayoutProperties(**props))
    output.automaticstyles.addElement(page)
    master = MasterPage(name='Standard', pagelayoutname=page)
    for name, cls in (('header', Header), ('footer', Footer)):
        value = layout.get(name)
        if value:
            region = cls()
            region.addElement(P(text=value if isinstance(value, str) else value.get('text', '')))
            master.addElement(region)
    output.masterstyles.addElement(master)
    for kind in ('bullet', 'ordered'):
        style = ListStyle(name='ByaxList' + kind)
        for level in range(1, 10):
            style.addElement(
                ListLevelStyleNumber(level=level, numformat='1', numsuffix='.')
                if kind == 'ordered'
                else ListLevelStyleBullet(level=level, bulletchar='•')
            )
        output.styles.addElement(style)

    def add_text(parent, block):
        text, marks = block.get('text', ''), block.get('marks', [])
        boundaries = sorted(
            {
                0,
                len(text),
                *[max(0, min(len(text), int(m.get(k, 0)))) for m in marks for k in ('start', 'end')],
            }
        )
        for start, end in itertools.pairwise(boundaries):
            values = {
                m['type']: m.get('value', True)
                for m in marks
                if m.get('start', 0) <= start and m.get('end', 0) >= end
            }
            props = text_properties(values)
            container = parent
            if props:
                key = json.dumps(props, sort_keys=True)
                if key not in mark_styles:
                    style = Style(name='ByaxInline' + str(len(mark_styles)), family='text')
                    style.addElement(TextProperties(**props))
                    output.automaticstyles.addElement(style)
                    mark_styles[key] = style
                container = Span(stylename=mark_styles[key])
                parent.addElement(container)
            if values.get('link') and _safe_url(str(values['link'])):
                anchor = A(href=str(values['link']))
                container.addElement(anchor)
                container = anchor
            from odf.teletype import addTextToElement

            addTextToElement(container, text[start:end])

    notes = {}
    for index, note in enumerate(b for b in doc['blocks'] if b.get('type') == 'footnote'):
        notes.setdefault(note.get('attrs', {}).get('reference'), []).append((index + 1, note))
    list_stack, last_list_kind = [], None
    for block in doc['blocks']:
        kind, attrs = block.get('type'), block.get('attrs', {})
        if kind == 'footnote' and attrs.get('reference'):
            continue
        if kind == 'table':
            table = Table(name=block['id'])
            covered = set()
            headers = TableHeaderRows() if attrs.get('header_rows') else None
            if headers is not None:
                table.addElement(headers)
            rows = block.get('rows', [])
            for r, row in enumerate(rows):
                tr = TableRow()
                for c, cell in enumerate(row.get('cells', [])):
                    if (r, c) in covered:
                        tr.addElement(CoveredTableCell())
                        continue
                    ca = cell.get('attrs', {})
                    # Spans never reach past the table, so covered stays bounded.
                    rs = max(1, min(int(ca.get('rowspan', 1)), len(rows) - r))
                    cs = max(1, min(int(ca.get('colspan', 1)), len(row.get('cells', [])) - c))
                    td = TableCell(numberrowsspanned=rs, numbercolumnsspanned=cs)
                    paragraph = P()
                    add_text(paragraph, cell)
                    td.addElement(paragraph)
                    tr.addElement(td)
                    covered.update(
                        (rr, cc) for rr in range(r, r + rs) for cc in range(c, c + cs) if (rr, cc) != (r, c)
                    )
                (headers if headers is not None and r < attrs['header_rows'] else table).addElement(tr)
            output.text.addElement(table)
            list_stack = []
            continue
        if kind == 'image':
            asset = doc.get('assets', {}).get(attrs.get('asset'))
            paragraph = P()
            if asset:
                href = output.addPictureFromString(
                    base64.b64decode(asset['data']), asset.get('mime', 'image/png')
                )
                frame = Frame(
                    anchortype='as-char',
                    width=f'{attrs.get("width", 288)}pt',
                    height=f'{attrs.get("height", 216)}pt',
                )
                frame.addElement(Image(href=href, type='simple', show='embed', actuate='onLoad'))
                paragraph.addElement(frame)
            else:
                paragraph.addText('[Image: ' + attrs.get('alt', '') + ']')
            output.text.addElement(paragraph)
            continue
        element = H(outlinelevel=attrs.get('level', 1)) if kind == 'heading' else P()
        if block.get('style') in style_map:
            element.setAttribute('stylename', style_map[block['style']])
        add_text(element, block)
        for number, note in notes.get(block.get('id'), []):
            note_element = Note(id=note['id'], noteclass='footnote')
            note_element.addElement(NoteCitation(text=str(number)))
            note_body = NoteBody()
            p = P()
            add_text(p, note)
            note_body.addElement(p)
            note_element.addElement(note_body)
            element.addElement(note_element)
        if attrs.get('list'):
            list_kind = 'ordered' if _list_kind(attrs) in ('ordered', 'numbered') else 'bullet'
            level = max(1, min(9, _list_level(attrs)))
            if last_list_kind != list_kind:
                list_stack = []
            while len(list_stack) > level:
                list_stack.pop()
            while len(list_stack) < level:
                listing = List(stylename='ByaxList' + list_kind)
                if list_stack:
                    item = ListItem()
                    list_stack[-1].addElement(item)
                    item.addElement(listing)
                else:
                    output.text.addElement(listing)
                list_stack.append(listing)
            item = ListItem()
            item.addElement(element)
            list_stack[-1].addElement(item)
            last_list_kind = list_kind
        else:
            output.text.addElement(element)
            list_stack = []
    output.save(str(path), addsuffix=False)


def _rtf_escape(text):
    result = []
    for char in text:
        if char in '\\{}':
            result.append('\\' + char)
        elif char == '\n':
            result.append('\\line ')
        elif char == '\t':
            result.append('\\tab ')
        elif 32 <= ord(char) < 127:
            result.append(char)
        else:
            raw = char.encode('utf-16-le')
            for index in range(0, len(raw), 2):
                code = int.from_bytes(raw[index : index + 2], 'little', signed=True)
                result.append(f'\\u{code}?')
    return ''.join(result)


def _rtf(doc):
    output = ['{\\rtf1\\ansi\\deff0{\\fonttbl{\\f0 Times New Roman;}}\\uc1 ']
    for block in doc['blocks']:
        if block.get('type') == 'table':
            for row in block.get('rows', []):
                output.append(
                    _rtf_escape('\t'.join(c.get('text', '') for c in row.get('cells', []))) + '\\par\n'
                )
            continue
        text = block.get('text', '')
        marks = block.get('marks', [])
        boundaries = sorted(
            {
                0,
                len(text),
                *[max(0, min(len(text), int(m.get(k, 0)))) for m in marks for k in ('start', 'end')],
            }
        )
        for start, end in itertools.pairwise(boundaries):
            controls = []
            for mark in marks:
                if mark.get('start', 0) <= start and mark.get('end', 0) >= end and mark.get('value', True):
                    control = {
                        'bold': 'b',
                        'italic': 'i',
                        'underline': 'ul',
                        'strikethrough': 'strike',
                        'superscript': 'super',
                        'subscript': 'sub',
                    }.get(mark.get('type'))
                    if control:
                        controls.append('\\' + control + ' ')
            output.append('{' + ''.join(controls) + _rtf_escape(text[start:end]) + '}')
        output.append('\\par\n')
    return ''.join(output) + '}'


def _exchange_losses(doc, fmt, losses):
    if fmt == 'byax':
        return
    for category in ('review', 'references', 'fields', 'layers'):
        value = doc.get(category)
        if category == 'review':
            value = any(value.values()) if value else False
        if value and not (fmt == 'pdf' and category == 'layers'):
            _loss(
                losses,
                category + '_not_exported',
                f'{category.title()} state is not preserved by this {fmt.upper()} adapter.',
                category,
            )
    if fmt != 'pdf':
        if doc.get('styles') and fmt not in ('docx', 'odt'):
            _loss(
                losses,
                'named_styles',
                'Named style definitions and inheritance are not preserved; supported direct formatting is exported.',
                'styles',
            )
        if fmt not in ('docx', 'odt'):
            _loss(
                losses,
                'page_layout',
                'Page geometry, masters, repeated regions and exact pagination are not represented in this adapter.',
                'layout',
            )
    supported_marks = {
        'md': {'bold', 'italic', 'code', 'link'},
        'markdown': {'bold', 'italic', 'code', 'link'},
        'html': {'bold', 'italic', 'underline', 'strikethrough', 'superscript', 'subscript', 'code', 'link'},
        'htm': {'bold', 'italic', 'underline', 'strikethrough', 'superscript', 'subscript', 'code', 'link'},
        'odt': {
            'bold',
            'italic',
            'underline',
            'strikethrough',
            'superscript',
            'subscript',
            'font_family',
            'font_size',
            'color',
            'highlight',
            'link',
        },
        'docx': {
            'link',
            'bold',
            'italic',
            'underline',
            'strikethrough',
            'superscript',
            'subscript',
            'font_family',
            'font_size',
            'font',
            'size',
            'color',
        },
        'rtf': {'bold', 'italic', 'underline', 'strikethrough', 'superscript', 'subscript'},
    }.get(fmt, set())
    for block in doc['blocks']:
        kind, attrs = block.get('type'), block.get('attrs', {})
        if fmt != 'pdf':
            omitted_marks = {m.get('type') for m in block.get('marks', [])} - supported_marks
            if omitted_marks:
                _loss(
                    losses,
                    'inline_formatting',
                    'Unsupported inline marks: ' + ', '.join(sorted(omitted_marks)),
                    block['id'],
                )
        if (
            kind in ('footnote', 'toc', 'section_break')
            and fmt != 'pdf'
            and not (kind == 'footnote' and fmt == 'odt')
        ):
            _loss(
                losses,
                'structural_object',
                f'{kind} has no active equivalent in this export adapter; only available text is retained.',
                block['id'],
            )
        if kind == 'page_break' and fmt not in ('pdf', 'docx'):
            _loss(losses, 'page_break', 'Explicit page break is not represented.', block['id'])
        if kind == 'table':
            if fmt in ('txt', 'rtf'):
                _loss(
                    losses,
                    'table_flattened',
                    'Table exported as tab-separated text; editable grid and formatting are lost.',
                    block['id'],
                )
            elif fmt not in ('pdf',):
                _loss(
                    losses,
                    'table_subset',
                    'Table text/grid exported; widths, spans, cell formatting and pagination policies may be lost.',
                    block['id'],
                )
        if kind == 'image' and fmt not in ('pdf', 'docx', 'odt', 'html', 'htm'):
            _loss(
                losses,
                'image_not_embedded',
                'Embedded image bytes are not included; source/alternative text is retained where representable.',
                block['id'],
            )
        if attrs.get('list') and fmt not in ('pdf', 'md', 'markdown'):
            _loss(
                losses,
                'list_subset',
                'List nesting, continuation and restart settings are not preserved by this adapter.',
                block['id'],
            )
        if fmt in ('txt', 'rtf') and kind == 'heading':
            _loss(
                losses, 'heading_flattened', 'Heading text retained without heading semantics.', block['id']
            )
        if attrs.get('alignment') and fmt not in ('pdf', 'docx'):
            _loss(losses, 'paragraph_format', 'Paragraph alignment is not represented.', block['id'])


def export_document(doc, path, format=None, profile=None, *, owner=None):
    destination = Path(path)
    fmt = (format or destination.suffix.lstrip('.')).lower()
    settings = doc.get('profiles', {}).get(profile) if isinstance(profile, str) else profile or {}
    if settings is None:
        raise ByaxError('missing_target', 'Unknown export profile', profile=profile)
    if not fmt:
        fmt = settings.get('format', '')
    if fmt not in ('byax', 'txt', 'md', 'markdown', 'html', 'htm', 'rtf', 'docx', 'odt', 'pdf'):
        raise ByaxError('unsupported_format', f'Unsupported export format: {fmt}')
    if settings.get('pages') and fmt != 'pdf':
        raise ByaxError(
            'unsupported_feature',
            'Page selection is supported for controlled PDF output; editable exports retain whole semantic content',
        )
    original = doc.get('metadata', {}).get('import', {}).get('original_path')
    if original and destination.resolve() == Path(original).resolve():
        raise ByaxError(
            'conflict', 'Conversion must preserve the imported original; choose a separate output path'
        )
    candidate = copy.deepcopy(doc)
    candidate['blocks'] = [b for b in candidate['blocks'] if not b.get('attrs', {}).get('excluded')]
    if settings.get('layout'):
        candidate['layout'].update(settings['layout'])
    if settings.get('metadata') == 'omit' or settings.get('include_metadata') is False:
        candidate['metadata'] = {'title': '', 'author': '', 'language': 'en'}
    losses = []
    _exchange_losses(candidate, fmt, losses)
    checks = None
    required = settings.get('required_suites', [])
    if required:
        from .checks import run_checks

        checks = run_checks(candidate, required)
        if checks.get('validation') != 'passed':
            raise ByaxError(
                'failed_requirement' if checks.get('validation') == 'failed' else 'incomplete_coverage',
                'Selected export did not satisfy required checks',
                report=checks,
            )
    if fmt == 'pdf':
        from .layout import build_document, resolve_layout

        shaped = build_document(candidate)
        resolve_layout(shaped, candidate)
        findings = shaped.byax_findings
        if findings and any(suite in required for suite in ('print', 'publication')):
            raise ByaxError(
                'incomplete_coverage',
                'Required print coverage is incomplete for the selected renderer',
                findings=findings,
            )
    destination.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix='.byax-export-', dir=destination.parent) as temporary:
        stage = Path(temporary) / ('export.' + fmt)
        result = None
        if fmt == 'byax':
            from .storage import save_document

            save_document(candidate, stage)
        elif fmt == 'pdf':
            from .layout import export_pdf

            result = export_pdf(candidate, stage, pages=settings.get('pages'))
            losses.extend(result.get('losses', []))
        elif fmt == 'docx':
            _export_docx(_xml_safe(candidate, losses), stage)
        elif fmt == 'odt':
            _export_odt(candidate, stage)
        else:
            if fmt == 'txt':
                content = _plain(candidate)
                provenance = candidate.get('metadata', {}).get('import', {})
                if (
                    provenance.get('format') == 'txt'
                    and [b['id'] for b in candidate['blocks']] == provenance.get('block_order')
                    and all(
                        _fingerprint(b) == provenance.get('fingerprints', {}).get(b['id'])
                        for b in candidate['blocks']
                    )
                ):
                    content = provenance['source']
            elif fmt in ('md', 'markdown'):
                content = _markdown(candidate, losses)
            elif fmt in ('html', 'htm'):
                content = _html(candidate)
            else:
                content = _rtf(candidate)
            if fmt == 'txt' and settings.get('line_ending'):
                ending = {'lf': '\n', 'crlf': '\r\n'}.get(settings['line_ending'], settings['line_ending'])
                content = content.replace('\r\n', '\n').replace('\r', '\n').replace('\n', ending)
            encoding = settings.get('encoding', 'utf-8') if fmt == 'txt' else 'utf-8'
            stage.write_bytes(content.encode(encoding))
        from .storage import DocumentLock

        with DocumentLock(destination, owner=owner):
            os.replace(stage, destination)
    report = _report(fmt, losses, [str(destination.resolve())])
    if result:
        report.update(
            {
                key: value
                for key, value in result.items()
                if key not in ('outputs', 'losses', 'findings', 'coverage')
            }
        )
    if checks:
        report['checks'] = checks
        report['validation'] = checks['validation']
    return report


def _rtf_group(source, opening):
    """Return one balanced RTF group so table regexes never scan the whole body."""
    start = source.find(opening)
    if start < 0:
        return ''
    depth, index = 0, start
    while index < len(source):
        char = source[index]
        if char == '\\':
            index += 2
            continue
        if char == '{':
            depth += 1
        elif char == '}':
            depth -= 1
            if depth == 0:
                return source[start : index + 1]
        index += 1
    return source[start:]


def _import_rtf(doc, source, losses):
    """Parse the editable RTF text/style/table subset without object execution."""
    if not source.lstrip().startswith('{\\rtf'):
        raise ByaxError('invalid_input', 'File is not an RTF document')
    fonts = {
        int(match.group(1)): match.group(2).strip()
        for match in re.finditer(
            r'\{\\f(\d+)(?:\\[a-z]+-?\d* ?)*([^;{}\\]*);', _rtf_group(source, '{\\fonttbl')
        )
    }
    colors = [None]
    color_match = re.search(r'\{\\colortbl\s*;(.*?)\}', source, re.DOTALL)
    if color_match:
        for value in color_match.group(1).split(';'):
            numbers = [re.search(r'\\' + name + r'(\d+)', value) for name in ('red', 'green', 'blue')]
            if all(numbers):
                colors.append('#' + ''.join(f'{min(255, int(n.group(1))):02x}' for n in numbers))
    state = {'marks': {}, 'attrs': {}, 'uc': 1, 'skip': False, 'star': False}
    tails = {}
    stack, block, cells, table = [], _block(), [], None
    fallback, row_active, position = 0, False, 0
    skipped = set()
    destinations = {
        'fonttbl',
        'colortbl',
        'stylesheet',
        'info',
        'pict',
        'object',
        'objdata',
        'fldinst',
        'header',
        'footer',
        'headerl',
        'headerr',
        'footerl',
        'footerr',
        'footnote',
        'annotation',
        'datastore',
        'xmlopen',
        'xmlattrname',
        'xmlattrvalue',
        'themedata',
        'colorschememapping',
        'generator',
        'filetbl',
        'listtable',
        'listoverridetable',
    }

    def append(value):
        nonlocal fallback
        if state['skip']:
            return
        if fallback:
            cut = min(fallback, len(value))
            value, fallback = value[cut:], fallback - cut
        if not value:
            return
        start = len(block['text'])
        block['text'] += value
        latest = tails.setdefault(block['id'], {})
        for name, setting in state['marks'].items():
            if setting is not None and setting is not False:
                # Only the newest mark of a kind can end where this run starts;
                # scanning every mark was quadratic for alternating runs.
                key = (name, repr(setting))
                existing = latest.get(key)
                if existing is not None and existing['end'] == start:
                    existing['end'] += len(value)
                else:
                    latest[key] = {'start': start, 'end': start + len(value), 'type': name, 'value': setting}
                    block['marks'].append(latest[key])

    def normalized():
        raw = block['text']

        def scalar(value):
            return value.encode('utf-16-le', errors='surrogatepass').decode('utf-16-le', errors='replace')

        output = copy.deepcopy(block)
        output['text'] = scalar(raw)
        if len(output['text']) != len(raw):
            # Map UTF-16 offsets to scalar offsets in one pass; a surrogate pair
            # becomes one scalar. Re-encoding each prefix per mark was quadratic.
            offsets, count, index = [0] * (len(raw) + 1), 0, 0
            while index < len(raw):
                offsets[index] = count
                pair = (
                    '\ud800' <= raw[index] <= '\udbff'
                    and index + 1 < len(raw)
                    and '\udc00' <= raw[index + 1] <= '\udfff'
                )
                if pair:
                    offsets[index + 1] = count + 1
                index += 2 if pair else 1
                count += 1
            offsets[len(raw)] = count
            for mark in output['marks']:
                mark['start'], mark['end'] = offsets[mark['start']], offsets[mark['end']]
        output['attrs'].update(state['attrs'])
        return output

    def paragraph(force=False):
        nonlocal block, table
        if block['text'] or force:
            doc['blocks'].append(normalized())
            block = _block()
            table = None

    def cell():
        nonlocal block
        value = normalized()
        cells.append({'id': _id('cell'), 'text': value['text'], 'marks': value['marks'], 'attrs': {}})
        block = _block()

    token_count = 0
    control_pattern = re.compile(r'([a-zA-Z]+)(-?\d+)? ?')
    while position < len(source):
        token_count += 1
        if token_count > 500000:
            raise ByaxError('resource_limit', 'RTF exceeds 500,000 parser tokens')
        char = source[position]
        if char == '{':
            stack.append(copy.deepcopy(state))
            position += 1
            continue
        if char == '}':
            if stack:
                state = stack.pop()
            position += 1
            continue
        if char != '\\':
            end = position
            while end < len(source) and source[end] not in '{}\\':
                end += 1
            append(source[position:end].replace('\r', '').replace('\n', ''))
            position = end
            continue
        position += 1
        if position >= len(source):
            break
        char = source[position]
        if char in '{}\\':
            append(char)
            position += 1
            continue
        if char == "'" and position + 2 < len(source):
            try:
                append(
                    bytes([int(source[position + 1 : position + 3], 16)]).decode('cp1252', errors='replace')
                )
            except ValueError:
                pass
            position += 3
            continue
        if not (char.isascii() and char.isalpha()):
            if char == '*':
                state['star'] = True
            elif char == '~':
                append('\u00a0')
            elif char == '_':
                append('\u2011')
            elif char == '-':
                append('\u00ad')
            position += 1
            continue
        match = control_pattern.match(source, position)
        if match is None:
            raise ByaxError('invalid_input', 'Invalid RTF control word', target=f'offset:{position}')
        word, number = match.group(1), match.group(2)
        argument = int(number) if number is not None else 1
        position = match.end()
        if word == 'bin':
            if argument < 0 or argument > MAX_INPUT:
                raise ByaxError('resource_limit', 'RTF binary payload exceeds limit')
            position += argument
            continue
        if word in destinations or state['star']:
            state['skip'], state['star'] = True, False
            if word not in {'fonttbl', 'colortbl', 'info', 'generator'} and word not in skipped:
                skipped.add(word)
                _loss(
                    losses,
                    'rtf_' + word,
                    f'RTF {word} destination is retained in original but not imported as editable content.',
                    f'offset:{position}',
                )
            continue
        if state['skip']:
            continue
        if word == 'uc':
            state['uc'] = max(0, min(16, argument))
        elif word == 'u':
            append(chr(argument % 65536))
            fallback = state['uc']
        elif word in ('b', 'i', 'ul', 'strike', 'super', 'sub'):
            state['marks'][
                {
                    'b': 'bold',
                    'i': 'italic',
                    'ul': 'underline',
                    'strike': 'strikethrough',
                    'super': 'superscript',
                    'sub': 'subscript',
                }[word]
            ] = argument != 0
        elif word == 'ulnone':
            state['marks']['underline'] = False
        elif word == 'nosupersub':
            state['marks'].update(superscript=False, subscript=False)
        elif word == 'plain':
            state['marks'] = {}
        elif word == 'fs':
            state['marks']['font_size'] = argument / 2
        elif word == 'f' and argument in fonts:
            state['marks']['font_family'] = fonts[argument]
        elif word in ('cf', 'highlight') and 0 <= argument < len(colors):
            state['marks']['color' if word == 'cf' else 'highlight'] = colors[argument]
        elif word in ('ql', 'qr', 'qc', 'qj'):
            state['attrs']['alignment'] = {'ql': 'left', 'qr': 'right', 'qc': 'center', 'qj': 'justify'}[word]
        elif word in ('li', 'fi', 'sb', 'sa'):
            state['attrs'][
                {'li': 'indent', 'fi': 'first_line_indent', 'sb': 'space_before', 'sa': 'space_after'}[word]
            ] = argument / 20
        elif word == 'pard':
            state['attrs'] = {}
        elif word == 'par':
            if row_active:
                append('\n')
            else:
                paragraph(force=True)
        elif word == 'line':
            append('\n')
        elif word == 'tab':
            append('\t')
        elif word == 'page':
            paragraph()
            doc['blocks'].append(_block(kind='page_break'))
        elif word == 'trowd':
            if block['text']:
                paragraph()
            row_active, cells = True, []
        elif word == 'cell':
            cell()
        elif word == 'row':
            if block['text']:
                cell()
            if table is None:
                table = _block(kind='table', header_rows=0)
                table['rows'] = []
                doc['blocks'].append(table)
            table['rows'].append({'id': _id('row'), 'cells': cells})
            cells, row_active = [], False
        elif word in ('emdash', 'endash', 'bullet', 'lquote', 'rquote', 'ldblquote', 'rdblquote'):
            append(
                {
                    'emdash': '—',
                    'endash': '–',
                    'bullet': '•',
                    'lquote': '‘',
                    'rquote': '’',
                    'ldblquote': '“',
                    'rdblquote': '”',
                }[word]
            )
    paragraph()
