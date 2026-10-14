import zipfile

import pytest

from byax.errors import ByaxError
from byax.exchange import export_document, import_document
from byax.model import new_document, validate_document


def test_rtf_non_ascii_escape_and_seeded_mutations(tmp_path):
    import random

    from byax.cli import main

    path = tmp_path / 'mutated.rtf'
    original = '{\\rtf1\\ansi Hello \\é world}'
    path.write_text(original, encoding='utf-8')
    assert main(['import', str(path), '--out', str(tmp_path / 'valid.byax')]) in (0, 2)
    rng = random.Random(54)
    base = '{\\rtf1\\ansi Hello \\b bold\\b0 \\u233? café\\par Next}'
    symbols = ['é', '文', 'Я', 'α', '\\', '{', '}', '0', '*', "'", '\n']
    for _ in range(200):
        chars = list(base)
        for _ in range(rng.randint(1, 4)):
            chars.insert(rng.randrange(len(chars) + 1), rng.choice(symbols))
        path.write_text(''.join(chars), encoding='utf-8')
        try:
            doc, _ = import_document(path)
            validate_document(doc)
        except ByaxError:
            pass


def test_markdown_roundtrip_preserves_source_and_local_edits(tmp_path):
    source = b'# Title\r\n\r\nA **bold** paragraph.\r\n\r\nSecond line.\r\n'
    original = tmp_path / 'source.md'
    original.write_bytes(source)
    doc, _report = import_document(original)
    validate_document(doc)
    assert doc['blocks'][1]['marks'][0]['type'] == 'bold'
    exported = tmp_path / 'export.md'
    export_document(doc, exported)
    assert exported.read_bytes() == source
    doc['blocks'][2]['text'] = 'Changed line.'
    export_document(doc, exported)
    assert exported.read_bytes() == source.replace(b'Second line.', b'Changed line.')
    assert original.read_bytes() == source
    with pytest.raises(ByaxError, match='preserve'):
        export_document(doc, original)


def test_markdown_semantic_table_lists_and_unicode(tmp_path):
    path = tmp_path / 'example.md'
    path.write_text('# مرحبا\n\n- first\n- second\n\n| A | B |\n|---|---|\n| é | 文 |\n', encoding='utf-8')
    doc, _ = import_document(path)
    assert doc['blocks'][1]['attrs']['list'] == 'bullet'
    table = doc['blocks'][-1]
    assert table['type'] == 'table'
    assert table['rows'][1]['cells'][1]['text'] == '文'
    validate_document(doc)


def test_html_does_not_import_active_content(tmp_path):
    path = tmp_path / 'input.html'
    path.write_text(
        '<h1>Safe</h1><script>SECRET()</script><p>A <b>bold</b> <a href="javascript:alert(1)">bad</a></p><img src="https://example.invalid/a.png" alt="remote"><table><tr><th>H</th></tr><tr><td>cell</td></tr></table>'
    )
    doc, report = import_document(path)
    assert all('SECRET' not in block.get('text', '') for block in doc['blocks'])
    assert not any(mark.get('type') == 'link' for b in doc['blocks'] for mark in b.get('marks', []))
    assert doc['blocks'][-1]['rows'][1]['cells'][0]['text'] == 'cell'
    assert any(item['code'] == 'html_inactive_content' for item in report['losses'])


@pytest.mark.parametrize('fmt', ['txt', 'md', 'html', 'rtf', 'docx', 'odt'])
def test_editable_exchange_retains_unicode_text(tmp_path, fmt):
    doc = new_document('Test', 'Hello café 文')
    path = tmp_path / ('output.' + fmt)
    report = export_document(doc, path)
    imported, result = import_document(path)
    assert 'Hello café 文' in '\n'.join(b.get('text', '') for b in imported['blocks'])
    assert report['validation'] == 'not_run'
    assert result['direction'] == 'import'
    validate_document(imported)


def test_zip_traversal_and_entities_rejected(tmp_path):
    path = tmp_path / 'unsafe.docx'
    with zipfile.ZipFile(path, 'w') as archive:
        archive.writestr('../bad.xml', '<a/>')
    with pytest.raises(ByaxError) as error:
        import_document(path)
    assert error.value.code == 'unsafe_archive'
    with zipfile.ZipFile(path, 'w') as archive:
        archive.writestr('word/document.xml', '<!DOCTYPE x [<!ENTITY a "abc">]><x>&a;</x>')
    with pytest.raises(ByaxError) as error:
        import_document(path)
    assert error.value.code == 'unsafe_xml'


def test_failed_export_keeps_previous_output(tmp_path):
    target = tmp_path / 'out.txt'
    target.write_text('Original')
    with pytest.raises(UnicodeEncodeError):
        export_document(new_document('Title', '文'), target, profile={'encoding': 'ascii'})
    assert target.read_text() == 'Original'


def test_doc_requires_local_adapter(tmp_path, monkeypatch):
    path = tmp_path / 'old.doc'
    path.write_bytes(b'fake')
    monkeypatch.delenv('BYAX_LIBREOFFICE', raising=False)
    with pytest.raises(ByaxError) as error:
        import_document(path)
    assert error.value.code == 'adapter_unavailable'


@pytest.mark.parametrize('fmt', ['docx', 'odt'])
def test_rich_exchange_preserves_named_styles_marks_and_hyperlinks(tmp_path, fmt):
    doc = new_document('Rich', 'Bold linked words')
    doc['styles']['Special'] = {
        'based_on': 'body',
        'font_family': 'DejaVu Sans',
        'font_size': 16,
        'bold': True,
    }
    block = doc['blocks'][0]
    block['style'] = 'Special'
    block['marks'] = [
        {'start': 0, 'end': 4, 'type': 'bold', 'value': True},
        {'start': 5, 'end': 11, 'type': 'link', 'value': 'https://example.com'},
        {'start': 12, 'end': 17, 'type': 'italic', 'value': True},
    ]
    path = tmp_path / ('rich.' + fmt)
    export_document(doc, path)
    loaded, _ = import_document(path)
    paragraph = loaded['blocks'][0]
    assert paragraph['text'] == block['text']
    assert 'Special' in loaded['styles']
    assert loaded['styles']['Special']['font_size'] == 16
    assert any(m['type'] == 'link' and m['value'] == 'https://example.com' for m in paragraph['marks'])
    assert any(m['type'] == 'italic' and m['value'] for m in paragraph['marks'])


def test_rtf_import_formats_tables_unicode_and_does_not_execute_objects(tmp_path):
    path = tmp_path / 'formatted.rtf'
    path.write_text(
        r'{\rtf1\ansi{\fonttbl{\f0 Arial;}}\f0\fs28 Plain {\b bold} and {\i italic} \u-10179?\u-8704?\par\trowd\cellx2000\cellx4000\intbl A\cell B\cell\row{\object\objdata SECRET}}'
    )
    doc, report = import_document(path)
    paragraph = doc['blocks'][0]
    assert paragraph['text'] == 'Plain bold and italic 😀'
    assert any(
        m['type'] == 'bold' and paragraph['text'][m['start'] : m['end']] == 'bold' for m in paragraph['marks']
    )
    assert doc['blocks'][1]['rows'][0]['cells'][1]['text'] == 'B'
    assert any(loss['code'] == 'rtf_object' for loss in report['losses'])
    assert 'SECRET' not in '\n'.join(b.get('text', '') for b in doc['blocks'])


def test_odt_preserves_footnote_association(tmp_path):
    doc = new_document('Notes', 'Reference')
    doc['blocks'].append(
        {
            'id': 'note_test',
            'type': 'footnote',
            'text': 'Retained note',
            'style': 'body',
            'marks': [],
            'attrs': {'reference': doc['blocks'][0]['id']},
        }
    )
    path = tmp_path / 'notes.odt'
    export_document(doc, path)
    loaded, _ = import_document(path)
    note = next(b for b in loaded['blocks'] if b['type'] == 'footnote')
    assert note['text'] == 'Retained note'
    assert note['attrs']['reference'] == loaded['blocks'][0]['id']


def _odt(tmp_path, body, name='hostile.odt'):
    content = (
        '<office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
        'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" '
        'xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0" '
        'xmlns:style="urn:oasis:names:tc:opendocument:xmlns:style:1.0" '
        'xmlns:fo="urn:oasis:names:tc:opendocument:xmlns:xsl-fo-compatible:1.0">'
        '<office:body><office:text>' + body + '</office:text></office:body></office:document-content>'
    )
    path = tmp_path / name
    with zipfile.ZipFile(path, 'w') as archive:
        archive.writestr('mimetype', 'application/vnd.oasis.opendocument.text')
        archive.writestr('content.xml', content)
    return path


def test_odt_expansion_bombs_fail_before_allocating(tmp_path):
    spaces = _odt(tmp_path, '<text:p>' + '<text:s text:c="100000"/>' * 20 + '</text:p>', 'spaces.odt')
    with pytest.raises(ByaxError) as error:
        import_document(spaces)
    assert error.value.code == 'resource_limit'
    cells = ''.join(
        '<table:table-cell table:number-columns-repeated="1000"><text:p>x</text:p></table:table-cell>'
        for _ in range(10)
    )
    table = _odt(
        tmp_path,
        f'<table:table><table:table-row table:number-rows-repeated="1000">{cells}</table:table-row></table:table>',
        'table.odt',
    )
    with pytest.raises(ByaxError) as error:
        import_document(table)
    assert error.value.code == 'resource_limit'
    nested = _odt(
        tmp_path, '<text:p>' + '<text:span>' * 1500 + 'x' + '</text:span>' * 1500 + '</text:p>', 'nested.odt'
    )
    with pytest.raises(ByaxError) as error:
        import_document(nested)
    assert error.value.code == 'resource_limit'


def test_odt_malformed_numbers_and_deep_headings(tmp_path):
    bad = _odt(tmp_path, '<text:p><text:s text:c="lots"/></text:p>', 'bad.odt')
    with pytest.raises(ByaxError) as error:
        import_document(bad)
    assert error.value.code == 'invalid_input'
    deep = _odt(tmp_path, '<text:h text:outline-level="8">Deep</text:h>', 'deep.odt')
    doc, report = import_document(deep)
    assert doc['blocks'][0]['attrs']['level'] == 6
    assert any(loss['code'] == 'heading_level' for loss in report['losses'])


def test_docx_deep_heading_split_grapheme_and_control_characters(tmp_path):
    from docx import Document

    word = Document()
    word.add_paragraph('Deep', style='Heading 7')
    paragraph = word.add_paragraph()
    paragraph.add_run('e')
    paragraph.add_run('́x').bold = True
    path = tmp_path / 'word.docx'
    word.save(path)
    doc, _ = import_document(path)
    assert doc['blocks'][0]['attrs']['level'] == 6
    bold = next(m for m in doc['blocks'][1]['marks'] if m['type'] == 'bold')
    assert (bold['start'], bold['end']) == (0, 3)
    control = new_document('Escapes', 'Red \x1b[31m text')
    out = tmp_path / 'control.docx'
    report = export_document(control, out)
    assert any(loss['code'] == 'xml_control_characters' for loss in report['losses'])
    assert Document(out).paragraphs[-1].text == 'Red �[31m text'


def test_html_bare_attributes_optional_tags_and_title(tmp_path):
    path = tmp_path / 'loose.html'
    path.write_text(
        '<html><head><title>Ignored title</title><body><p><a href>link</a> <img src>'
        '<table><tr><td colspan>a<td colspan="²">b<tr><td rowspan="999999999">c<td>d</table>'
        '<p>Body text',
        encoding='utf-8',
    )
    doc, _ = import_document(path)
    texts = [b.get('text', '') for b in doc['blocks']]
    assert 'Ignored title' not in ''.join(texts)
    assert any('Body text' in text for text in texts)
    table = next(b for b in doc['blocks'] if b['type'] == 'table')
    assert [[c['text'] for c in row['cells']] for row in table['rows']] == [['a', 'b'], ['c', 'd']]
    assert table['rows'][1]['cells'][0]['attrs']['rowspan'] <= 1000
    export_document(doc, tmp_path / 'spans.odt')
    exported = tmp_path / 'round.html'
    export_document(new_document('MyTitle', 'hello'), exported)
    again, _ = import_document(exported)
    assert [b['text'] for b in again['blocks']] == ['hello']


def test_rtf_font_table_scan_is_linear(tmp_path):
    import time

    path = tmp_path / 'fonts.rtf'
    path.write_text(
        '{\\rtf1{\\fonttbl{\\f0\\froman\\fcharset0 Times New Roman;}}{\\f1 ' + 'a ' * 60000 + '}}',
        encoding='ascii',
    )
    started = time.monotonic()
    import_document(path)
    assert time.monotonic() - started < 3


def test_many_formatting_runs_validate_in_linear_time():
    import time

    doc = new_document(text='ab' * 20000)
    doc['blocks'][0]['marks'] = [
        {'type': 'bold', 'start': i, 'end': i + 1, 'value': True} for i in range(0, 40000, 2)
    ]
    started = time.monotonic()
    validate_document(doc)
    assert time.monotonic() - started < 3


def test_rtf_alternating_runs_import_quickly_with_scalar_offsets(tmp_path):
    import time

    path = tmp_path / 'runs.rtf'
    path.write_text(
        '{\\rtf1 '
        + ''.join(('\\b ' if i % 2 else '\\b0 ') + 'w ' for i in range(5000))
        + '\\u-10179?\\u-8704?\\b z}',
        encoding='ascii',
    )
    started = time.monotonic()
    doc, _ = import_document(path)
    assert time.monotonic() - started < 3
    block = doc['blocks'][0]
    assert block['text'].endswith('w \U0001f600z')
    assert (block['marks'][-1]['start'], block['marks'][-1]['end']) == (
        len(block['text']) - 4,
        len(block['text']),
    )
