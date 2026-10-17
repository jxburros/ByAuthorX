import threading
from concurrent.futures import ThreadPoolExecutor

import pytest

from byax.errors import ByaxError
from byax.libraries import LibraryIndex, LibraryRegistry


@pytest.fixture
def library(tmp_path):
    root = tmp_path / 'sources'
    root.mkdir()
    registry = LibraryRegistry(tmp_path / 'config')
    definition = registry.create('Research', [root])
    return root, registry, LibraryIndex(definition, tmp_path / 'cache')


def test_mixed_files_identity_and_relink(library, tmp_path):
    root, registry, index = library
    (root / 'notes.md').write_text('# People\n\nAria is the captain.')
    (root / 'copy.md').write_bytes((root / 'notes.md').read_bytes())
    (root / 'unknown.xyz').write_text('opaque')
    (root / '.hidden.txt').write_text('secret Aria')
    result = index.refresh()
    assert result['documents'] == 1
    assert any(f['status'] == 'unsupported' for f in result['files'])
    hit = index.search('Aria')['results'][0]
    assert hit['paths'] == ['copy.md', 'notes.md']
    assert hit['locator']['heading_path'] == ['People']
    (root / 'notes.md').rename(root / 'renamed.md')
    index.refresh()
    renamed = index.search('Aria')['results'][0]
    assert renamed['chunk_id'] == hit['chunk_id']
    assert 'renamed.md' in renamed['paths']
    assert index.read(hit['document_id'], hit['chunk_id'])['results'][0]['text'] == hit['text']
    registry.remove(index.library['library_id'])
    assert (root / 'renamed.md').exists()
    assert registry.relink(index.library['library_id'], [root], name='Restored')['name'] == 'Restored'


def test_stale_changed_and_deleted_content(library):
    root, _, index = library
    path = root / 'notes.txt'
    path.write_text('Aria captain')
    index.refresh()
    hit = index.search('Aria')['results'][0]
    path.write_text('Aria pilot with longer text')
    assert index.search('Aria')['results'][0]['stale']
    index.refresh()
    assert index.search('pilot')['results'][0]['document_id'] != hit['document_id']
    assert not index.search('captain')['results']
    path.unlink()
    assert not index.search('Aria')['results']
    with pytest.raises(ByaxError):
        index.read(hit['document_id'])


def test_symlinks_are_not_authorization(library, tmp_path):
    root, _, index = library
    outside = tmp_path / 'outside.txt'
    outside.write_text('private classified secret')
    (root / 'escape.txt').symlink_to(outside)
    index.refresh()
    assert not index.search('private')['results']
    (root / 'normal.txt').write_text('original words')
    index.refresh()
    hit = index.search('original')['results'][0]
    (root / 'normal.txt').unlink()
    (root / 'normal.txt').symlink_to(outside)
    with pytest.raises(ByaxError):
        index.read(hit['document_id'])


def test_cancellation_limits_and_resume(library):
    root, _, index = library
    (root / 'notes.txt').write_text('Aria captain')
    event = threading.Event()
    event.set()
    assert index.refresh(cancel_event=event)['cancelled']
    event.clear()
    assert index.refresh(cancel_event=event)['documents'] == 1
    (root / 'large.txt').write_text('x' * 100)
    index.limits['file_bytes'] = 50
    result = index.refresh()
    assert any(f['status'] == 'error' and 'size' in f['error'] for f in result['files'])
    assert index.search('Aria')['total'] == 1


def test_concurrent_indexers_claim_a_file_once(library, monkeypatch):
    root, _, index = library
    (root / 'notes.txt').write_text('Aria captain')
    started, release = threading.Event(), threading.Event()
    original = index._parse
    calls = []

    def slow(*args):
        calls.append(args)
        started.set()
        assert release.wait(5)
        return original(*args)

    monkeypatch.setattr(index, '_parse', slow)
    with ThreadPoolExecutor(2) as pool:
        first = pool.submit(index.refresh)
        assert started.wait(5)
        second = pool.submit(index.refresh)
        second.result(timeout=5)
        assert index.search('Aria')['results'] == []
        release.set()
        first.result(timeout=5)
    assert len(calls) == 1
    assert index.search('Aria')['total'] == 1


def test_pdf_text_and_no_text(library):
    from pypdf import PdfWriter

    root, _, index = library
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    with (root / 'scan.pdf').open('wb') as handle:
        writer.write(handle)
    result = index.refresh()
    assert result['files'][0]['status'] == 'no_text'
    assert result['chunks'] == 0


def test_invalid_ids_and_cache_inside_sources(library):
    root, registry, index = library
    with pytest.raises(ByaxError):
        registry.get('../../escape')
    with pytest.raises(ByaxError):
        LibraryIndex(index.library, root / 'cache')


def test_safe_docx_and_byax_parsers_preserve_locators(library):
    from docx import Document

    from byax.model import new_document
    from byax.storage import save_document

    root, _, index = library
    word = Document()
    word.add_heading('Research', 1)
    word.add_paragraph('Aria keeps the lighthouse.')
    word.save(root / 'research.docx')
    doc = new_document(text='Mira keeps the observatory.')
    save_document(doc, root / 'draft.byax')
    result = index.refresh()
    assert result['documents'] == 2
    assert index.search('Aria')['results'][0]['locator']['heading_path'] == ['Research']
    assert index.search('Mira')['results'][0]['locator']['block'] == doc['blocks'][0]['id']
    (root / 'broken.docx').write_bytes(b'not a zip file')
    result = index.refresh()
    assert any(f['path'] == 'broken.docx' and f['status'] == 'error' for f in result['files'])
    assert index.search('Mira')['total'] == 1


def test_read_locator_pagination_and_live_status(library):
    root, _, index = library
    path = root / 'notes.md'
    path.write_text('# One\n\nFirst Aria paragraph.\n\n# Two\n\nSecond Aria paragraph.')
    index.refresh()
    hit = index.search('Aria')['results'][0]
    first = index.read(hit['document_id'], limit=25)
    assert first['truncated']
    # A chunk never crosses a heading, so each section is its own chunk with its own ordinal.
    second = index.read(hit['document_id'], locator={'ordinal': 2})
    assert second['results'][0]['text'].startswith('Two')
    assert 'Second Aria paragraph.' in second['results'][0]['text']
    path.write_text('changed content')
    assert index.status()['stale']


def test_expired_lease_reclaimed_and_parser_deadline(library):
    root, _, index = library
    path = root / 'notes.txt'
    path.write_text('Aria captain')
    index.refresh()
    with index._db() as db:
        db.execute("UPDATE files SET lease='crashed',lease_until=0,status='indexing'")
    assert index.refresh()['processed'] == 1
    path.write_text('Aria changed captain')
    index.limits['parse_seconds'] = 0.00001
    result = index.refresh()
    assert result['files'][0]['status'] == 'error'
    assert 'parse time' in result['files'][0]['error']
    assert index.search('captain')['results'][0]['stale']


def test_refresh_reclaims_obsolete_chunks_before_enforcing_limit(library):
    root, _, index = library
    path = root / 'notes.txt'
    index.limits['chunks'] = 1
    for version in range(4):
        path.write_text(f'Aria version {version}')
        result = index.refresh()
        assert result['files'][0]['status'] == 'ready'
        assert result['chunks'] == 1
        assert str(version) in index.search('Aria')['results'][0]['text']
    path.unlink()
    assert index.refresh()['chunks'] == 0


def test_frozen_parser_dispatch_and_result_file(tmp_path, monkeypatch):
    import json
    import sys

    from byax.libraries import _parser_command
    from byax.library_parser import dispatch_worker

    source, output = tmp_path / 'source.txt', tmp_path / 'output.json'
    source.write_text('Aria lighthouse')
    monkeypatch.setattr(sys, 'frozen', True, raising=False)
    command = _parser_command(source, 100, output)
    assert command == [sys.executable, '--byax-library-parser', str(source), '100', str(output)]
    # Do not set process-wide resource limits inside the test runner.
    monkeypatch.setattr('byax.library_parser.sys.platform', 'win32')
    with pytest.raises(SystemExit) as result:
        dispatch_worker(command[1:])
    assert result.value.code == 0
    assert json.loads(output.read_text())['passages'][0]['text'] == 'Aria lighthouse'
