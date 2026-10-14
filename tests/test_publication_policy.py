"""Export profile inclusion categories, staged publication failures and serialized residue inspection."""

import copy
import os
import zipfile

import pytest

from byax.errors import ByaxError
from byax.exchange import export_document, import_document
from byax.model import new_document
from byax.publication import PUBLICATION_CATEGORIES, export_clean, inspect_residue, save_sanitized
from byax.storage import load_document

SECRET = 'Project Nightjar'


def _private():
    doc = new_document('Private master', 'Public intro.')
    target = doc['blocks'][0]['id']
    doc['blocks'].append(
        {
            'id': 'hidden',
            'type': 'paragraph',
            'text': 'Hidden ' + SECRET,
            'style': 'body',
            'marks': [],
            'attrs': {'excluded': True},
        }
    )
    doc['review']['comments'].append(
        {'id': 'c1', 'target': target, 'text': 'Comment ' + SECRET, 'status': 'open'}
    )
    doc['review']['tracked_changes'].append(
        {'id': 't1', 'target': target, 'before': SECRET, 'after': 'Public intro.', 'status': 'pending'}
    )
    doc['metadata']['author'] = 'Author ' + SECRET
    doc['references']['r1'] = {'title': 'Source', 'snapshot': SECRET}
    doc['settings']['transcripts'] = [{'text': SECRET}]
    return doc


def test_every_inclusion_category_is_independent_and_reported(tmp_path):
    doc = _private()
    for category in PUBLICATION_CATEGORIES:
        report = export_clean(doc, tmp_path / f'{category}.byax', profile={category: True})
        assert report['inclusion']['included'] == [category]
        assert category not in report['inclusion']['excluded']
    with_comments = load_document(tmp_path / 'include_comments.byax')
    assert with_comments['review']['comments'] and not with_comments['review']['tracked_changes']
    assert all(b['id'] != 'hidden' for b in with_comments['blocks'])
    hidden = load_document(tmp_path / 'include_hidden.byax')
    assert any(b['id'] == 'hidden' for b in hidden['blocks']) and not hidden['review']['comments']
    with pytest.raises(ByaxError) as error:
        export_clean(doc, tmp_path / 'x.byax', profile={'include_secrets': True})
    assert error.value.code == 'invalid_input'


def test_saved_clean_profile_routes_through_publication_policy(tmp_path):
    doc = _private()
    doc['profiles']['share'] = {'clean': True, 'format': 'docx', 'verify_absent': [SECRET]}
    report = export_document(doc, tmp_path / 'share.docx', profile='share')
    assert report['residue_inspection']['status'] == 'clean'
    assert report['inclusion']['included'] == [] and report['validation'] == 'not_run'
    assert inspect_residue(tmp_path / 'share.docx', [SECRET])['status'] == 'clean'


@pytest.mark.parametrize('fmt', ['docx', 'odt', 'html', 'md', 'rtf', 'txt', 'byax'])
def test_residue_inspection_blocks_publication_of_retained_terms(tmp_path, fmt):
    doc = _private()
    out = tmp_path / ('leak.' + fmt)
    with pytest.raises(ByaxError) as error:
        export_clean(doc, out, profile={'include_hidden': True, 'verify_absent': [SECRET]})
    assert error.value.code == 'failed_requirement' and error.value.details['locations']
    assert not out.exists() and not [p for p in tmp_path.iterdir() if p.name.startswith('.byax-')]


def test_residue_finds_terms_split_across_runs_and_reports_pdf_honestly(tmp_path):
    doc = new_document('Split', 'Night')
    doc['blocks'][0]['marks'] = [{'start': 0, 'end': 3, 'type': 'bold', 'value': True}]
    export_document(doc, tmp_path / 'split.docx')
    assert inspect_residue(tmp_path / 'split.docx', ['Night'])['status'] == 'found'
    export_document(doc, tmp_path / 'split.pdf')
    assert inspect_residue(tmp_path / 'split.pdf', ['Nightjar'])['status'] in ('unverifiable', 'clean')
    assert inspect_residue(tmp_path / 'split.pdf', ['Night'])['status'] in ('found', 'unverifiable')


def test_overwrite_refusal_and_interrupted_publication_keep_previous_deliverable(tmp_path, monkeypatch):
    doc = _private()
    out = tmp_path / 'deliverable.odt'
    export_clean(doc, out)
    before = out.read_bytes()
    with pytest.raises(ByaxError) as error:
        export_clean(doc, out, overwrite=False)
    assert error.value.code == 'conflict' and out.read_bytes() == before
    from byax import exchange

    def interrupted(candidate, stage):
        stage.write_bytes(b'PK\x03\x04half')
        raise KeyboardInterrupt

    monkeypatch.setattr(exchange, '_export_odt', interrupted)
    with pytest.raises(KeyboardInterrupt):
        export_clean(doc, out)
    with pytest.raises(KeyboardInterrupt):
        export_clean(doc, tmp_path / 'fresh.odt')
    assert out.read_bytes() == before and not (tmp_path / 'fresh.odt').exists()
    assert sorted(p.name for p in tmp_path.iterdir()) == ['deliverable.odt']


def test_complete_publication_and_redaction_scenario_preserves_private_master(tmp_path):
    from docx import Document

    word = Document()
    word.add_heading('Case report', 1)
    paragraph = word.add_paragraph('Witness ' + SECRET + ' was interviewed.')
    word.add_comment(paragraph.runs, text='Verify ' + SECRET, author='Editor')
    word.add_paragraph('Unrelated public conclusion.')
    word.core_properties.author = SECRET
    source = tmp_path / 'case.docx'
    word.save(source)
    original_bytes = source.read_bytes()
    master, _ = import_document(source)
    master['history']['revisions'].append({'revision': 'rev_old', 'snapshot': {'blocks': [{'text': SECRET}]}})
    private = copy.deepcopy(master)
    sanitized = tmp_path / 'sanitized.byax'
    report = save_sanitized(master, sanitized, [], terms=[SECRET], preserve=['styles', 'layout'])
    assert report['validation'] == 'passed' and report['coverage']['serialized_output'] == 'passed'
    clean = load_document(sanitized)
    assert any('Unrelated public conclusion.' == b.get('text') for b in clean['blocks'])
    for fmt in ('docx', 'odt', 'pdf', 'md'):
        published = export_clean(clean, tmp_path / ('public.' + fmt), profile={'verify_absent': [SECRET]})
        assert published['residue_inspection']['status'] in ('clean', 'unverifiable')
        with open(tmp_path / ('public.' + fmt), 'rb') as handle:
            assert SECRET.encode() not in handle.read()
    assert master == private and source.read_bytes() == original_bytes
    with zipfile.ZipFile(sanitized) as archive:
        assert all(SECRET.encode() not in archive.read(name) for name in archive.namelist())
    assert not os.path.exists(tmp_path / 'public.docx.lock')
